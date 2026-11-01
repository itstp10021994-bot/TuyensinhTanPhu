"""Phân hệ Nội trú: phòng, giáo viên quản nhiệm, đánh giá THEO THÁNG từng học sinh theo 6 mục.

Mục chọn mức (Tốt / Khá / Trung bình / Chưa tốt) kèm lời nhận xét: Nề nếp, Học tập,
Sức khỏe - Tâm lý. Mục viết lời: Điểm sáng trong tháng, Quá trình trao đổi và hỗ trợ,
Định hướng giáo dục.

- Phòng (list NoiTru_Phong): mỗi năm học, mỗi phòng có 1 giáo viên quản nhiệm và danh sách
  học sinh (ID hồ sơ nhập học).
- Đánh giá (list NoiTru_DanhGiaThang): mỗi bản ghi là MỘT HỌC SINH trong MỘT THÁNG, mỗi mục
  một cột (lọc / xem trực tiếp trên SharePoint được). List theo ngày cũ (NoiTru_DanhGia) không
  dùng nữa và được giữ nguyên.

Quyền: nhập/sửa đánh giá: giáo viên quản nhiệm của phòng (phòng mình), Tổ trưởng quản nhiệm và
Quản trị (mọi phòng); xem đánh giá: mọi người vào được
trang; quản lý phòng và báo cáo tổng hợp: Tổ trưởng quản nhiệm, Quản trị.
"""
from __future__ import annotations

from datetime import date

import pandas as pd

from .schema import MUC_DANH_GIA, NOI_TRU_DANH_GIA, NOI_TRU_PHONG
from .storage.base import Storage

DIEM = {"Tốt": 4, "Khá": 3, "Trung bình": 2, "Chưa tốt": 1}
CAN_CHU_Y = ("Trung bình", "Chưa tốt")
# Mục chọn mức (kèm lời nhận xét) — key trong JSON -> tên cột hiển thị
MUC = ("NeNep", "HocTap", "SucKhoe")
TEN_MUC = {"NeNep": "Nề nếp", "HocTap": "Học tập", "SucKhoe": "Sức khỏe - Tâm lý"}
NX = {m: f"{m}_NX" for m in MUC}
TEN_NX = {m: f"Nhận xét {TEN_MUC[m].lower()}" for m in MUC}
# Mục chỉ viết lời
LOI = ("DiemSang", "TraoDoi", "DinhHuong")
TEN_LOI = {"DiemSang": "Điểm sáng trong tháng", "TraoDoi": "Quá trình trao đổi và hỗ trợ",
           "DinhHuong": "Định hướng giáo dục"}
COT_MUC = [TEN_MUC[m] for m in MUC]
COT_BANG = (["Tháng", "PhongID", "Phòng", "HocSinhID", "Học sinh", "Lớp"]
            + [c for m in MUC for c in (TEN_MUC[m], TEN_NX[m])]
            + [TEN_LOI[k] for k in LOI] + ["Người đánh giá", "Ngày đánh giá"])
COT_LOI = [TEN_NX[m] for m in MUC] + [TEN_LOI[k] for k in LOI]
MAX_LOI = 1000  # ký tự mỗi ô lời nhận xét (cột nhiều dòng SharePoint)
VAI_QUAN_LY = ("Quản trị", "Tổ trưởng quản nhiệm")


# ------------------------------------------------------------------ quyền
def la_quan_ly(tk: dict | None) -> bool:
    """Tổ trưởng quản nhiệm / Quản trị (app chưa dùng tài khoản: ai cũng là quản lý)."""
    return not tk or tk.get("VaiTro") in VAI_QUAN_LY


def duoc_nhap(phong: dict, tk: dict | None) -> bool:
    """Giáo viên quản nhiệm nhập phòng mình; Tổ trưởng quản nhiệm / Quản trị nhập mọi phòng."""
    if la_quan_ly(tk):
        return True
    if tk.get("id") and str(phong.get("GiaoVienID") or "") == str(tk["id"]):
        return True
    # phòng nhập từ Excel trước khi giáo viên có tài khoản: khớp theo họ tên
    from .danh_muc import fold
    return (not _s(phong.get("GiaoVienID")) and bool(_s(phong.get("GiaoVien")))
            and fold(phong.get("GiaoVien")) == fold(tk.get("HoTen")))


# ------------------------------------------------------------------ phòng
def _s(v) -> str:
    return "" if v is None or (isinstance(v, float) and v != v) else str(v).strip()


def hs_ids(phong: dict) -> list[str]:
    return [x.strip() for x in _s(phong.get("HocSinhIDs")).split(";") if x.strip()]


def phong_nam(phongs: list[dict], nam_hoc: str) -> list[dict]:
    ds = [p for p in phongs if _s(p.get("NamHoc")) == nam_hoc]
    return sorted(ds, key=lambda p: (_s(p.get("Khu")), _so_tu_nhien(_s(p.get("TenPhong")))))


def _so_tu_nhien(s: str) -> tuple:
    import re
    return tuple(int(x) if x.isdigit() else x.lower() for x in re.split(r"(\d+)", s))


def phong_cua_hs(phongs: list[dict]) -> dict[str, dict]:
    return {i: p for p in phongs for i in hs_ids(p)}


def luu_phong(storage: Storage, phongs: list[dict], nam_hoc: str, ten: str, khu: str,
              gv: dict | None, hoc_sinh: list[str], ghi_chu: str = "",
              item_id: str | None = None) -> dict:
    """Thêm / sửa phòng. Tên phòng không trùng trong năm; mỗi học sinh chỉ ở một phòng."""
    ten = _s(ten)
    if not ten:
        raise ValueError("Cần nhập **Tên phòng**.")
    cung_nam = [p for p in phongs if _s(p.get("NamHoc")) == nam_hoc
                and str(p.get("id")) != str(item_id)]
    if any(_s(p.get("TenPhong")).lower() == ten.lower() for p in cung_nam):
        raise ValueError(f"Phòng **{ten}** đã có trong năm học {nam_hoc}.")
    o_phong_khac = phong_cua_hs(cung_nam)
    trung = [i for i in hoc_sinh if i in o_phong_khac]
    if trung:
        ten_p = sorted({_s(o_phong_khac[i].get("TenPhong")) for i in trung})
        raise ValueError(f"{len(trung)} học sinh đang ở phòng khác ({', '.join(ten_p)}). "
                         "Bỏ các em khỏi phòng cũ trước.")
    data = {"NamHoc": nam_hoc, "TenPhong": ten, "Khu": _s(khu),
            "GiaoVienID": _s((gv or {}).get("id")), "GiaoVien": _s((gv or {}).get("HoTen")),
            "HocSinhIDs": ";".join(dict.fromkeys(hoc_sinh)), "GhiChu": _s(ghi_chu)}
    if item_id:
        return storage.update_item(NOI_TRU_PHONG.name, item_id, data)
    return storage.create_item(NOI_TRU_PHONG.name, data)


def chuyen_ban_tru(storage: Storage, nh: list[dict], ts: list[dict], hoc_sinh: list[str],
                   progress=None) -> int:
    """Chuyển học sinh sang Bán trú: hồ sơ nhập học (N.trú, B.trú) và liên hệ Data tuyển sinh
    liên kết (Chế độ)."""
    from concurrent.futures import ThreadPoolExecutor

    from .schema import NHAP_HOC, TUYEN_SINH
    chon = {str(i) for i in hoc_sinh}
    ts_ids = {str(t.get("id")) for t in ts}
    viec = []
    for r in nh:
        if str(r.get("id")) not in chon:
            continue
        viec.append((NHAP_HOC.name, r["id"], {"NoiTruBanTru": "Bán trú"}))
        tid = _s(r.get("TuyenSinhID"))
        if tid in ts_ids:
            viec.append((TUYEN_SINH.name, tid, {"CheDo": "Bán trú"}))
    done = 0
    with ThreadPoolExecutor(max_workers=4) as pool:
        for _ in pool.map(lambda v: storage.update_item(*v), viec):
            done += 1
            if progress:
                progress(done, len(viec))
    return len(chon)


def bo_khoi_phong(storage: Storage, phongs: list[dict], hoc_sinh: list[str]) -> int:
    """Bỏ các học sinh khỏi mọi phòng đang chứa các em."""
    bo, n = {str(i) for i in hoc_sinh}, 0
    for p in phongs:
        con = [i for i in hs_ids(p) if i not in bo]
        if len(con) != len(hs_ids(p)):
            storage.update_item(NOI_TRU_PHONG.name, p["id"], {"HocSinhIDs": ";".join(con)})
            n += 1
    return n


def them_vao_phong(storage: Storage, phongs: list[dict], nam_hoc: str, phong: dict | None,
                   ten_moi: str, hoc_sinh: list[str]) -> dict:
    """Xếp thêm học sinh vào phòng có sẵn (giữ giáo viên, học sinh cũ) hoặc phòng mới."""
    if not hoc_sinh:
        raise ValueError("Chưa chọn học sinh nào.")
    if not phong:
        return luu_phong(storage, phongs, nam_hoc, ten_moi, "", None, hoc_sinh)
    gv = {"id": _s(phong.get("GiaoVienID")), "HoTen": _s(phong.get("GiaoVien"))}
    return luu_phong(storage, phongs, nam_hoc, _s(phong.get("TenPhong")), _s(phong.get("Khu")),
                     gv, hs_ids(phong) + [i for i in hoc_sinh if i not in hs_ids(phong)],
                     _s(phong.get("GhiChu")), phong.get("id"))


# ------------------------------------------------------------------ đánh giá theo tháng
def thang_cua_nam(nam_hoc: str) -> list[str]:
    """Các tháng của năm học "2026-2027": 2026-08 … 2027-07 (dạng YYYY-MM)."""
    try:
        y = int(str(nam_hoc)[:4])
    except ValueError:
        y = date.today().year
    return [f"{y + (m < 8)}-{m:02d}" for m in (*range(8, 13), *range(1, 8))]


def thang_mac_dinh(nam_hoc: str, hom_nay: date | None = None) -> str:
    """Tháng hiện tại nếu thuộc năm học; năm học đã qua -> tháng cuối, chưa tới -> tháng đầu."""
    ds = thang_cua_nam(nam_hoc)
    t = f"{(hom_nay or date.today()):%Y-%m}"
    return t if t in ds else (ds[-1] if t > ds[-1] else ds[0])


def ten_thang(t: str) -> str:
    """'2026-10' -> 'Tháng 10/2026'."""
    t = _s(t)
    return f"Tháng {int(t[5:7])}/{t[:4]}" if len(t) >= 7 and t[4] == "-" else t


def chuan_thang(v) -> str:
    """Đọc tháng từ ô Excel: '2026-10', '10/2026', 'Tháng 10/2026', ngày bất kỳ trong tháng."""
    import re
    if isinstance(v, (date, pd.Timestamp)):
        return f"{v:%Y-%m}"
    s = _s(v).lower().removeprefix("tháng").strip()
    m = re.fullmatch(r"(\d{4})[-/.](\d{1,2})(?:[-/.]\d{1,2})?(?:[ t].*)?", s)
    if m and 1 <= int(m[2]) <= 12:
        return f"{m[1]}-{int(m[2]):02d}"
    m = re.fullmatch(r"(?:\d{1,2}[-/.])?(\d{1,2})[-/.](\d{4})", s)
    if m and 1 <= int(m[1]) <= 12:
        return f"{m[2]}-{int(m[1]):02d}"
    return ""


def tim_danh_gia(danh_gia: list[dict], hs_id: str, thang: str,
                 nam_hoc: str | None = None) -> dict | None:
    for r in danh_gia:
        if (str(r.get("HocSinhID")) == str(hs_id) and _s(r.get("Thang")) == thang
                and (not nam_hoc or _s(r.get("NamHoc")) == nam_hoc)):
            return r
    return None


def can_chu_y(r: dict) -> bool:
    """Có mục nào Trung bình / Chưa tốt."""
    return any(_s(r.get(m)) in CAN_CHU_Y for m in MUC)


def luu_danh_gia_thang(storage: Storage, danh_gia: list[dict], nam_hoc: str, thang: str,
                       hs: dict, phong: dict | None, values: dict, nguoi: str) -> dict:
    """Ghi đánh giá tháng của MỘT học sinh (tạo mới hoặc ghi đè bản ghi tháng đó).

    `hs`: {"id", "ten", "lop"}; `values`: {<mục>: mức, <mục>_NX: lời, DiemSang/TraoDoi/DinhHuong}"""
    hid = _s(hs.get("id"))
    if not hid:
        raise ValueError("Thiếu học sinh cần đánh giá.")
    if thang not in thang_cua_nam(nam_hoc):
        raise ValueError(f"Tháng {thang} không thuộc năm học {nam_hoc}.")
    data = {"NamHoc": nam_hoc, "Thang": thang, "HocSinhID": hid, "HoTen": _s(hs.get("ten")),
            "Lop": _s(hs.get("lop")), "PhongID": _s((phong or {}).get("id")),
            "Phong": _s((phong or {}).get("TenPhong")), "NguoiDanhGia": _s(nguoi),
            "NgayDanhGia": date.today().isoformat()}
    for m in MUC:
        v = _s(values.get(m)) or "Tốt"
        if v not in MUC_DANH_GIA:
            raise ValueError(f"Mức đánh giá không hợp lệ: {v}")
        data[m] = v
    for k in [NX[m] for m in MUC] + list(LOI):
        v = _s(values.get(k))
        if len(v) > MAX_LOI:
            ten = TEN_LOI.get(k) or TEN_NX[k.removesuffix("_NX")]
            raise ValueError(f"{ten} quá dài ({len(v)} ký tự, tối đa {MAX_LOI}).")
        data[k] = v
    cu = tim_danh_gia(danh_gia, hid, thang, nam_hoc)
    if cu:
        return storage.update_item(NOI_TRU_DANH_GIA.name, cu["id"], data)
    return storage.create_item(NOI_TRU_DANH_GIA.name, data)


def danh_gia_nhanh(storage: Storage, danh_gia: list[dict], nam_hoc: str, thang: str,
                   ds: list[tuple[dict, dict | None]], nguoi: str) -> int:
    """Đánh giá "Tốt" cả 3 mục cho các em CHƯA có đánh giá tháng `thang` (em đã có: bỏ qua).
    `ds`: [(hs {"id","ten","lop"}, phòng)]."""
    from concurrent.futures import ThreadPoolExecutor

    viec = [(h, p) for h, p in ds if not tim_danh_gia(danh_gia, h.get("id"), thang, nam_hoc)]
    with ThreadPoolExecutor(max_workers=4) as pool:
        list(pool.map(lambda hp: luu_danh_gia_thang(storage, [], nam_hoc, thang, hp[0], hp[1],
                                                    {}, nguoi), viec))
    return len(viec)


def nhap_danh_gia(storage: Storage, df: pd.DataFrame, phongs_all: list[dict],
                  danh_gia: list[dict], nam_hoc: str, nh: list[dict], nguoi: str,
                  progress=None) -> dict:
    """Nhập đánh giá tháng từ file mẫu (mỗi dòng = 1 học sinh trong 1 tháng; cột theo key:
    Thang, Phong, HoTen, Lop, <mục>, <mục>_NX, DiemSang / TraoDoi / DinhHuong, NguoiDanhGia).

    Học sinh tìm trong phòng ghi ở cột Phòng (để trống: mọi phòng của năm học) theo họ tên,
    trùng tên thì xét lớp. Học sinh + tháng đã có đánh giá: ghi đè theo file (ô trống giữ cũ)."""
    from .danh_muc import fold

    ps = phong_nam(phongs_all, nam_hoc)
    phongs = {_s(p.get("TenPhong")).lower(): p for p in ps}
    o = phong_cua_hs(ps)
    hs = {str(r.get("id")): r for r in nh}
    loi, viec = [], {}
    for i, row in df.iterrows():
        d = {k: _s(v) for k, v in row.to_dict().items()}
        dong = i + 2
        thang = chuan_thang(row.get("Thang"))
        if thang not in thang_cua_nam(nam_hoc):
            loi.append({"Dòng": dong, "Lỗi": f"Tháng không hợp lệ hoặc ngoài năm học {nam_hoc}: "
                                             f"\"{d.get('Thang', '')}\" (ghi vd 10/2026)"})
            continue
        ten_p = d.get("Phong", "").lower().removeprefix("phòng ").strip()
        p = phongs.get(ten_p) if ten_p else None
        if ten_p and p is None:
            loi.append({"Dòng": dong, "Lỗi": f"Không có phòng \"{d.get('Phong', '')}\" trong "
                                             f"năm học {nam_hoc}"})
            continue
        cho = hs_ids(p) if p else list(o)
        ung = [x for x in cho if fold((hs.get(x) or {}).get("HoTen")) == fold(d.get("HoTen"))]
        if len(ung) > 1 and d.get("Lop"):
            ung = [x for x in ung if _chuan_lop(hs[x].get("LopHoc")) == _chuan_lop(d["Lop"])] or ung
        noi = f"phòng {p.get('TenPhong')}" if p else "các phòng nội trú"
        if len(ung) != 1:
            loi.append({"Dòng": dong, "Lỗi": (f"{noi.capitalize()} có {len(ung)} em tên "
                                              f"\"{d.get('HoTen')}\" — ghi thêm lớp") if ung else
                        f"Không thấy \"{d.get('HoTen', '')}\" trong {noi}"})
            continue
        sai = [d[m] for m in MUC if d.get(m) and d[m] not in MUC_DANH_GIA]
        if sai:
            loi.append({"Dòng": dong, "Lỗi": f"Mức không hợp lệ: {sai[0]} (chọn: "
                                             f"{', '.join(MUC_DANH_GIA)})"})
            continue
        if any(len(d.get(k, "")) > MAX_LOI for k in [NX[m] for m in MUC] + list(LOI)):
            loi.append({"Dòng": dong, "Lỗi": f"Nhận xét quá dài (tối đa {MAX_LOI} ký tự)"})
            continue
        viec[(ung[0], thang)] = (dong, d, p or o.get(ung[0]))
    danh_gia = list(danh_gia)
    ok = 0
    for n, ((hid, thang), (dong, d, p)) in enumerate(viec.items(), 1):
        r = hs[hid]
        cu = tim_danh_gia(danh_gia, hid, thang, nam_hoc) or {}
        vals = {k: d.get(k) or _s(cu.get(k)) for k in list(MUC) + [NX[m] for m in MUC] + list(LOI)}
        try:
            luu_danh_gia_thang(storage, danh_gia, nam_hoc, thang,
                               {"id": hid, "ten": r.get("HoTen"), "lop": r.get("LopHoc")}, p,
                               vals, d.get("NguoiDanhGia") or nguoi)
            ok += 1
        except Exception as e:
            loi.append({"Dòng": dong, "Lỗi": str(e)[:300]})
        if progress:
            progress(n, len(viec))
    return {"ok": ok, "ban_ghi": ok, "errors": sorted(loi, key=lambda e: e["Dòng"]),
            "total": len(df)}


def bang_thang(danh_gia: list[dict], nam_hoc: str | None = None) -> pd.DataFrame:
    """Mỗi dòng = một học sinh trong một tháng (cột hiển thị COT_BANG)."""
    rows = []
    for r in danh_gia:
        if nam_hoc and _s(r.get("NamHoc")) != nam_hoc:
            continue
        thang = _s(r.get("Thang"))
        if not thang or not _s(r.get("HocSinhID")):
            continue
        row = {"Tháng": thang, "PhongID": _s(r.get("PhongID")), "Phòng": _s(r.get("Phong")),
               "HocSinhID": _s(r.get("HocSinhID")), "Học sinh": _s(r.get("HoTen")),
               "Lớp": _s(r.get("Lop"))}
        for m in MUC:
            row[TEN_MUC[m]] = _s(r.get(m))
            row[TEN_NX[m]] = _s(r.get(NX[m]))
        for k in LOI:
            row[TEN_LOI[k]] = _s(r.get(k))
        row["Người đánh giá"] = _s(r.get("NguoiDanhGia"))
        row["Ngày đánh giá"] = _s(r.get("NgayDanhGia"))[:10]
        rows.append(row)
    return pd.DataFrame(rows, columns=COT_BANG)


def theo_ho_so(long: pd.DataFrame, nh: list[dict]) -> pd.DataFrame:
    """Tên và lớp lấy theo hồ sơ nhập học hiện tại (đánh giá lưu tên / lớp lúc đánh giá); bỏ
    đánh giá của học sinh đã bị xóa hồ sơ."""
    hs = {str(r.get("id")): r for r in nh}
    out = long[long["HocSinhID"].isin(hs)].copy()
    if len(out):
        out["Học sinh"] = [_s(hs[i].get("HoTen")) or t for i, t in
                           zip(out["HocSinhID"], out["Học sinh"])]
        out["Lớp"] = [_s(hs[i].get("LopHoc")) or (f"Khối {_s(hs[i].get('Khoi'))}"
                                                 if _s(hs[i].get("Khoi")) else l)
                     for i, l in zip(out["HocSinhID"], out["Lớp"])]
    return out


# ------------------------------------------------------------------ báo cáo
def _ty_le_tot(s: pd.Series) -> float | None:
    """% "Tốt"; chưa có đánh giá nào thì None (hiện "—", không phải 0%)."""
    s = s[s.fillna("") != ""]
    return round(100 * (s == "Tốt").mean(), 1) if len(s) else None


def tong_hop_phong(x: pd.DataFrame, phongs: list[dict], dang_nt: set[str]) -> pd.DataFrame:
    """Một tháng, theo phòng: số em đã / chưa đánh giá, % Tốt từng mục, số em cần chú ý.
    `dang_nt`: ID học sinh nội trú đang học (em đã rời nội trú không tính)."""
    out = []
    for p in phongs:
        ids = [i for i in hs_ids(p) if i in dang_nt]
        y = x[x["HocSinhID"].isin(ids)]
        da = y["HocSinhID"].nunique()
        r = {"Phòng": _s(p.get("TenPhong")), "Giáo viên": _s(p.get("GiaoVien")),
             "Số HS": len(ids), "Đã đánh giá": da, "Chưa đánh giá": max(len(ids) - da, 0),
             "Hoàn thành (%)": round(100 * da / len(ids), 1) if ids else None}
        for c in COT_MUC:
            r[f"{c} Tốt (%)"] = _ty_le_tot(y[c])
        r["Em cần chú ý"] = int(y[COT_MUC].isin(CAN_CHU_Y).any(axis=1).sum())
        out.append(r)
    return pd.DataFrame(out)


def hoc_sinh_can_chu_y(x: pd.DataFrame) -> pd.DataFrame:
    """Học sinh có mục Trung bình / Chưa tốt (mỗi lượt học sinh-tháng một dòng)."""
    cols = ["Tháng", "Học sinh", "Lớp", "Phòng", *COT_MUC, "Nhận xét"]
    if x.empty:
        return pd.DataFrame(columns=cols)
    y = x[x[COT_MUC].isin(CAN_CHU_Y).any(axis=1)].copy()
    if y.empty:
        return pd.DataFrame(columns=cols)
    y["Nhận xét"] = y[[TEN_NX[m] for m in MUC]].apply(lambda r: " · ".join(v for v in r if v),
                                                      axis=1)
    y["_d"] = [sum(DIEM.get(r, 4) for r in t) for t in zip(*(y[c] for c in COT_MUC))]
    y = y.sort_values(["Tháng", "_d", "Phòng", "Học sinh"], ascending=[False, True, True, True])
    y["Tháng"] = y["Tháng"].map(ten_thang)
    return y[cols].reset_index(drop=True)


def nhan_xet(x: pd.DataFrame) -> pd.DataFrame:
    """Các lượt có lời nhận xét (để đọc nhanh)."""
    cols = ["Tháng", "Phòng", "Học sinh", "Lớp"] + COT_LOI
    if x.empty:
        return pd.DataFrame(columns=cols)
    return x.loc[x[COT_LOI].fillna("").ne("").any(axis=1), cols]


def xu_huong(x: pd.DataFrame, thangs: list[str]) -> pd.DataFrame:
    """% Tốt của từng mục theo tháng (chỉ tháng đã có đánh giá)."""
    rows = []
    for t in thangs:
        y = x[x["Tháng"] == t]
        if y.empty:
            continue
        for c in COT_MUC:
            rows.append({"Tháng": ten_thang(t).removeprefix("Tháng "), "Mục": c,
                         "Tỷ lệ Tốt (%)": _ty_le_tot(y[c])})
    return pd.DataFrame(rows, columns=["Tháng", "Mục", "Tỷ lệ Tốt (%)"])


def tien_do_thang(x: pd.DataFrame, thangs: list[str], n_hs: int) -> pd.DataFrame:
    """Số học sinh đã đánh giá mỗi tháng (đến tháng đang xem)."""
    return pd.DataFrame([{"Tháng": ten_thang(t).removeprefix("Tháng "),
                          "Đã đánh giá": int(x.loc[x["Tháng"] == t, "HocSinhID"].nunique()),
                          "Số HS": n_hs} for t in thangs])


def phan_bo_muc(x: pd.DataFrame) -> pd.DataFrame:
    """Số học sinh ở từng mức của từng mục."""
    rows = []
    for col in COT_MUC:
        vc = x[col].value_counts()
        for m in MUC_DANH_GIA:
            rows.append({"Mục": col, "Mức": m, "Số HS": int(vc.get(m, 0))})
    return pd.DataFrame(rows)


# ------------------------------------------------------------------ nhập phân phòng từ Excel
TT_KHOP, TT_KHONG_THAY, TT_TRUNG = "Khớp", "Không tìm thấy", "Trùng tên, chưa rõ em nào"


def _chuan_lop(v) -> str:
    """9.3 / "9,3" / "10A1ESL" / "10A1E" -> dạng so khớp: "9.3", "10A1" (bỏ đuôi E/ESL)."""
    import re
    if v is None or (isinstance(v, float) and v != v):
        return ""
    t = (f"{v:g}" if isinstance(v, float) else str(v)).strip().upper().replace(",", ".")
    t = re.sub(r"\s+", "", t)
    return re.sub(r"(ESL|E|IEP)$", "", t)


def _khoi_lop(lop: str) -> str:
    import re
    m = re.match(r"^(\d{1,2})", lop or "")
    return m.group(1) if m else ""


def _tim_cot(nhan: list[str]) -> dict | None:
    def tim(*dk):
        for f in dk:
            j = next((j for j, v in enumerate(nhan) if v and f(v)), None)
            if j is not None:
                return j
        return None
    c = {"ten": tim(lambda v: v in ("ho va ten", "ho ten", "ho ten hs")),
         "phong": tim(lambda v: v.startswith("phong") and ("nt" in v.split() or "noi tru" in v),
                      lambda v: v.startswith("phong") and "hoc" not in v),
         "lop": tim(lambda v: v == "lop", lambda v: v == "lop moi"),
         "gt": tim(lambda v: v.startswith("gioi tinh")),
         "gv": tim(lambda v: v.startswith("qn") or "quan ly" in v or "quan nhiem" in v,
                   lambda v: v.startswith("giao vien"))}
    return c if c["ten"] is not None and c["phong"] is not None else None


def sheet_phong(src) -> list[str]:
    """Các sheet có cột Họ và tên + Phòng (nội trú), sheet đầy đủ cột nhất đứng trước."""
    import openpyxl
    from .danh_muc import fold

    wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
    ds = []
    for ws in wb.worksheets:
        for r in ws.iter_rows(min_row=1, max_row=12, values_only=True):
            c = _tim_cot([fold(v).replace("\n", " ") if v is not None else "" for v in r])
            if c:
                ds.append((-sum(v is not None for v in c.values()), len(ds), ws.title))
                break
    wb.close()
    return [t for *_, t in sorted(ds)]


def doc_ds_phong(src, sheet: str | None = None) -> tuple[pd.DataFrame, str]:
    """Đọc danh sách nội trú: cột Họ và tên, Lớp, Giới tính, Phòng (NT), Giáo viên quản lý.

    Trả về (bảng: Họ tên, Lớp, Giới tính, Phòng, Giáo viên; tên sheet đã đọc)."""
    import openpyxl
    from .danh_muc import fold

    ds = sheet_phong(src)
    if not ds:
        raise ValueError("Không tìm thấy sheet có cột \"Họ và tên\" và cột \"Phòng\".")
    sheet = sheet if sheet in ds else ds[0]
    if hasattr(src, "seek"):
        src.seek(0)
    wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
    ws = wb[sheet]
    c, tieu_de = None, 0
    for i, r in enumerate(ws.iter_rows(min_row=1, max_row=12, values_only=True), 1):
        c = _tim_cot([fold(v).replace("\n", " ") if v is not None else "" for v in r])
        if c:
            tieu_de = i
            break
    out = []
    for r in ws.iter_rows(min_row=tieu_de + 1, values_only=True):
        def lay(k):
            j = c[k]
            return r[j] if j is not None and j < len(r) else None
        ten = _s(lay("ten"))
        phong = lay("phong")
        phong = f"{phong:g}" if isinstance(phong, float) else _s(phong)
        if not ten or not phong:  # chỉ lấy học sinh có phòng nội trú
            continue
        out.append({"Họ tên": " ".join(ten.split()), "Lớp": _chuan_lop(lay("lop")),
                    "Giới tính": _s(lay("gt")), "Phòng": phong,
                    "Giáo viên": " ".join(_s(lay("gv")).split())})
    wb.close()
    return pd.DataFrame(out, columns=["Họ tên", "Lớp", "Giới tính", "Phòng", "Giáo viên"]), sheet


def ghep_hs(ds: pd.DataFrame, nh: list[dict], nam_hoc: str) -> pd.DataFrame:
    """Ghép từng dòng danh sách với hồ sơ nhập học của năm: theo họ tên, rồi khối / lớp /
    giới tính khi trùng tên. Thêm cột Kết quả, HocSinhID, Lớp (app), Chế độ (app)."""
    from .danh_muc import fold

    ung = [r for r in nh if _s(r.get("NamHoc")) == nam_hoc
           and _s(r.get("TinhTrangHS")) not in ("Rút hồ sơ", "Không học tiếp")]
    theo_ten: dict[str, list[dict]] = {}
    for r in ung:
        theo_ten.setdefault(fold(r.get("HoTen")), []).append(r)
    kq = []
    for _, d in ds.iterrows():
        cands = theo_ten.get(fold(d["Họ tên"]), [])
        lop, khoi = d["Lớp"], _khoi_lop(d["Lớp"])
        for loc in (
                lambda c: _chuan_lop(c.get("LopHoc")) == lop and lop,
                lambda c: (_s(c.get("Khoi")) == khoi or _khoi_lop(_chuan_lop(c.get("LopHoc")))
                           == khoi) and khoi,
                lambda c: fold(c.get("GioiTinh")) == fold(d["Giới tính"]) and d["Giới tính"],
                # vẫn trùng (vd 2 em cùng tên, cùng lớp): chọn em đang là Nội trú
                lambda c: _s(c.get("NoiTruBanTru")).startswith("Nội trú")):
            if len(cands) > 1:
                hep = [c for c in cands if loc(c)]
                cands = hep or cands
        if len(cands) == 1:
            c = cands[0]
            kq.append((TT_KHOP, str(c["id"]), _s(c.get("LopHoc")) or _s(c.get("Khoi")),
                       _s(c.get("NoiTruBanTru"))))
        else:
            kq.append((TT_TRUNG if cands else TT_KHONG_THAY, "", "", ""))
    out = ds.copy()
    out["Kết quả"] = [k[0] for k in kq]
    out["HocSinhID"] = [k[1] for k in kq]
    out["Lớp (app)"] = [k[2] for k in kq]
    out["Chế độ (app)"] = [k[3] for k in kq]
    # một hồ sơ bị ghép cho 2 dòng -> không chắc, bỏ ghép
    trung_id = out["HocSinhID"][out["HocSinhID"] != ""].duplicated(keep=False)
    for i in trung_id[trung_id].index:
        out.loc[i, ["Kết quả", "HocSinhID"]] = [TT_TRUNG, ""]
    return out


def ke_hoach_phong(ghep: pd.DataFrame, tks: list[dict]) -> list[dict]:
    """Mỗi phòng: tên, giáo viên (nhiều tên thì lấy tên xuất hiện nhiều nhất), tài khoản khớp
    theo họ tên, danh sách HocSinhID đã ghép."""
    from .danh_muc import fold

    theo_ten = {fold(t.get("HoTen")): t for t in tks if _s(t.get("HoTen"))}
    out = []
    for phong, g in ghep.groupby("Phòng", sort=False):
        if not _s(phong):
            continue
        gv_ds = g["Giáo viên"][g["Giáo viên"] != ""]
        gv = gv_ds.value_counts().index[0] if len(gv_ds) else ""
        out.append({"TenPhong": _s(phong), "GiaoVien": gv, "TaiKhoan": theo_ten.get(fold(gv)),
                    "NhieuGV": sorted(set(gv_ds)) if gv_ds.nunique() > 1 else [],
                    "HocSinhIDs": [i for i in g["HocSinhID"] if i]})
    return sorted(out, key=lambda x: _so_tu_nhien(x["TenPhong"]))


def ap_dung_phan_phong(storage: Storage, phongs_all: list[dict], nam_hoc: str,
                       ghep: pd.DataFrame, ke_hoach: list[dict], nh_name: str,
                       progress=None) -> dict:
    """Ghi lên SharePoint: hồ sơ ghép được -> "Nội trú"; tạo / cập nhật phòng theo file (học
    sinh của phòng = theo file); em có trong file bị bỏ khỏi phòng khác của năm."""
    from concurrent.futures import ThreadPoolExecutor

    doi_che_do = ghep[(ghep["HocSinhID"] != "")
                      & ~ghep["Chế độ (app)"].str.startswith("Nội trú")]["HocSinhID"].tolist()
    cung_nam = [p for p in phongs_all if _s(p.get("NamHoc")) == nam_hoc]
    theo_ten = {_s(p.get("TenPhong")).lower(): p for p in cung_nam}
    trong_file = {i for k in ke_hoach for i in k["HocSinhIDs"]}
    viec = [("nh", i) for i in doi_che_do]
    phong_file = set()
    for k in ke_hoach:
        cu = theo_ten.get(k["TenPhong"].lower())
        phong_file.add(k["TenPhong"].lower())
        tk = k["TaiKhoan"] or {}
        data = {"NamHoc": nam_hoc, "TenPhong": k["TenPhong"],
                "GiaoVienID": _s(tk.get("id")), "GiaoVien": _s(tk.get("HoTen")) or k["GiaoVien"],
                "HocSinhIDs": ";".join(dict.fromkeys(k["HocSinhIDs"]))}
        if cu:
            data["Khu"] = _s(cu.get("Khu"))
        viec.append(("phong", (cu.get("id") if cu else None, data)))
    for p in cung_nam:  # bỏ học sinh có trong file khỏi các phòng khác
        if _s(p.get("TenPhong")).lower() in phong_file:
            continue
        con = [i for i in hs_ids(p) if i not in trong_file]
        if len(con) != len(hs_ids(p)):
            viec.append(("phong", (p["id"], {"HocSinhIDs": ";".join(con)})))

    def lam(v):
        loai, x = v
        if loai == "nh":
            return storage.update_item(nh_name, x, {"NoiTruBanTru": "Nội trú"})
        item_id, data = x
        if item_id:
            return storage.update_item(NOI_TRU_PHONG.name, item_id, data)
        return storage.create_item(NOI_TRU_PHONG.name, data)

    done = 0
    with ThreadPoolExecutor(max_workers=4) as pool:
        for _ in pool.map(lam, viec):
            done += 1
            if progress:
                progress(done, len(viec))
    return {"doi_che_do": len(doi_che_do), "phong": len(ke_hoach),
            "xep": len(trong_file), "viec": len(viec)}
