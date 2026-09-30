"""Xuất danh sách HS nhập học theo đúng biểu mẫu "DANH SÁCH HỌC SINH" (import VEMIS)."""
from __future__ import annotations

import io
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from . import danh_muc

# (nhóm tiêu đề dòng 7, tiêu đề dòng 8, key trong NHAP_HOC). Nhóm None = gộp 2 dòng.
COLUMNS: list[tuple[str | None, str, str | None]] = [
    (None, "STT", None),
    (None, "Lớp học", "LopHoc"),
    (None, "Mã học sinh", "MaHocSinh"),
    (None, "Mã VEMIS", "MaVEMIS"),
    (None, "Mã MOET", "MaMOET"),
    (None, "Sổ đăng bộ", "SoDangBo"),
    (None, "Họ và tên", "HoTen"),
    (None, "Ngày sinh", "NgaySinh"),
    (None, "Ngày vào trường", "NgayVaoTruong"),
    (None, "Giới tính", "GioiTinh"),
    (None, "Quốc tịch", "QuocTich"),
    ("Chỗ ở hiện nay", "SN/Xóm", "ChoO_SoNha"),
    ("Chỗ ở hiện nay", "Khu dân cư", "ChoO_KhuDanCu"),
    ("Chỗ ở hiện nay", "Xã/Phường", "ChoO_Xa"),
    ("Chỗ ở hiện nay", "Tỉnh/Tp", "ChoO_Tinh"),
    ("Hộ khẩu thường trú", "SN/Xóm", "HK_SoNha"),
    ("Hộ khẩu thường trú", "Khu dân cư", "HK_KhuDanCu"),
    ("Hộ khẩu thường trú", "Xã/Phường", "HK_Xa"),
    ("Hộ khẩu thường trú", "Tỉnh/Tp", "HK_Tinh"),
    ("Nơi sinh", "Thông tin nơi sinh", "NoiSinh_ThongTin"),
    ("Nơi sinh", "Xã/Phường", "NoiSinh_Xa"),
    ("Nơi sinh", "Tỉnh/Tp", "NoiSinh_Tinh"),
    ("Quê quán", "Thông tin quê quán", "QueQuan_ThongTin"),
    ("Quê quán", "Xã/Phường", "QueQuan_Xa"),
    ("Quê quán", "Tỉnh/Tp", "QueQuan_Tinh"),
    ("Nơi khai sinh", "Xã/Phường", "NoiKS_Xa"),
    ("Nơi khai sinh", "Tỉnh/Tp", "NoiKS_Tinh"),
    (None, "Căn cước", "CanCuoc"),
    (None, "Ngày cấp Căn cước", "NgayCapCanCuoc"),
    (None, "Nơi cấp Căn cước", "NoiCapCanCuoc"),
    (None, "Dân tộc", "DanToc"),
    (None, "Tôn giáo", "TonGiao"),
    (None, "Diện chính sách", "DienChinhSach"),
    (None, "Cận nghèo", "CanNgheo"),
    (None, "Đoàn viên", "DoanVien"),
    (None, "Đội viên", "DoiVien"),
    (None, "Tên cha", "TenCha"),
    (None, "Nghề nghiệp cha", "NgheNghiepCha"),
    (None, "Năm sinh cha", "NamSinhCha"),
    (None, "Căn cước cha", "CanCuocCha"),
    (None, "Đơn vị công tác cha", "DonViCongTacCha"),
    (None, "Tên mẹ", "TenMe"),
    (None, "Nghề nghiệp mẹ", "NgheNghiepMe"),
    (None, "Năm sinh mẹ", "NamSinhMe"),
    (None, "Căn cước mẹ", "CanCuocMe"),
    (None, "Đơn vị công tác mẹ", "DonViCongTacMe"),
    (None, "Điện thoại SLL", "DienThoaiSLL"),
    (None, "Email SLL", "EmailSLL"),
    (None, "Điện thoại bố", "DienThoaiBo"),
    (None, "Điện thoại mẹ", "DienThoaiMe"),
    (None, "Điện thoại HS", "DienThoaiHS"),
    (None, "Khuyết tật", "KhuyetTat"),
    (None, "N.trú, B.trú", "NoiTruBanTru"),
    (None, "Ghi chú", "GhiChu"),
]
HEADER_ROW = 7  # dòng tiêu đề (1-based), dữ liệu bắt đầu từ dòng 9


def _fmt(key: str, v):
    if v is None or v is pd.NaT or (isinstance(v, float) and pd.isna(v)):
        return ""
    if key in ("CanNgheo", "DoanVien", "DoiVien"):
        return "x" if bool(v) else ""
    if isinstance(v, date):
        return v.strftime("%d/%m/%Y")
    if key in ("NgaySinh", "NgayVaoTruong", "NgayCapCanCuoc") and v:
        try:
            return pd.to_datetime(v).strftime("%d/%m/%Y")
        except (ValueError, TypeError):
            return str(v)
    return str(v)


def build_workbook(df: pd.DataFrame, school: str = "TH, THCS VÀ THPT TÂN PHÚ") -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    n = len(COLUMNS)
    bold = Font(name="Times New Roman", bold=True, size=11)
    normal = Font(name="Times New Roman", size=11)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=7)
    ws.merge_cells(start_row=1, start_column=8, end_row=1, end_column=n)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=7)
    ws.merge_cells(start_row=2, start_column=8, end_row=2, end_column=n)
    ws.merge_cells(start_row=3, start_column=8, end_row=3, end_column=n)
    for (r, c, text) in [(1, 1, "SỞ GIÁO DỤC VÀ ĐÀO TẠO TP. HỒ CHÍ MINH"), (2, 1, school),
                         (1, 8, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM"),
                         (2, 8, "Độc lập - Tự do - Hạnh phúc"), (3, 8, "DANH SÁCH HỌC SINH")]:
        cell = ws.cell(r, c, text)
        cell.font, cell.alignment = bold, center

    # Tiêu đề 2 dòng
    h1, h2 = HEADER_ROW, HEADER_ROW + 1
    col = 1
    while col <= n:
        group, label, _ = COLUMNS[col - 1]
        if group is None:
            ws.merge_cells(start_row=h1, start_column=col, end_row=h2, end_column=col)
            ws.cell(h1, col, label)
            col += 1
            continue
        end = col
        while end < n and COLUMNS[end][0] == group:
            end += 1
        ws.merge_cells(start_row=h1, start_column=col, end_row=h1, end_column=end)
        ws.cell(h1, col, group)
        for c in range(col, end + 1):
            ws.cell(h2, c, COLUMNS[c - 1][1])
        col = end + 1
    for r in (h1, h2):
        for c in range(1, n + 1):
            cell = ws.cell(r, c)
            cell.font, cell.alignment, cell.border = bold, center, border

    # Dữ liệu — lưu dạng chữ để giữ số 0 đầu SĐT / căn cước
    for i, (_, row) in enumerate(df.iterrows(), start=1):
        r = h2 + i
        for c, (_, _, key) in enumerate(COLUMNS, start=1):
            v = i if key is None else _fmt(key, row.get(key))
            cell = ws.cell(r, c, v)
            cell.font, cell.border = normal, border
            if key is not None:
                cell.number_format = "@"

    for c, (_, label, _) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(c)].width = 6 if c == 1 else max(12, len(label) + 4)
    ws.column_dimensions["G"].width = 26
    ws.freeze_panes = ws.cell(h2 + 1, 8)

    # Sheet hướng dẫn + danh mục như file mẫu
    dm = danh_muc.load()
    g = wb.create_sheet("Hướng dẫn dẫn nhập")
    g.append(["Tên cột", "Hướng dẫn nhập", "Ghi chú"])
    for row in dm["huong_dan_nhap"]:
        g.append(row)
    for title, key in [("Danh mục tỉnh", "tinh"), ("Danh mục dân tộc", "dan_toc"),
                       ("Danh mục quốc tịch", "quoc_tich"), ("Khuyết tật", "khuyet_tat"),
                       ("N.trú, B.trú", "noi_tru_ban_tru"), ("Danh mục tôn giáo", "ton_giao"),
                       ("Diện chính sách", "dien_chinh_sach")]:
        s = wb.create_sheet(title)
        s.append(["STT", "Tên trường dữ liệu"])
        for k, v in enumerate(dm[key], start=1):
            s.append([k, v])
    s = wb.create_sheet("Danh mục xã")
    s.append(["STT", "Tên", "Tỉnh / Thành Phố"])
    k = 0
    for tinh, xas in dm["xa_theo_tinh"].items():
        for xa in xas:
            k += 1
            s.append([k, xa, tinh])
    return wb


def to_bytes(df: pd.DataFrame) -> bytes:
    buf = io.BytesIO()
    build_workbook(df).save(buf)
    return buf.getvalue()
