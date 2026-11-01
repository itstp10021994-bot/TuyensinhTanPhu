"""Thiết lập list SharePoint và nhập dữ liệu Excel — dùng chung cho trang "Cài đặt" và scripts/."""
from __future__ import annotations

import re
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from typing import Callable

import pandas as pd

from . import config, danh_muc, services
from .export_vemis import COLUMNS, HEADER_ROW
from .schema import (ALL_LISTS, NHAP_DUOC, BOOL, CHOICE, DATE, GIAY_TO, MUC_DANH_GIA, NHAP_HOC, NOI_TRU_DANH_GIA,
                     NOI_TRU_PHONG, NOTE, NUMBER, TAI_KHOAN, TUYEN_SINH, VAI_TRO, Field, ListDef)
from .storage.base import Storage
from .storage.convert import ColumnMap


# ------------------------------------------------------------------ thiết lập list
def check_lists(storage: Storage) -> list[dict]:
    """Tình trạng từng list: có chưa, cột nào đã khớp, cột nào còn thiếu."""
    out = []
    for ld in ALL_LISTS:
        row = {"list": ld, "ten": config.list_name(ld.name), "co_list": False,
               "khop": {}, "thieu": [], "loi": ""}
        try:
            row["co_list"] = storage.list_exists(ld.name)
            if row["co_list"]:
                cm = ColumnMap(ld.name, storage.columns(ld.name))
                row["khop"], row["thieu"] = cm.internal, cm.missing
            else:
                row["thieu"] = list(ld.keys)
                row["site"] = getattr(storage, "site_seen", "")
                if hasattr(storage, "list_titles"):
                    try:
                        row["lists_on_site"] = storage.list_titles()
                    except Exception:
                        pass
        except Exception as e:  # lỗi kết nối
            row["loi"] = str(e)
        out.append(row)
    return out


def setup_lists(storage: Storage, log: Callable[[str], None] = print,
                lists=None) -> None:
    """Tạo list chưa có và thêm các cột còn thiếu (không xóa / đổi cột có sẵn)."""
    for ld in lists or ALL_LISTS:
        name = config.list_name(ld.name)
        if not storage.list_exists(ld.name):
            log(f"Tạo list {name}")
            storage.create_list(ld.name, ld.title)
        cm = ColumnMap(ld.name, storage.columns(ld.name))
        for key in cm.missing:
            f = ld.get(key)
            log(f"{name}: thêm cột “{f.sp_title}” ({key})")
            storage.add_column(ld.name, f, key)
        if hasattr(storage, "show_columns"):  # hiện đủ cột khi mở list trên SharePoint
            cm = ColumnMap(ld.name, storage.columns(ld.name))
            n = storage.show_columns(ld.name, [cm.internal[f.key] for f in ld.fields
                                               if f.key in cm.internal])
            if n:
                log(f"{name}: hiện thêm {n} cột trong dạng xem “Tất cả các mục”")
    log("Xong.")


# ------------------------------------------------------------------ đọc Excel
def _blank(v) -> bool:
    try:
        if pd.isna(v):
            return True
    except (TypeError, ValueError):
        pass
    return str(v).strip() in ("", "nan", "NaN", "NaT", "None")


def _excel_engine(src) -> str | None:
    """File .xls cũ xuất từ VEMIS: xlrd hay báo lỗi "Workbook corruption" -> dùng calamine."""
    if str(getattr(src, "name", src)).lower().endswith(".xls"):
        from importlib.util import find_spec
        return "calamine" if find_spec("python_calamine") else None
    return None


def read_vemis(src) -> pd.DataFrame:
    raw = pd.read_excel(src, header=None, skiprows=HEADER_ROW + 1, dtype=str,
                        engine=_excel_engine(src))
    raw = raw.iloc[:, :len(COLUMNS)]
    raw.columns = [key or "STT" for _, _, key in COLUMNS][:raw.shape[1]]
    raw = raw.dropna(subset=["HoTen"])
    for k in ("CanNgheo", "DoanVien", "DoiVien"):
        raw[k] = raw[k].fillna("").str.strip().str.lower().eq("x")
    for k in ("NgaySinh", "NgayVaoTruong", "NgayCapCanCuoc"):
        raw[k] = pd.to_datetime(raw[k], dayfirst=True, errors="coerce").dt.date
    return raw.drop(columns=["STT"])


def guess_list(sheet_names: list[str]) -> ListDef:
    """Nhận list theo tên sheet (file mẫu đặt tên sheet = tên list)."""
    if TUYEN_SINH.name in sheet_names:
        return TUYEN_SINH
    for ld in NHAP_DUOC:
        if ld.name in sheet_names:
            return ld
    if "DSHS" in sheet_names:  # file danh sách nội trú của Tổ quản nhiệm
        return NOI_TRU_PHONG
    return TUYEN_SINH


def read_frame(src, ld: ListDef, sheet: str | None = None, vemis: bool = False,
               mapping: dict[str, str] | None = None) -> tuple[pd.DataFrame, list[str]]:
    """Đọc file → DataFrame với cột = key trong app. Trả về (dữ liệu, cột bỏ qua)."""
    if vemis:
        df = read_vemis(src)
    elif str(getattr(src, "name", src)).lower().endswith(".csv"):
        df = pd.read_csv(src, dtype=str)
    else:
        sheets = pd.ExcelFile(src).sheet_names
        sheet = sheet or (ld.name if ld.name in sheets else sheets[0])
        df = pd.read_excel(src, sheet_name=sheet, dtype=str)
    # nhận key, nhãn trong app và tên hiển thị trên SharePoint (file "Export to Excel")
    ld = mau(ld) if ld.name in MAU or ld.name == NOI_TRU_DANH_GIA.name else ld
    rename = {f.label: f.key for f in ld.fields}
    rename.update({f.sp_title: f.key for f in ld.fields})
    rename.update(mapping or {})
    df = df.rename(columns=rename)
    keep = [c for c in df.columns if c in ld.keys or c == "Title"]
    skipped = [str(c) for c in df.columns if c not in keep]
    df = df[keep].dropna(how="all")
    for f in ld.fields:  # ngày gõ dạng 01/05/2012 (kiểu Việt Nam) -> 2012-05-01
        if f.type == DATE and f.key in df.columns:
            df[f.key] = df[f.key].map(ngay_chuan)
    return df, skipped


def ngay_chuan(v):
    """Chuẩn hóa ngày trong file: "2012-05-01 00:00:00" / "01/05/2012" / "1-5-2012" -> ISO."""
    if _blank(v):
        return v
    s = str(v).strip()
    if re.match(r"^\d{4}-\d{2}-\d{2}", s):
        return s[:10]
    d = pd.to_datetime(s, dayfirst=True, errors="coerce")
    return s if pd.isna(d) else d.date().isoformat()


# ------------------------------------------------------------------ file mẫu
BO_QUA_MAU = {"TuyenSinhID"}  # trường app tự quản lý, không cho nhập tay


def _mau_danh_gia() -> ListDef:
    from . import noi_tru as NT
    cot = [Field("Ngay", "Ngày", DATE, required=True), Field("Phong", "Phòng", required=True),
           Field("HoTen", "Học sinh", required=True), Field("Lop", "Lớp")]
    for m in NT.MUC:
        cot += [Field(m, NT.TEN_MUC[m], CHOICE, MUC_DANH_GIA), Field(NT.NX[m], NT.TEN_NX[m], NOTE)]
    cot += [Field(k, NT.TEN_LOI[k], NOTE) for k in NT.LOI]
    cot.append(Field("NguoiDanhGia", "Người đánh giá"))
    return ListDef(NOI_TRU_DANH_GIA.name, NOI_TRU_DANH_GIA.title, tuple(cot))


# List có cột app tự quản lý (ID, mật khẩu đã mã hóa, JSON): file mẫu dùng cột dễ điền hơn,
# tên sheet vẫn = tên list để trang Cài đặt tự nhận đúng list
MAU = {
    TAI_KHOAN.name: ListDef(TAI_KHOAN.name, TAI_KHOAN.title, (
        Field("HoTen", "Họ tên", required=True),
        Field("TenDangNhap", "Tên đăng nhập", required=True),
        Field("MatKhau", "Mật khẩu", required=True),
        Field("VaiTro", "Vai trò", CHOICE, VAI_TRO, required=True),
        Field("HoatDong", "Hoạt động", CHOICE, ("Có", "Không")))),
    # cùng tiêu đề với file DSHS của Tổ quản nhiệm (noi_tru.doc_ds_phong đọc được)
    NOI_TRU_PHONG.name: ListDef(NOI_TRU_PHONG.name, "Danh sách nội trú & phân phòng", (
        Field("HoTen", "Họ và tên", required=True),
        Field("Lop", "Lớp"),
        Field("GioiTinh", "Giới tính", CHOICE, ("Nam", "Nữ")),
        Field("Phong", "Phòng NT", required=True),
        Field("GiaoVien", "QN quản lý"))),
}


def mau(ld: ListDef) -> ListDef:
    """Cột của file mẫu nhập cho list `ld`."""
    if ld.name == NOI_TRU_DANH_GIA.name:
        return _mau_danh_gia()
    return MAU.get(ld.name, ld)


def truong_mau(ld: ListDef) -> list:
    return [f for f in mau(ld).fields if f.key not in BO_QUA_MAU and f.group != "_"]


# Hướng dẫn riêng từng list (dòng 4–6 của sheet HuongDan)
_HD = {
    GIAY_TO.name: [
        "4. Mỗi dòng là một giấy tờ cần nộp của một khối. Khối ghi như trong app (vd 6, 10, "
        "10-IEP). Thứ tự: số nhỏ đứng trước.",
        "6. Nhập lại cùng file vẫn an toàn: dòng trùng khối + tên giấy tờ được bỏ qua."],
    TAI_KHOAN.name: [
        "4. Mật khẩu gõ dạng thường, ít nhất 6 ký tự — app chỉ lưu dạng đã mã hóa. XÓA file này "
        "sau khi nhập xong. Trang được vào lấy mặc định theo vai trò (sửa sau ở tab Tài khoản).",
        "6. Tên đăng nhập đã có thì bỏ qua (không đổi mật khẩu tài khoản cũ)."],
    NOI_TRU_PHONG.name: [
        "4. Mỗi học sinh nội trú một dòng. App ghép với Hồ sơ nhập học của năm học chọn lúc nhập "
        "theo họ tên (trùng tên thì xét lớp, giới tính), đánh dấu Nội trú và tạo phòng. QN quản "
        "lý = họ tên giáo viên quản nhiệm (trùng họ tên tài khoản thì tự gắn tài khoản).",
        "6. Nhập lại an toàn: phòng trùng tên trong năm được cập nhật theo file, không tạo trùng."],
    NOI_TRU_DANH_GIA.name: [
        "4. Mỗi dòng là một học sinh trong một ngày; học sinh phải đang ở phòng ghi ở cột Phòng "
        "(năm học chọn lúc nhập). Mức để trống = Tốt. Mỗi ô nhận xét tối đa 300 ký tự. File tải "
        "ở Nội trú → Xem → Tải Excel cũng nhập lại được.",
        "6. Phòng + ngày đã có đánh giá: em có trong file được cập nhật, các em khác giữ nguyên."],
}


def _lua_chon(f) -> list[str]:
    """Danh sách chọn cố định của trường (bỏ danh mục phụ thuộc như phường/xã, trường học)."""
    if f.type != CHOICE:
        return []
    if isinstance(f.options, str) and (f.options.startswith("@xa:")
                                       or f.options.startswith("@truong:")):
        return []
    try:
        return [str(x) for x in danh_muc.options_for(f)]
    except Exception:
        return []


def file_mau(ld: ListDef) -> bytes:
    """File Excel mẫu để nhập: sheet nhập liệu (tiêu đề = tên cột trong app, có ô chọn) và
    sheet hướng dẫn. Tên sheet nhập liệu = tên list -> trang Cài đặt tự nhận đúng list."""
    import io

    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    fields = truong_mau(ld)
    ld = mau(ld)
    wb = Workbook()
    ws = wb.active
    ws.title = ld.name
    dm = wb.create_sheet("DanhMuc")
    hd = wb.create_sheet("HuongDan")
    bat_buoc = PatternFill("solid", fgColor="FFE0B2")
    thuong = PatternFill("solid", fgColor="DCEBFF")
    n_dong = 2000
    cot_dm = 0
    for i, f in enumerate(fields, 1):
        c = ws.cell(1, i, f.label)
        c.font = Font(bold=True)
        c.fill = bat_buoc if f.required else thuong
        c.alignment = Alignment(wrap_text=True, vertical="center")
        col = get_column_letter(i)
        ws.column_dimensions[col].width = max(14, min(36, len(f.label) + 4))
        if f.type == DATE:
            for r in range(2, n_dong + 1):
                ws.cell(r, i).number_format = "DD/MM/YYYY"
            c.comment = Comment("Ngày dạng dd/mm/yyyy, vd 01/05/2012", "App tuyển sinh")
        ds = _lua_chon(f)
        if ds:
            cot_dm += 1
            dcol = get_column_letter(cot_dm)
            dm.cell(1, cot_dm, f.label).font = Font(bold=True)
            for r, v in enumerate(ds, 2):
                dm.cell(r, cot_dm, v)
            dv = DataValidation(type="list", allow_blank=True,
                                formula1=f"=DanhMuc!${dcol}$2:${dcol}${len(ds) + 1}",
                                showErrorMessage=not f.free, errorStyle="warning")
            dv.add(f"{col}2:{col}{n_dong}")
            ws.add_data_validation(dv)
    ws.freeze_panes = "B2"
    ws.row_dimensions[1].height = 42
    dm.sheet_state = "hidden"

    hd.column_dimensions["A"].width = 34
    hd.column_dimensions["B"].width = 11
    hd.column_dimensions["C"].width = 12
    hd.column_dimensions["D"].width = 70
    hd4, hd6 = _HD.get(ld.name, [
        "4. Năm học để trống thì app lấy năm học chọn lúc nhập.",
        "6. Nhập lại cùng file vẫn an toàn: dòng đã có (trùng họ tên + "
        + ("ngày sinh" if ld.name == NHAP_HOC.name else "SĐT + ngày liên hệ")
        + " + năm học) được bỏ qua."])
    loi = [f"FILE MẪU NHẬP: {ld.title.upper()} (list {config.list_name(ld.name)})", "",
           f"1. Nhập mỗi dòng một bản ghi ở sheet \"{ld.name}\", bắt đầu từ dòng 2. Không đổi "
           "tên, không xóa dòng tiêu đề.",
           "2. Cột tô CAM là bắt buộc; cột XANH để trống được. Cột không dùng có thể xóa cả cột.",
           "3. Ô có mũi tên: chọn trong danh sách. Ngày gõ dạng dd/mm/yyyy (vd 01/05/2012).",
           hd4,
           "5. Vào app: Cài đặt & đồng bộ → Nhập / xuất dữ liệu → Nhập dữ liệu từ Excel → tải file "
           "lên → kiểm tra bảng xem trước → bấm Nhập.",
           hd6, ""]
    for r, t in enumerate(loi, 1):
        hd.cell(r, 1, t).font = Font(bold=r == 1, size=13 if r == 1 else 11)
    r0 = len(loi) + 1
    for j, t in enumerate(("Cột", "Bắt buộc", "Kiểu", "Giá trị hợp lệ / ghi chú"), 1):
        c = hd.cell(r0, j, t)
        c.font = Font(bold=True)
        c.fill = thuong
    kieu = {DATE: "Ngày", NUMBER: "Số", CHOICE: "Chọn", NOTE: "Văn bản dài", BOOL: "Có/Không"}
    for k, f in enumerate(fields, r0 + 1):
        ds = _lua_chon(f)
        ghi = ", ".join(ds[:15]) + (f" … ({len(ds)} giá trị, xem ô chọn)" if len(ds) > 15 else "")
        if not ds and isinstance(f.options, str) and f.options.startswith(("@xa:", "@truong:")):
            ghi = "Tên theo địa giới mới (gõ tay), vd Phường Tân Phú"
        if f.free and ds:
            ghi += " — được gõ giá trị khác"
        hd.cell(k, 1, f.label)
        hd.cell(k, 2, "Có" if f.required else "")
        hd.cell(k, 3, kieu.get(f.type, "Chữ"))
        hd.cell(k, 4, ghi).alignment = Alignment(wrap_text=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ nhập
def signature(ld: ListDef, rec: dict) -> tuple:
    """Dấu hiệu nhận biết một dòng đã có (để chạy lại không bị trùng)."""
    def d(v):
        return str(v)[:10] if not _blank(v) else ""

    if ld is NHAP_HOC:
        return (rec.get("NamHoc") or "", danh_muc.fold(rec.get("HoTen")), d(rec.get("NgaySinh")))
    if ld is GIAY_TO:
        return (danh_muc.fold(rec.get("Khoi")), danh_muc.fold(rec.get("TenGiayTo")))
    return (rec.get("NamHoc") or "", danh_muc.fold(rec.get("HoTenHS")),
            services.normalize_phone(rec.get("SDT")), d(rec.get("NgayLienHe")))


def import_rows(storage: Storage, ld: ListDef, df: pd.DataFrame, nam_hoc: str,
                keep_all: bool = True, skip_existing: bool = True, workers: int = 4,
                progress: Callable[[int, int], None] | None = None) -> dict:
    """Ghi các dòng vào list.

    keep_all: chuyển dữ liệu cũ — ghi cả dòng thiếu trường bắt buộc (không kiểm tra).
    skip_existing: bỏ qua dòng đã có trên list (chạy lại an toàn khi bị ngắt giữa chừng).
    """
    rows = []
    for i, row in df.iterrows():
        data = {k: v for k, v in row.to_dict().items() if not _blank(v)}
        data.setdefault("NamHoc", nam_hoc)
        rows.append((i, data))

    linked = 0
    if ld is NHAP_HOC:
        linked = link_tuyen_sinh(storage, [d for _, d in rows])
    if hasattr(storage, "colmap"):  # đọc cấu trúc cột 1 lần trước khi ghi song song
        storage.colmap(ld.name)
    # Đếm số bản đã có trên list theo dấu hiệu; chỉ bỏ qua đúng số đó (dòng trùng nhau
    # ngay trong file dữ liệu cũ vẫn được giữ đủ)
    existing = Counter()
    if skip_existing:
        existing = Counter(signature(ld, it) for it in storage.list_items(ld.name))

    todo, skipped = [], 0
    for i, data in rows:
        clean = services._clean(ld, data)
        if keep_all and data.get("Title"):
            clean["Title"] = data["Title"]
        sig = signature(ld, clean)
        if existing[sig] > 0:
            existing[sig] -= 1
            skipped += 1
            continue
        todo.append((i, data, clean))

    def write(item):
        i, data, clean = item
        if keep_all or ld not in (TUYEN_SINH, NHAP_HOC):
            return storage.create_item(ld.name, clean)
        if ld is TUYEN_SINH:
            return services.save_tuyen_sinh(storage, data)
        return services.save_nhap_hoc(storage, data)

    ok, errors, done = 0, [], 0
    total = len(todo)
    with ThreadPoolExecutor(max_workers=max(1, workers)) as pool:
        futures = {pool.submit(write, item): item for item in todo}
        for fut in as_completed(futures):
            i = futures[fut][0]
            try:
                fut.result()
                ok += 1
            except Exception as e:  # ghi tiếp các dòng khác, báo lỗi cuối cùng
                errors.append({"Dòng": i + 2, "Lỗi": str(e)[:300]})
            done += 1
            if progress:
                progress(done, total)
    return {"ok": ok, "skipped": skipped, "errors": sorted(errors, key=lambda e: e["Dòng"]),
            "total": len(rows), "linked": linked}


def fill_blanks(storage: Storage, ld: ListDef, df: pd.DataFrame, nam_hoc: str,
                workers: int = 4, progress: Callable[[int, int], None] | None = None) -> dict:
    """Bổ sung ô còn trống của các dòng đã có trên list từ file (không ghi đè ô đã có dữ liệu).

    Dùng khi tải list lên SharePoint bằng tay bị mất một cột (vd Khối)."""
    if hasattr(storage, "colmap"):
        storage.colmap(ld.name)
    by_sig: dict[tuple, list[dict]] = {}
    for it in storage.list_items(ld.name):
        by_sig.setdefault(signature(ld, it), []).append(it)
    todo, not_found, cells = [], 0, 0
    for _, row in df.iterrows():
        data = {k: v for k, v in row.to_dict().items() if not _blank(v)}
        data.setdefault("NamHoc", nam_hoc)
        clean = services._clean(ld, data)
        items = by_sig.get(signature(ld, clean))
        if not items:
            not_found += 1
            continue
        it = items.pop(0)
        upd = {k: v for k, v in clean.items()
               if k in ld.keys and not _blank(v) and v != "" and _blank(it.get(k))}
        if upd:
            todo.append((it["id"], upd))
            cells += len(upd)
    done = 0
    with ThreadPoolExecutor(max_workers=max(1, workers)) as pool:
        futures = [pool.submit(storage.update_item, ld.name, i, u) for i, u in todo]
        for fut in as_completed(futures):
            fut.result()
            done += 1
            if progress:
                progress(done, len(todo))
    return {"updated": len(todo), "cells": cells, "not_found": not_found, "total": len(df)}


def nhap_tai_khoan(storage: Storage, df: pd.DataFrame, items: list[dict]) -> dict:
    """Tạo tài khoản từ file mẫu; tên đăng nhập đã có thì bỏ qua (không đổi mật khẩu)."""
    from . import tai_khoan
    items = list(items)
    ok, skipped, errors = 0, 0, []
    for i, row in df.iterrows():
        d = {k: ("" if _blank(v) else str(v).strip()) for k, v in row.to_dict().items()}
        ten = d.get("TenDangNhap", "").lower()
        if ten and any(str(t.get("TenDangNhap") or "").strip().lower() == ten for t in items):
            skipped += 1
            continue
        try:
            vai = d.get("VaiTro") or ""
            if vai not in VAI_TRO:
                raise ValueError(f"Vai trò không hợp lệ: \"{vai}\" (chọn: {', '.join(VAI_TRO)})")
            items.append(tai_khoan.luu(storage, items, d.get("HoTen", ""), ten, vai,
                                       d.get("MatKhau") or None,
                                       hoat_dong=d.get("HoatDong") != "Không"))
            ok += 1
        except Exception as e:
            errors.append({"Dòng": i + 2, "Lỗi": str(e).replace("**", "")[:300]})
    return {"ok": ok, "skipped": skipped, "errors": errors, "total": len(df)}


# ------------------------------------------------------------------ đối chiếu học sinh hiện hữu
KQ_MA, KQ_TEN_NS, KQ_TEN = "Có (trùng mã)", "Có (họ tên + ngày sinh)", "Có (chỉ trùng họ tên)"
KQ_KHONG = "Không có trong file"


def doi_chieu_hien_huu(nh: list[dict], ds: pd.DataFrame, nam_hoc: str
                       ) -> tuple[pd.DataFrame, pd.DataFrame]:
    """So hồ sơ nhập học năm `nam_hoc` với danh sách học sinh hiện hữu (file VEMIS đã đọc).

    Khớp theo Mã VEMIS / Mã học sinh, rồi họ tên + ngày sinh, rồi họ tên (một em duy nhất).
    Trả về (hồ sơ app kèm cột Kết quả, các dòng trong file không có hồ sơ trong app)."""
    def d(v):
        return "" if _blank(v) else str(v)[:10]

    def ma(v):
        return "" if _blank(v) else str(v).strip().lstrip("0")

    from .noi_tru import _chuan_lop, _khoi_lop
    con = ds.reset_index(drop=True).copy()
    for k in ("HoTen", "NgaySinh", "LopHoc", "MaVEMIS", "MaHocSinh"):
        if k not in con.columns:
            con[k] = None
    theo_ma: dict[str, int] = {}
    theo_ten_ns: dict[tuple, list[int]] = {}
    theo_ten: dict[str, list[int]] = {}
    for i, r in con.iterrows():
        for k in ("MaVEMIS", "MaHocSinh"):
            if ma(r.get(k)):
                theo_ma[k + ma(r.get(k))] = i
        ten = danh_muc.fold(r.get("HoTen"))
        theo_ten_ns.setdefault((ten, d(r.get("NgaySinh"))), []).append(i)
        theo_ten.setdefault(ten, []).append(i)
    dung: set[int] = set()
    rows = []
    for rec in [r for r in nh if str(r.get("NamHoc") or "") == nam_hoc]:
        kq, j = KQ_KHONG, None
        for k in ("MaVEMIS", "MaHocSinh"):
            j = theo_ma.get(k + ma(rec.get(k))) if ma(rec.get(k)) else None
            if j is not None:
                kq = KQ_MA
                break
        ten = danh_muc.fold(rec.get("HoTen"))
        if j is None:
            ung = [x for x in theo_ten_ns.get((ten, d(rec.get("NgaySinh"))), []) if x not in dung]
            if ung and d(rec.get("NgaySinh")):
                kq, j = KQ_TEN_NS, ung[0]
        if j is None:  # trùng tên nhiều em: xét lớp, rồi khối
            ung = [x for x in theo_ten.get(ten, []) if x not in dung]
            lop = _chuan_lop(rec.get("LopHoc"))
            khoi = str(rec.get("Khoi") or "").split("-")[0] or _khoi_lop(lop)
            for loc in (lambda x: lop and _chuan_lop(con.at[x, "LopHoc"]) == lop,
                        lambda x: khoi and _khoi_lop(_chuan_lop(con.at[x, "LopHoc"])) == khoi):
                if len(ung) > 1:
                    ung = [x for x in ung if loc(x)] or ung
            # chỉ trùng họ tên mà khác khối: coi là em khác (không khớp nhầm)
            if khoi:
                ung = [x for x in ung
                       if _khoi_lop(_chuan_lop(con.at[x, "LopHoc"])) in ("", khoi)]
            if len(ung) == 1:
                kq, j = KQ_TEN, ung[0]
        if j is not None:
            dung.add(j)
        rows.append({"id": str(rec.get("id")), "Họ tên": rec.get("HoTen") or "",
                     "Ngày sinh": d(rec.get("NgaySinh")), "Khối": rec.get("Khoi") or "",
                     "Lớp (app)": rec.get("LopHoc") or "",
                     "Tình trạng": rec.get("TinhTrangHS") or "",
                     "Lớp (file)": "" if j is None else str(con.at[j, "LopHoc"] or ""),
                     "Kết quả": kq, "_dong": -1 if j is None else int(j)})
    bang = pd.DataFrame(rows, columns=["id", "Họ tên", "Ngày sinh", "Khối", "Lớp (app)",
                                       "Tình trạng", "Lớp (file)", "Kết quả", "_dong"])
    thua = con.loc[[i for i in con.index if i not in dung]]
    thua = pd.DataFrame({"Họ tên": thua["HoTen"], "Ngày sinh": thua["NgaySinh"].map(d),
                         "Lớp": thua["LopHoc"], "_dong": list(thua.index)})
    return bang, thua


NGHI = ("Không học tiếp", "Rút hồ sơ")


def tao_tu_file(storage: Storage, ds: pd.DataFrame, dong: list[int], nam_hoc: str,
                workers: int = 4, progress: Callable[[int, int], None] | None = None) -> int:
    """Tạo hồ sơ nhập học cho các em có trong file VEMIS mà app chưa có (mọi cột theo file)."""
    from .noi_tru import _chuan_lop, _khoi_lop
    con = ds.reset_index(drop=True)
    viec = []
    for j in dong:
        f = con.loc[j]
        data = {k: _gia_tri(NHAP_HOC.get(k), f.get(k)) for k in COT_CAP_NHAT
                if k in f.index and k != "Khoi"}
        data = {k: v for k, v in data.items() if v not in ("", False)}
        khoi = _khoi_lop(_chuan_lop(f.get("LopHoc")))
        if khoi in KHOI_HOP_LE:
            data["Khoi"] = khoi
        data.update(NamHoc=nam_hoc, TinhTrangHS="Học tiếp")
        viec.append(services._clean(NHAP_HOC, data))
    done = 0
    with ThreadPoolExecutor(max_workers=max(1, workers)) as pool:
        for _ in pool.map(lambda x: storage.create_item(NHAP_HOC.name, x), viec):
            done += 1
            if progress:
                progress(done, len(viec))
    return len(viec)


# Cột của biểu mẫu VEMIS được cập nhật theo file (Ghi chú là ghi chú nội bộ của app: giữ nguyên)
COT_CAP_NHAT = {k: (f"{g} - {l}" if g else l) for g, l, k in COLUMNS if k and k != "GhiChu"}
COT_CAP_NHAT["Khoi"] = "Khối"  # không có trong file: suy ra từ lớp (10A1 -> 10)
COT_CAP_NHAT["TinhTrangHS"] = "Tình trạng"  # có trong file = đang học


KHOI_HOP_LE = {str(i) for i in range(1, 13)}


def _gia_tri(f, v):
    """Giá trị đã chuẩn hóa của trường `f` (để so sánh và để ghi)."""
    if _blank(v):
        return False if f.type == BOOL else ""
    if f.type == BOOL:
        return v if isinstance(v, bool) else str(v).strip().lower() in ("true", "x", "1", "có")
    if f.type == DATE:
        return str(v)[:10]
    return " ".join(str(v).split())


def thay_doi_tu_file(bang: pd.DataFrame, ds: pd.DataFrame, nh: list[dict],
                     xoa_o_trong: bool = False) -> list[dict]:
    """Các ô của hồ sơ đã khớp (bảng của doi_chieu_hien_huu) khác với file VEMIS: mọi cột của
    biểu mẫu theo file. Ô trống trong file chỉ xóa ô trong app khi `xoa_o_trong`."""
    con = ds.reset_index(drop=True)
    theo_id = {str(r.get("id")): r for r in nh}
    out = []
    for _, b in bang[bang["_dong"] >= 0].iterrows():
        rec, f = theo_id[b["id"]], con.loc[b["_dong"]]
        upd = {}
        for k in COT_CAP_NHAT:
            if k not in f.index:
                continue
            fd = NHAP_HOC.get(k)
            moi, cu = _gia_tri(fd, f.get(k)), _gia_tri(fd, rec.get(k))
            if moi == cu or (moi in ("", False) and not xoa_o_trong):
                continue
            upd[k] = (cu, moi)
        from .noi_tru import _chuan_lop, _khoi_lop
        khoi = _khoi_lop(_chuan_lop(f.get("LopHoc")))
        if khoi in KHOI_HOP_LE and khoi != str(rec.get("Khoi") or "").strip():
            upd["Khoi"] = (str(rec.get("Khoi") or ""), khoi)
        if str(rec.get("TinhTrangHS") or "") in NGHI:  # có trong danh sách đang học
            upd["TinhTrangHS"] = (str(rec.get("TinhTrangHS")), "Học tiếp")
        if upd:
            out.append({"id": b["id"], "Họ tên": b["Họ tên"], "Kết quả": b["Kết quả"],
                        "thay": upd})
    return out


def _bo_khoi_phong(phongs: list[dict], ids: set[str]) -> list[tuple]:
    """Việc cập nhật phòng nội trú để bỏ các học sinh `ids`."""
    viec = []
    for p in phongs:
        cu = [x.strip() for x in str(p.get("HocSinhIDs") or "").split(";") if x.strip()]
        con = [x for x in cu if x not in ids]
        if len(con) != len(cu):
            viec.append(("phong", (p["id"], ";".join(con))))
    return viec


def cap_nhat_tu_file(storage: Storage, thay: list[dict], phongs: list[dict] | None = None,
                     workers: int = 4, progress: Callable[[int, int], None] | None = None) -> int:
    """Ghi các thay đổi; em đổi sang không còn Nội trú được bỏ khỏi phòng nội trú."""
    viec = [("nh", (t["id"], {k: v[1] for k, v in t["thay"].items()})) for t in thay]
    het_nt = {t["id"] for t in thay if "NoiTruBanTru" in t["thay"]
              and not str(t["thay"]["NoiTruBanTru"][1]).startswith("Nội trú")}
    viec += _bo_khoi_phong(phongs or [], het_nt)

    def lam(v):
        loai, x = v
        if loai == "phong":
            return storage.update_item(NOI_TRU_PHONG.name, x[0], {"HocSinhIDs": x[1]})
        return storage.update_item(NHAP_HOC.name, *x)

    done = 0
    with ThreadPoolExecutor(max_workers=max(1, workers)) as pool:
        for _ in pool.map(lam, viec):
            done += 1
            if progress:
                progress(done, len(viec))
    return len(thay)


def xoa_ho_so(storage: Storage, ids: list[str], ngung: bool, phongs: list[dict],
              workers: int = 4, progress: Callable[[int, int], None] | None = None) -> int:
    """Xóa hồ sơ nhập học (hoặc chỉ đánh dấu "Không học tiếp" khi `ngung`) và bỏ các em khỏi
    phòng nội trú."""
    ids = [str(i) for i in ids]
    viec = [("nh", i) for i in ids] + _bo_khoi_phong(phongs, set(ids))

    def lam(v):
        loai, x = v
        if loai == "phong":
            return storage.update_item(NOI_TRU_PHONG.name, x[0], {"HocSinhIDs": x[1]})
        if ngung:
            return storage.update_item(NHAP_HOC.name, x, {"TinhTrangHS": "Không học tiếp"})
        return storage.delete_item(NHAP_HOC.name, x)

    done = 0
    with ThreadPoolExecutor(max_workers=max(1, workers)) as pool:
        for _ in pool.map(lam, viec):
            done += 1
            if progress:
                progress(done, len(viec))
    return len(ids)


def link_existing(storage: Storage, workers: int = 4,
                  progress: Callable[[int, int], None] | None = None) -> dict:
    """Liên kết các hồ sơ nhập học đã có trên list (vd tải tay từ Excel) với Data tuyển sinh."""
    items = [dict(it) for it in storage.list_items(NHAP_HOC.name)]
    todo = [it for it in items if _blank(it.get("TuyenSinhID"))]
    n = link_tuyen_sinh(storage, todo)
    changed = [it for it in todo if not _blank(it.get("TuyenSinhID"))]
    done = 0
    with ThreadPoolExecutor(max_workers=max(1, workers)) as pool:
        futures = [pool.submit(storage.update_item, NHAP_HOC.name, it["id"],
                               {"TuyenSinhID": it["TuyenSinhID"]}) for it in changed]
        for fut in as_completed(futures):
            fut.result()
            done += 1
            if progress:
                progress(done, len(changed))
    return {"total": len(items), "chua_lien_ket": len(todo), "lien_ket": n}


def link_tuyen_sinh(storage: Storage, rows: list[dict]) -> int:
    """Gán TuyenSinhID cho hồ sơ nhập học chưa có: cùng họ tên và SĐT (hoặc ngày sinh) với
    một học sinh ở Data tuyển sinh, ưu tiên cùng năm học. Trả về số hồ sơ đã liên kết."""
    todo = [d for d in rows if _blank(d.get("TuyenSinhID"))]
    if not todo:
        return 0
    index: dict[tuple, list[dict]] = {}
    for ts in storage.list_items(TUYEN_SINH.name):
        name = danh_muc.fold(ts.get("HoTenHS"))
        if not name:
            continue
        keys = [("sdt", name, services.normalize_phone(ts.get("SDT")))]
        if not _blank(ts.get("NgaySinh")):
            keys.append(("ns", name, str(ts["NgaySinh"])[:10]))
        for k in keys:
            if k[2]:
                index.setdefault(k, []).append(ts)
    n = 0
    for d in todo:
        name = danh_muc.fold(d.get("HoTen"))
        phones = {services.normalize_phone(d.get(k)) for k in
                  ("DienThoaiSLL", "DienThoaiBo", "DienThoaiMe", "DienThoaiNGH")} - {""}
        cands = [ts for p in phones for ts in index.get(("sdt", name, p), [])]
        if not cands and not _blank(d.get("NgaySinh")):
            cands = index.get(("ns", name, str(d["NgaySinh"])[:10]), [])
        if not cands:
            continue
        same_year = [ts for ts in cands if ts.get("NamHoc") == d.get("NamHoc")]
        best = max(same_year or cands, key=lambda ts: str(ts.get("NamHoc") or ""))
        d["TuyenSinhID"] = str(best["id"])
        n += 1
    return n
