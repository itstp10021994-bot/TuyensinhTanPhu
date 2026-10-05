"""Các mẫu báo cáo (tính toán thuần pandas — trang Báo cáo chỉ hiển thị và xuất Excel).

Mỗi báo cáo nhận `Ctx` (dữ liệu đã lọc) và trả về `Result`: bảng số liệu, chỉ số tóm tắt và
gợi ý biểu đồ. Thêm mẫu mới: viết hàm `_ten(ctx) -> Result` và đăng ký trong `REPORTS`.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Callable

import pandas as pd

from . import dia_ban as DB
from . import services
from .schema import GIAY_TO_NHAP_HOC, KHOI, TRANG_THAI

G_TS, G_KT, G_NH = "Tuyển sinh", "Tài chính", "Nhập học"
G_XH, G_XNH = "Xu hướng tuyển sinh", "Xu hướng nhập học"


@dataclass
class Ctx:
    ts: pd.DataFrame            # Data tuyển sinh của năm học (đã lọc khối / ngày)
    nh: pd.DataFrame            # Hồ sơ nhập học của năm học (đã lọc khối)
    ts_all: pd.DataFrame        # Data tuyển sinh mọi năm (so sánh năm học)
    nam_hoc: str
    params: dict = field(default_factory=dict)
    nh_all: pd.DataFrame | None = None  # Hồ sơ nhập học mọi năm (xu hướng nhập học)


@dataclass
class Result:
    table: pd.DataFrame
    kpis: list[tuple[str, object]] = field(default_factory=list)
    chart: dict | None = None     # {"kind": bar|hbar|stack|group|line, "x", "y", "color"...}
    note: str = ""
    sheets: dict[str, pd.DataFrame] | None = None  # xuất nhiều sheet (vd danh sách theo lớp)


@dataclass
class Report:
    id: str
    group: str
    title: str
    desc: str
    fn: Callable[[Ctx], Result]
    params: tuple = ()  # tham số riêng: "ky" (tuần/tháng), "top", "ngay"


# ------------------------------------------------------------------ tiện ích
def _pct(a, b) -> float:
    return round(100 * a / b, 1) if b else 0.0


def _khoi_order(values) -> list[str]:
    vals = [v for v in dict.fromkeys(list(values)) if v != ""]
    return sorted(vals, key=lambda s: (int(s.split("-")[0]) if str(s).split("-")[0].isdigit()
                                       else 99, str(s)))


def _blank(s: pd.Series, label: str = "Chưa rõ") -> pd.Series:
    return s.fillna("").astype(str).str.strip().replace("", label)


def _steps(df: pd.DataFrame, by: str) -> pd.DataFrame:
    """Số HS theo `by` × bước + tổng + tỷ lệ nộp hồ sơ / nhập học."""
    t = pd.crosstab(df[by], df["TrangThai"]).reindex(columns=list(TRANG_THAI), fill_value=0)
    t["Tổng"] = t.sum(axis=1)
    nop = t["Nộp hồ sơ"] + t["Nhập học"]
    t["Tỷ lệ nộp hồ sơ (%)"] = [_pct(a, b) for a, b in zip(nop, t["Tổng"])]
    t["Tỷ lệ nhập học (%)"] = [_pct(a, b) for a, b in zip(t["Nhập học"], t["Tổng"])]
    return t


def _total_row(t: pd.DataFrame, label_col: str, pct_cols=()) -> pd.DataFrame:
    tot = t.select_dtypes("number").sum()
    row = {c: tot.get(c, "") for c in t.columns}
    row[label_col] = "Tổng cộng"
    for c, (num, den) in dict(pct_cols).items():
        row[c] = _pct(tot[num], tot[den])
    return pd.concat([t, pd.DataFrame([row])], ignore_index=True)


def _long(t: pd.DataFrame, id_col: str, cols: list[str], var="Nhóm", val="Số HS"):
    return t[t[id_col] != "Tổng cộng"].melt(id_col, cols, var_name=var, value_name=val)


def _num(s: pd.Series) -> pd.Series:
    return pd.to_numeric(s, errors="coerce").fillna(0)


# ------------------------------------------------------------------ tuyển sinh
def ts_khoi(c: Ctx) -> Result:
    t = _steps(c.ts.assign(Khoi=_blank(c.ts["Khoi"])), "Khoi")
    t = t.reindex(_khoi_order(t.index.tolist()) + (["Chưa rõ"] if "Chưa rõ" in t.index else []))
    t = t.rename_axis("Khối").reset_index()
    t = _total_row(t, "Khối", {"Tỷ lệ nộp hồ sơ (%)": ("Nộp hồ sơ", "Tổng"),
                               "Tỷ lệ nhập học (%)": ("Nhập học", "Tổng")})
    tot = t.iloc[-1]
    return Result(t, [("Liên hệ", int(tot["Tổng"])), ("Nhập học", int(tot["Nhập học"])),
                      ("Tỷ lệ nhập học", f"{tot['Tỷ lệ nhập học (%)']}%"),
                      ("Rút hồ sơ", int(tot["Rút hồ sơ"]))],
                  {"kind": "stack", "data": _long(t, "Khối", list(TRANG_THAI), "Bước"),
                   "x": "Khối", "y": "Số HS", "color": "Bước"})


def ts_nguon(c: Ctx) -> Result:
    t = _steps(c.ts.assign(Nguon=_blank(c.ts["Nguon"])), "Nguon")
    t = t.sort_values("Tổng", ascending=False).rename_axis("Nguồn").reset_index()
    t = t[["Nguồn", "Tổng", "Nộp hồ sơ", "Nhập học", "Rút hồ sơ", "Tỷ lệ nộp hồ sơ (%)",
           "Tỷ lệ nhập học (%)"]].rename(columns={"Tổng": "Liên hệ"})
    best = t[t["Liên hệ"] >= 10].sort_values("Tỷ lệ nhập học (%)", ascending=False)
    t = _total_row(t, "Nguồn", {"Tỷ lệ nộp hồ sơ (%)": ("Nộp hồ sơ", "Liên hệ"),
                                "Tỷ lệ nhập học (%)": ("Nhập học", "Liên hệ")})
    return Result(t, [("Số nguồn", len(t) - 1),
                      ("Nguồn nhiều liên hệ nhất", t.iloc[0]["Nguồn"] if len(t) > 1 else "—"),
                      ("Tỷ lệ nhập học cao nhất (≥10 liên hệ)",
                       f"{best.iloc[0]['Nguồn']} · {best.iloc[0]['Tỷ lệ nhập học (%)']}%"
                       if len(best) else "—")],
                  {"kind": "group", "data": _long(t, "Nguồn", ["Liên hệ", "Nhập học"], "Chỉ số"),
                   "x": "Nguồn", "y": "Số HS", "color": "Chỉ số", "horizontal": True})


def ts_thoigian(c: Ctx) -> Result:
    ky = c.params.get("ky", "Tháng")
    d = pd.to_datetime(c.ts["NgayLienHe"], errors="coerce")
    df = c.ts.assign(_d=d).dropna(subset=["_d"])
    per = df["_d"].dt.to_period("M" if ky == "Tháng" else "W").dt.start_time
    g = df.assign(Ky=per).groupby("Ky").agg(
        **{"Liên hệ mới": ("id", "count"),
           "Đã nộp hồ sơ": ("TrangThai", lambda s: int(s.isin(["Nộp hồ sơ", "Nhập học"]).sum())),
           "Đã nhập học": ("TrangThai", lambda s: int((s == "Nhập học").sum()))}).reset_index()
    g["Lũy kế liên hệ"] = g["Liên hệ mới"].cumsum()
    fmt = "%m/%Y" if ky == "Tháng" else "%d/%m/%Y"
    t = g.assign(**{ky: g["Ky"].dt.strftime(fmt)}).drop(columns="Ky")
    t = t[[ky, "Liên hệ mới", "Đã nộp hồ sơ", "Đã nhập học", "Lũy kế liên hệ"]]
    peak = g.loc[g["Liên hệ mới"].idxmax()] if len(g) else None
    return Result(t, [(f"Số {ky.lower()}", len(g)),
                      (f"{ky} nhiều liên hệ nhất",
                       f"{peak['Ky'].strftime(fmt)} · {peak['Liên hệ mới']}" if peak is not None
                       else "—"),
                      (f"Trung bình / {ky.lower()}",
                       round(g["Liên hệ mới"].mean(), 1) if len(g) else 0)],
                  {"kind": "line", "data": g.melt("Ky", ["Liên hệ mới", "Đã nhập học"],
                                                  var_name="Chỉ số", value_name="Số HS"),
                   "x": "Ky", "y": "Số HS", "color": "Chỉ số", "fmt": fmt},
                  note="Tính theo Ngày liên hệ; cột Đã nộp / Đã nhập học là trạng thái hiện tại "
                       "của các liên hệ trong kỳ đó.")


def ts_tuvan(c: Ctx) -> Result:
    t = _steps(c.ts.assign(Nguoi=_blank(c.ts["NguoiNhanHoSo"], "Chưa phân công")), "Nguoi")
    t = t.sort_values("Tổng", ascending=False).rename_axis("Người nhận hồ sơ").reset_index()
    t = t[["Người nhận hồ sơ", "Tổng", "Tư vấn", "Nộp hồ sơ", "Nhập học", "Rút hồ sơ",
           "Tỷ lệ nhập học (%)"]].rename(columns={"Tổng": "Số HS phụ trách"})
    t = _total_row(t, "Người nhận hồ sơ", {"Tỷ lệ nhập học (%)": ("Nhập học", "Số HS phụ trách")})
    return Result(t, [("Số người phụ trách", len(t) - 1)],
                  {"kind": "stack", "data": _long(t, "Người nhận hồ sơ",
                                                  list(TRANG_THAI), "Bước"),
                   "x": "Người nhận hồ sơ", "y": "Số HS", "color": "Bước", "horizontal": True})


def ts_tinhtrang(c: Ctx) -> Result:
    df = c.ts.assign(TT=_blank(c.ts["TinhTrang"]), Khoi=_blank(c.ts["Khoi"]))
    t = pd.crosstab(df["TT"], df["Khoi"])
    t = t[_khoi_order(t.columns) + (["Chưa rõ"] if "Chưa rõ" in t.columns else [])]
    t["Tổng"] = t.sum(axis=1)
    t = t.sort_values("Tổng", ascending=False).rename_axis("Tình trạng tư vấn").reset_index()
    t.columns.name = None
    t = _total_row(t, "Tình trạng tư vấn")
    return Result(t, [],
                  {"kind": "hbar", "data": t[t["Tình trạng tư vấn"] != "Tổng cộng"],
                   "x": "Tình trạng tư vấn", "y": "Tổng"})


def ts_truongcu(c: Ctx) -> Result:
    top = int(c.params.get("top", 30))
    df = c.ts[c.ts["TruongCu"].str.strip() != ""]
    g = df.groupby(["TruongCu", "TruongCu_Tinh"]).agg(
        **{"Liên hệ": ("id", "count"),
           "Nhập học": ("TrangThai", lambda s: int((s == "Nhập học").sum()))}).reset_index()
    g["Tỷ lệ nhập học (%)"] = [_pct(a, b) for a, b in zip(g["Nhập học"], g["Liên hệ"])]
    g = g.sort_values(["Liên hệ", "Nhập học"], ascending=False).head(top)
    t = g.rename(columns={"TruongCu": "Trường cũ", "TruongCu_Tinh": "Tỉnh/Thành"})
    return Result(t, [("Số trường có học sinh liên hệ", df["TruongCu"].nunique()),
                      ("Liên hệ có ghi trường cũ", f"{len(df)}/{len(c.ts)}")],
                  {"kind": "group", "data": t.head(15).melt("Trường cũ", ["Liên hệ", "Nhập học"],
                                                             var_name="Chỉ số",
                                                             value_name="Số HS"),
                   "x": "Trường cũ", "y": "Số HS", "color": "Chỉ số", "horizontal": True})


# ------------------------------------------------------------------ địa bàn (3 cấp)
def _cap(c: Ctx) -> str:
    cap = c.params.get("cap")
    return cap if cap in DB.CAP else DB.CAP[1]


def _bang_quan(c: Ctx) -> dict:
    """Phường/xã mới → quận/huyện cũ, học từ Data tuyển sinh mọi năm (lưu trong params)."""
    if "_bang_quan" not in c.params:
        c.params["_bang_quan"] = DB.bang_phuong_quan(c.ts_all)
    return c.params["_bang_quan"]


def _theo_tinh(g: pd.DataFrame, c: Ctx) -> pd.Series:
    t = c.params.get("tinh")
    if not t or t == "Tất cả":
        return pd.Series(True, index=g.index)
    return g[DB.CAP[0]] == t


def _ts_dia_ban(c: Ctx, df: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    """(liên hệ đã lọc tỉnh, bảng địa bàn 3 cấp tương ứng)."""
    g = DB.dia_ban_ts(df, _bang_quan(c))
    m = _theo_tinh(g, c)
    return df[m], g[m]


def _ts_bang_cap(df: pd.DataFrame, g: pd.DataFrame, cap: str) -> pd.DataFrame:
    if df.empty:
        return pd.DataFrame(columns=[cap, "Liên hệ", "Nộp hồ sơ (gồm nhập học)", "Nhập học",
                                     "Rút hồ sơ", "Tỷ lệ nhập học (%)"])
    t = _steps(df.assign(_n=DB.nhan(g, cap)), "_n")
    t["Nộp hồ sơ (gồm nhập học)"] = t["Nộp hồ sơ"] + t["Nhập học"]
    t = t.sort_values(["Tổng", "Nhập học"], ascending=False).rename_axis(cap).reset_index()
    t = t.rename(columns={"Tổng": "Liên hệ"})[
        [cap, "Liên hệ", "Nộp hồ sơ (gồm nhập học)", "Nhập học", "Rút hồ sơ",
         "Tỷ lệ nhập học (%)"]]
    return _total_row(_chua_ro_cuoi(t, cap), cap,
                      {"Tỷ lệ nhập học (%)": ("Nhập học", "Liên hệ")})


def _chua_ro_cuoi(t: pd.DataFrame, col: str) -> pd.DataFrame:
    """Đưa các dòng "Chưa rõ…" xuống cuối (giữ thứ tự còn lại), trước dòng Tổng cộng."""
    ro = t[col].astype(str).str.startswith(DB.CHUA_RO)
    tong = t[col] == "Tổng cộng"
    return t.loc[[*t.index[~ro & ~tong], *t.index[ro], *t.index[tong]]]


def _kpi_dia_ban(t: pd.DataFrame, cap: str, ten: str) -> list:
    body = t[t[cap] != "Tổng cộng"]
    ro = body[~body[cap].str.startswith(DB.CHUA_RO)]
    k = [(f"Số {cap.split(' (')[0].lower()}", len(ro))]
    if len(ro):
        k.append((f"Nhiều {ten} nhất", f"{ro.iloc[0][cap].split(' · ')[0]}"))
    chua = body[body[cap].str.startswith(DB.CHUA_RO)]
    if len(chua):
        k.append((f"Chưa rõ {cap.split(' (')[0].lower()}", int(chua.iloc[:, 1].sum())))
    return k


def ts_diaban(c: Ctx) -> Result:
    cap = _cap(c)
    df, g = _ts_dia_ban(c, c.ts)
    t = _ts_bang_cap(df, g, cap)
    body = t[t[cap] != "Tổng cộng"].head(15)
    return Result(t, _kpi_dia_ban(t, cap, "liên hệ"),
                  {"kind": "group", "data": body.melt(cap, ["Liên hệ", "Nhập học"],
                                                       var_name="Chỉ số", value_name="Số HS"),
                   "x": cap, "y": "Số HS", "color": "Chỉ số", "horizontal": True},
                  note="Địa bàn theo TRƯỜNG CŨ của học sinh. Quận/huyện cũ lấy từ địa chỉ trước "
                       "sáp nhập; nếu thiếu thì suy ra từ phường/xã mới. Excel có đủ 3 cấp.",
                  sheets={k.split(" (")[0].replace("/", "_"): _ts_bang_cap(df, g, k)
                          for k in DB.CAP})


def ts_chedo(c: Ctx) -> Result:
    df = c.ts[c.ts["TrangThai"] == "Nhập học"]
    df = df.assign(Khoi=_blank(df["Khoi"]), CheDo=_blank(df["CheDo"]),
                   PhanHe=_blank(df["PhanHe"], "Thường"))
    t = pd.crosstab(df["Khoi"], [df["CheDo"]])
    t2 = pd.crosstab(df["Khoi"], df["PhanHe"])
    t = t.join(t2, rsuffix=" (phân hệ)")
    t["Tổng"] = df.groupby("Khoi").size() if len(df) else 0
    t = t.reindex(_khoi_order(t.index)).rename_axis("Khối").reset_index()
    t.columns.name = None
    t = _total_row(t.fillna(0), "Khối")
    long = df.groupby(["Khoi", "CheDo"]).size().reset_index(name="Số HS")
    return Result(t, [("Học sinh nhập học", len(df)),
                      ("Nội trú", int((df["CheDo"] == "Nội trú").sum())),
                      ("Bán trú", int((df["CheDo"] == "Bán trú").sum()))],
                  {"kind": "stack_cat", "data": long.rename(columns={"Khoi": "Khối",
                                                                    "CheDo": "Chế độ"}),
                   "x": "Khối", "y": "Số HS", "color": "Chế độ"},
                  note="Chỉ tính học sinh ở bước Nhập học.")


def ts_quahan(c: Ctx) -> Result:
    ngay = int(c.params.get("ngay", 14))
    d = pd.to_datetime(c.ts["NgayLienHe"], errors="coerce")
    days = (pd.Timestamp(date.today()) - d).dt.days
    df = c.ts.assign(SoNgay=days)
    df = df[(df["TrangThai"] == "Tư vấn") & (df["SoNgay"] > ngay)].sort_values("SoNgay",
                                                                               ascending=False)
    t = df[["HoTenHS", "Khoi", "SDT", "NgayLienHe", "SoNgay", "TinhTrang", "NguoiNhanHoSo",
            "Nguon", "GhiChu"]].rename(columns={
        "HoTenHS": "Học sinh", "Khoi": "Khối", "SDT": "SĐT", "NgayLienHe": "Ngày liên hệ",
        "SoNgay": "Số ngày chưa chuyển bước", "TinhTrang": "Tình trạng tư vấn",
        "NguoiNhanHoSo": "Người phụ trách", "Nguon": "Nguồn", "GhiChu": "Nội dung đã trao đổi"})
    by = _blank(df["NguoiNhanHoSo"], "Chưa phân công").value_counts().rename_axis(
        "Người phụ trách").reset_index(name="Số HS")
    return Result(t, [("Liên hệ quá hạn", len(t)),
                      ("Quá 30 ngày", int((df["SoNgay"] > 30).sum()))],
                  {"kind": "hbar", "data": by, "x": "Người phụ trách", "y": "Số HS"},
                  note=f"Học sinh còn ở bước Tư vấn hơn {ngay} ngày kể từ ngày liên hệ.")


def ts_rut(c: Ctx) -> Result:
    df = c.ts[c.ts["TrangThai"] == "Rút hồ sơ"]
    t = df[["HoTenHS", "Khoi", "SDT", "NgayLienHe", "GiuCho", "SoTienXacNhan", "NguoiNhanHoSo",
            "GhiChu"]].rename(columns={
        "HoTenHS": "Học sinh", "Khoi": "Khối", "SDT": "SĐT", "NgayLienHe": "Ngày liên hệ",
        "GiuCho": "Giữ chỗ", "SoTienXacNhan": "Tiền giữ chỗ", "NguoiNhanHoSo": "Người phụ trách",
        "GhiChu": "Lý do / nội dung trao đổi"})
    by = _blank(df["Khoi"]).value_counts().rename_axis("Khối").reset_index(name="Số HS")
    return Result(t, [("Rút hồ sơ", len(df)), ("Tỷ lệ trên liên hệ", f"{_pct(len(df), len(c.ts))}%"),
                      ("Đã giữ chỗ trước khi rút",
                       int(df["GiuCho"].isin(["Đã giữ chỗ", "Hủy giữ chỗ", "Đã hoàn phí"]).sum()))],
                  {"kind": "bar", "data": by, "x": "Khối", "y": "Số HS"})


# ------------------------------------------------------------------ tài chính
def kt_giucho(c: Ctx) -> Result:
    df = c.ts.assign(Khoi=_blank(c.ts["Khoi"]), Giu=_blank(c.ts["GiuCho"], "Chưa giữ chỗ"),
                     Tien=_num(c.ts["SoTienXacNhan"]))
    t = pd.crosstab(df["Khoi"], df["Giu"])
    t["Tiền đã xác nhận (đ)"] = df.groupby("Khoi")["Tien"].sum()
    t = t.reindex(_khoi_order(t.index)).fillna(0).rename_axis("Khối").reset_index()
    t.columns.name = None
    t = _total_row(t, "Khối")
    tien = df["Tien"].sum()
    return Result(t, [("Đã giữ chỗ", int((df["Giu"] == "Đã giữ chỗ").sum())),
                      ("Tổng tiền giữ chỗ", tien),
                      ("Hủy / hoàn phí", int(df["Giu"].isin(["Hủy giữ chỗ", "Đã hoàn phí"]).sum()))],
                  {"kind": "stack_cat", "data": df[df["Giu"] != "Chưa giữ chỗ"].groupby(
                      ["Khoi", "Giu"]).size().reset_index(name="Số HS").rename(
                      columns={"Khoi": "Khối", "Giu": "Giữ chỗ"}),
                   "x": "Khối", "y": "Số HS", "color": "Giữ chỗ"},
                  note="Theo Data tuyển sinh (kế toán xác nhận giữ chỗ).")


def kt_hocphi(c: Ctx) -> Result:
    df = c.nh.assign(Khoi=_blank(c.nh["Khoi"]), Phai=_num(c.nh["SoTienThanhToan"]),
                     Thu=_num(c.nh["TongDaThu"]), Con=_num(c.nh["SoTienConLai"]),
                     Giu=_num(c.nh["SoTienXacNhan"]))
    g = df.groupby("Khoi").agg(**{
        "Số HS": ("id", "count"),
        "HS đã ghi học phí": ("Phai", lambda s: int((s > 0).sum())),
        "Tiền giữ chỗ (đ)": ("Giu", "sum"), "PHHS thanh toán (đ)": ("Phai", "sum"),
        "Đã thu (đ)": ("Thu", "sum"), "Còn lại (đ)": ("Con", "sum"),
        "HS còn nợ": ("Con", lambda s: int((s > 0).sum()))})
    t = g.reindex(_khoi_order(g.index)).rename_axis("Khối").reset_index()
    t = _total_row(t, "Khối")
    tot = t.iloc[-1]
    return Result(t, [("PHHS thanh toán", tot["PHHS thanh toán (đ)"]),
                      ("Còn lại", tot["Còn lại (đ)"]), ("HS còn nợ", int(tot["HS còn nợ"]))],
                  {"kind": "bar", "data": t[t["Khối"] != "Tổng cộng"], "x": "Khối",
                   "y": "Còn lại (đ)", "money": True},
                  note="Theo Hồ sơ nhập học (nút Thanh toán học phí).")


def kt_congno(c: Ctx) -> Result:
    df = c.nh[_num(c.nh["SoTienConLai"]) > 0].sort_values(["Khoi", "LopHoc", "HoTen"])
    t = df[["HoTen", "Khoi", "LopHoc", "DienThoaiSLL", "SoTienThanhToan", "TongDaThu",
            "SoTienConLai", "NgayDongPhi", "KeToanXacNhan"]].rename(columns={
        "HoTen": "Học sinh", "Khoi": "Khối", "LopHoc": "Lớp", "DienThoaiSLL": "SĐT",
        "SoTienThanhToan": "PHHS thanh toán (đ)", "TongDaThu": "Đã thu (đ)",
        "SoTienConLai": "Còn lại (đ)", "NgayDongPhi": "Ngày đóng",
        "KeToanXacNhan": "Kế toán"})
    return Result(t, [("HS còn nợ", len(t)), ("Tổng còn lại", _num(df["SoTienConLai"]).sum())])


def kt_hoanphi(c: Ctx) -> Result:
    df = c.ts[c.ts["GiuCho"].isin(["Hủy giữ chỗ", "Đã hoàn phí"])].sort_values(["GiuCho",
                                                                               "HoTenHS"])
    t = df[["HoTenHS", "Khoi", "SDT", "GiuCho", "SoTienXacNhan", "TenChuTaiKhoan", "NganHang",
            "SoTaiKhoan", "NguoiXacNhan"]].rename(columns={
        "HoTenHS": "Học sinh", "Khoi": "Khối", "SDT": "SĐT", "GiuCho": "Tình trạng",
        "SoTienXacNhan": "Số tiền (đ)", "TenChuTaiKhoan": "Chủ tài khoản",
        "NganHang": "Ngân hàng", "SoTaiKhoan": "Số tài khoản", "NguoiXacNhan": "Người xác nhận"})
    cho = df[df["GiuCho"] == "Hủy giữ chỗ"]
    return Result(t, [("Chờ hoàn phí", len(cho)),
                      ("Tiền chờ hoàn", _num(cho["SoTienXacNhan"]).sum()),
                      ("Đã hoàn phí", int((df["GiuCho"] == "Đã hoàn phí").sum()))])


# ------------------------------------------------------------------ nhập học
def _nh_dang_hoc(c: Ctx) -> pd.DataFrame:
    return c.nh[~c.nh["TinhTrangHS"].isin(["Rút hồ sơ", "Không học tiếp"])]


def nh_siso(c: Ctx) -> Result:
    df = _nh_dang_hoc(c)
    df = df.assign(Lop=_blank(df["LopHoc"], "Chưa xếp lớp"), Khoi=_blank(df["Khoi"]))
    g = df.groupby(["Khoi", "Lop"]).agg(**{
        "Sĩ số": ("id", "count"),
        "Nam": ("GioiTinh", lambda s: int((s == "Nam").sum())),
        "Nữ": ("GioiTinh", lambda s: int((s == "Nữ").sum())),
        "Nội trú": ("NoiTruBanTru", lambda s: int(s.str.startswith("Nội trú").sum())),
        "Bán trú": ("NoiTruBanTru", lambda s: int(s.str.startswith("Bán trú").sum())),
        "IEP": ("PhanHe", lambda s: int((s == "IEP").sum())),
        "ESL": ("PhanHe", lambda s: int((s == "ESL").sum()))}).reset_index()
    order = {k: i for i, k in enumerate(_khoi_order(g["Khoi"]))}
    g = g.sort_values(["Khoi", "Lop"], key=lambda s: s.map(order) if s.name == "Khoi" else s)
    t = _total_row(g.rename(columns={"Khoi": "Khối", "Lop": "Lớp"}), "Khối")
    t.loc[t.index[-1], "Lớp"] = ""
    return Result(t, [("Học sinh đang theo học", len(df)),
                      ("Số lớp", df.loc[df["Lop"] != "Chưa xếp lớp", "Lop"].nunique()),
                      ("Chưa xếp lớp", int((df["Lop"] == "Chưa xếp lớp").sum()))],
                  {"kind": "bar", "data": g[g["Lop"] != "Chưa xếp lớp"].rename(
                      columns={"Lop": "Lớp"}), "x": "Lớp", "y": "Sĩ số"},
                  note="Không tính học sinh đã rút hồ sơ / không học tiếp.")


def nh_tinhtrang(c: Ctx) -> Result:
    df = c.nh.assign(Khoi=_blank(c.nh["Khoi"]), TT=_blank(c.nh["TinhTrangHS"]))
    t = pd.crosstab(df["Khoi"], df["TT"])
    t["Tổng"] = t.sum(axis=1)
    t = t.reindex(_khoi_order(t.index)).rename_axis("Khối").reset_index()
    t.columns.name = None
    t = _total_row(t, "Khối")
    return Result(t, [(k, int((df["TT"] == k).sum())) for k in
                      ("Đang nhập hồ sơ", "Đang đóng phí", "Đã đóng phí", "Rút hồ sơ")],
                  {"kind": "stack_cat", "data": df.groupby(["Khoi", "TT"]).size().reset_index(
                      name="Số HS").rename(columns={"Khoi": "Khối", "TT": "Tình trạng"}),
                   "x": "Khối", "y": "Số HS", "color": "Tình trạng"})


def nh_hoanthien(c: Ctx) -> Result:
    df = _nh_dang_hoc(c)
    miss = df.apply(services.missing_fields, axis=1)
    df = df.assign(Lop=_blank(df["LopHoc"], "Chưa xếp lớp"), _m=miss.map(len))
    g = df.groupby("Lop").agg(**{"Số HS": ("id", "count"),
                                 "Đủ thông tin": ("_m", lambda s: int((s == 0).sum())),
                                 "Còn thiếu": ("_m", lambda s: int((s > 0).sum()))})
    g["Tỷ lệ đủ (%)"] = [_pct(a, b) for a, b in zip(g["Đủ thông tin"], g["Số HS"])]
    t = g.rename_axis("Lớp").reset_index()
    t = _total_row(t, "Lớp", {"Tỷ lệ đủ (%)": ("Đủ thông tin", "Số HS")})
    thieu = pd.Series([services.NHAP_HOC_NHAN_NGAN[k] for m in miss for k in m])
    top = thieu.value_counts().rename_axis("Thông tin còn thiếu").reset_index(name="Số HS")
    return Result(t, [("Đủ thông tin VEMIS", int(t.iloc[-1]["Đủ thông tin"])),
                      ("Còn thiếu", int(t.iloc[-1]["Còn thiếu"]))],
                  {"kind": "hbar", "data": top, "x": "Thông tin còn thiếu", "y": "Số HS"},
                  sheets={"Theo_lop": t, "Thong_tin_thieu": top})


def nh_giayto(c: Ctx) -> Result:
    """Giấy tờ theo danh sách cần nộp của từng khối (giay_to_theo_khoi.json)."""
    df = _nh_dang_hoc(c)
    tien_do = [services.tien_do_giay_to(r) for r in df.to_dict("records")]
    rows = []
    for g in GIAY_TO_NHAP_HOC:
        can = sum(1 for cn, _, _ in tien_do if g in cn)
        goc = sum(1 for cn, d, _ in tien_do if g in cn and d.get(g) == "Bản gốc")
        sao = sum(1 for cn, d, _ in tien_do if g in cn and d.get(g) == "Bản sao")
        if can:
            rows.append({"Giấy tờ": g, "HS cần nộp": can, "Bản gốc": goc, "Bản sao": sao,
                         "Đã nộp": goc + sao, "Chưa nộp": can - goc - sao,
                         "Tỷ lệ đã nộp (%)": _pct(goc + sao, can)})
    t = pd.DataFrame(rows, columns=["Giấy tờ", "HS cần nộp", "Bản gốc", "Bản sao", "Đã nộp",
                                    "Chưa nộp", "Tỷ lệ đã nộp (%)"])
    khoi = df.assign(KhoiPH=[f"{k}-{p}" if k and p else k or p or "Chưa rõ"
                             for k, p in zip(df["Khoi"], df["PhanHe"])],
                     Can=[len(x[0]) for x in tien_do], Thieu=[len(x[2]) for x in tien_do])
    theo_khoi = khoi.groupby("KhoiPH").agg(**{
        "Số HS": ("id", "count"), "Số giấy tờ cần nộp": ("Can", "max"),
        "HS nộp đủ": ("Thieu", lambda s: int((s == 0).sum())),
        "HS còn thiếu": ("Thieu", lambda s: int((s > 0).sum()))})
    theo_khoi = theo_khoi.reindex(_khoi_order(theo_khoi.index)).rename_axis("Khối").reset_index()
    ds = khoi[khoi["Thieu"] > 0].assign(
        DaNop=[f"{len(cn) - len(th)}/{len(cn)}" for cn, _, th in tien_do if th],
        ConThieu=["; ".join(th) for _, _, th in tien_do if th])
    ds = ds.sort_values(["KhoiPH", "LopHoc", "HoTen"])[
        ["HoTen", "KhoiPH", "LopHoc", "DienThoaiSLL", "DaNop", "ConThieu"]].rename(columns={
            "HoTen": "Học sinh", "KhoiPH": "Khối", "LopHoc": "Lớp", "DienThoaiSLL": "SĐT",
            "DaNop": "Đã nộp", "ConThieu": "Còn thiếu"})
    du = int((khoi["Thieu"] == 0).sum()) if len(khoi) else 0
    return Result(t, [("Nộp đủ giấy tờ", du), ("Còn thiếu", len(ds)),
                      ("Chưa nộp giấy tờ nào",
                       sum(1 for _, d, _ in tien_do if not d))],
                  {"kind": "hbar", "data": t, "x": "Giấy tờ", "y": "Đã nộp"},
                  note="Theo danh sách giấy tờ cần nộp của từng khối. Xuất Excel có thêm sheet "
                       "tiến độ theo khối và danh sách học sinh còn thiếu.",
                  sheets={"Thong_ke_giay_to": t, "Theo_khoi": theo_khoi,
                          "HS_thieu_giay_to": ds})


def nh_xe(c: Ctx) -> Result:
    df = _nh_dang_hoc(c)
    df = df[(df["DangKyXe"] != "") & df["NoiTruBanTru"].str.startswith("Nội trú")]
    g = df.groupby("DangKyXe").agg(**{"Số HS": ("id", "count"),
                                      "Số điểm đón": ("DiemDonTra", "nunique")})
    g = g.sort_values("Số HS", ascending=False).rename_axis("Tuyến xe").reset_index()
    ds = df.sort_values(["DangKyXe", "DiemDonTra", "HoTen"])[
        ["DangKyXe", "DiemDonTra", "HoTen", "Khoi", "LopHoc", "DienThoaiSLL", "TenCha",
         "TenMe"]].rename(columns={"DangKyXe": "Tuyến xe", "DiemDonTra": "Điểm đón trả",
                                   "HoTen": "Học sinh", "Khoi": "Khối", "LopHoc": "Lớp",
                                   "DienThoaiSLL": "SĐT", "TenCha": "Cha", "TenMe": "Mẹ"})
    sheets = {"Tong_hop": g, **{f"Xe_{x}"[:31]: d for x, d in ds.groupby("Tuyến xe")}}
    return Result(ds, [("HS đăng ký xe", len(df)), ("Số tuyến", len(g))],
                  {"kind": "hbar", "data": g, "x": "Tuyến xe", "y": "Số HS"},
                  note="Chỉ học sinh nội trú. Xuất Excel: mỗi tuyến một sheet (danh sách đón trả).",
                  sheets=sheets)


def nh_tohop(c: Ctx) -> Result:
    df = _nh_dang_hoc(c)
    df = df[df["Khoi"].isin(["10", "11", "12"])]
    t = pd.crosstab(_blank(df["LuaChon1"], "Chưa chọn"), df["Khoi"])
    t["Tổng"] = t.sum(axis=1)
    t = t.sort_values("Tổng", ascending=False).rename_axis("Tổ hợp (lựa chọn 1)").reset_index()
    t.columns.name = None
    t = _total_row(t, "Tổ hợp (lựa chọn 1)")
    return Result(t, [("HS khối 10–12", len(df)),
                      ("Chưa chọn tổ hợp", int((df["LuaChon1"] == "").sum()))],
                  {"kind": "hbar", "data": t[t.iloc[:, 0] != "Tổng cộng"],
                   "x": "Tổ hợp (lựa chọn 1)", "y": "Tổng"})


def _nh_bang_cap(df: pd.DataFrame, g: pd.DataFrame, cap: str) -> pd.DataFrame:
    if df.empty:
        return pd.DataFrame(columns=[cap, "Số HS"])
    d = df.assign(_n=DB.nhan(g, cap), Khoi=_blank(df["Khoi"]))
    t = d.groupby("_n").agg(**{
        "Số HS": ("id", "count"),
        "HS mới": ("TinhTrangHS", lambda s: int((~s.isin(HS_CU)).sum())),
        "HS cũ học tiếp": ("TinhTrangHS", lambda s: int((s == "Học tiếp").sum())),
        "Nội trú": ("NoiTruBanTru", lambda s: int(s.fillna("").str.startswith("Nội trú").sum())),
        "Bán trú": ("NoiTruBanTru", lambda s: int(s.fillna("").str.startswith("Bán trú").sum())),
    })
    k = pd.crosstab(d["_n"], d["Khoi"])
    k = k[_khoi_order(k.columns) + (["Chưa rõ"] if "Chưa rõ" in k.columns else [])]
    t = t.join(k.add_prefix("Khối "))
    t = t.sort_values("Số HS", ascending=False).rename_axis(cap).reset_index()
    return _total_row(_chua_ro_cuoi(t, cap), cap)


def nh_diaban(c: Ctx) -> Result:
    cap = _cap(c)
    df = _nh_dang_hoc(c)
    g = DB.dia_ban_nh(df, _bang_quan(c))
    m = _theo_tinh(g, c)
    df, g = df[m], g[m]
    t = _nh_bang_cap(df, g, cap)
    body = t[t[cap] != "Tổng cộng"].head(15)
    return Result(t, _kpi_dia_ban(t, cap, "học sinh"),
                  {"kind": "hbar", "data": body, "x": cap, "y": "Số HS"},
                  note="Địa bàn theo CHỖ Ở HIỆN NAY của học sinh đang học (thiếu thì lấy hộ khẩu). "
                       "Quận/huyện cũ suy ra từ phường/xã mới. Excel có đủ 3 cấp.",
                  sheets={k.split(" (")[0].replace("/", "_"): _nh_bang_cap(df, g, k)
                          for k in DB.CAP})


def nh_hocluc(c: Ctx) -> Result:
    df = _nh_dang_hoc(c)
    df = df.assign(Khoi=_blank(df["Khoi"]), HL=_blank(df["HocLuc1"]))
    t = pd.crosstab(df["HL"], df["Khoi"])
    t = t[_khoi_order(t.columns) + (["Chưa rõ"] if "Chưa rõ" in t.columns else [])]
    t["Tổng"] = t.sum(axis=1)
    t = t.sort_values("Tổng", ascending=False).rename_axis("Học lực (kết quả 1)").reset_index()
    t.columns.name = None
    diem = df.groupby("Khoi")[["Toan1", "Van1", "Anh1"]].mean().round(2)
    diem = diem.reindex(_khoi_order(diem.index)).rename_axis("Khối").reset_index().rename(
        columns={"Toan1": "Toán TB", "Van1": "Văn TB", "Anh1": "Anh TB"})
    return Result(t, [("Có điểm Toán", int(df["Toan1"].notna().sum()))],
                  {"kind": "hbar", "data": t, "x": "Học lực (kết quả 1)", "y": "Tổng"},
                  sheets={"Hoc_luc": t, "Diem_TB_theo_khoi": diem})


def nh_danhsach(c: Ctx) -> Result:
    df = _nh_dang_hoc(c)
    df = df.assign(Lop=_blank(df["LopHoc"], "Chưa xếp lớp")).sort_values(["Lop", "HoTen"])
    cols = {"HoTen": "Họ và tên", "NgaySinh": "Ngày sinh", "GioiTinh": "Giới tính",
            "KhoiPH": "Khối", "NoiTruBanTru": "Chế độ", "DienThoaiSLL": "SĐT liên lạc",
            "TenCha": "Cha", "TenMe": "Mẹ", "TruongCu": "Trường cũ"}
    df = df.assign(KhoiPH=[f"{k}-{p}" if k and p else k or p
                           for k, p in zip(df["Khoi"], df["PhanHe"])])
    sheets = {}
    for lop, d in df.groupby("Lop", sort=True):
        x = d[list(cols)].rename(columns=cols)
        x.insert(0, "STT", range(1, len(x) + 1))
        sheets[str(lop)[:31]] = x
    t = df[["Lop", *cols]].rename(columns={"Lop": "Lớp", **cols})
    return Result(t, [("Số lớp", len(sheets)), ("Học sinh", len(df))],
                  note="Xuất Excel: mỗi lớp một sheet, có STT — dùng in danh sách lớp.",
                  sheets=sheets)


# ------------------------------------------------------------------ xu hướng các năm
def _xh(c: Ctx) -> tuple[pd.DataFrame, list[str]]:
    """Data tuyển sinh của các năm được chọn (params nam, mặc định mọi năm) + lọc khối."""
    df = c.ts_all.assign(NamHoc=_blank(c.ts_all["NamHoc"]))
    nams = [n for n in (c.params.get("nam") or sorted(df["NamHoc"].unique())) if n != "Chưa rõ"]
    df = df[df["NamHoc"].isin(nams)]
    if c.params.get("khoi"):
        df = df[df["Khoi"].isin(c.params["khoi"])]
    return df, sorted(nams)


def _bien_dong(t: pd.DataFrame, nams: list[str]) -> pd.DataFrame:
    """Thêm cột tăng/giảm (chuỗi "+12.5%") của năm cuối so với năm liền trước."""
    if len(nams) >= 2:
        t[f"Tăng/giảm {nams[-1]} so với {nams[-2]}"] = [
            _pct_doi(x, y) for x, y in zip(t[nams[-2]], t[nams[-1]])]
    return t


def _bien_dong_so(t: pd.DataFrame, nams: list[str], toi_thieu: int = 5) -> pd.Series:
    """Tăng/giảm (%) dạng số, chỉ tính các dòng năm trước có ít nhất `toi_thieu` HS."""
    if len(nams) < 2:
        return pd.Series(dtype=float)
    a, b = t[nams[-2]], t[nams[-1]]
    ok = a >= toi_thieu
    return (100 * (b[ok] - a[ok]) / a[ok]).round(1)


def _pivot_nam(df: pd.DataFrame, by: str, nams: list[str], label: str,
               chi_nhap_hoc: bool = False) -> pd.DataFrame:
    d = df[df["TrangThai"] == "Nhập học"] if chi_nhap_hoc else df
    t = pd.crosstab(_blank(d[by]), d["NamHoc"]).reindex(columns=nams, fill_value=0)
    t["Tổng"] = t.sum(axis=1)
    t = t.rename_axis(label).reset_index()
    t.columns.name = None
    return t


def _pct_doi(a, b) -> str:
    if not a:
        return "—"
    v = round(100 * (b - a) / a, 1)
    return f"{'+' if v > 0 else ''}{v}%"


def xh_tongquan(c: Ctx) -> Result:
    df, nams = _xh(c)
    t = _steps(df, "NamHoc").reindex(nams, fill_value=0).rename_axis("Năm học").reset_index()
    t["Nộp hồ sơ (gồm nhập học)"] = t["Nộp hồ sơ"] + t["Nhập học"]
    t = t[["Năm học", "Tổng", "Tư vấn", "Nộp hồ sơ (gồm nhập học)", "Nhập học", "Rút hồ sơ",
           "Tỷ lệ nộp hồ sơ (%)", "Tỷ lệ nhập học (%)"]].rename(columns={"Tổng": "Liên hệ"})
    for col in ("Liên hệ", "Nhập học"):
        t[f"{col} tăng/giảm"] = ["—"] + [_pct_doi(a, b) for a, b in
                                        zip(t[col].iloc[:-1], t[col].iloc[1:])]
    kpis = [("Số năm học", len(t))]
    if len(t) >= 2:
        a, b = t.iloc[-2], t.iloc[-1]
        kpis += [(f"Liên hệ {b['Năm học']}", f"{int(b['Liên hệ'])} ({_pct_doi(a['Liên hệ'], b['Liên hệ'])})"),
                 (f"Nhập học {b['Năm học']}", f"{int(b['Nhập học'])} ({_pct_doi(a['Nhập học'], b['Nhập học'])})"),
                 (f"Tỷ lệ nhập học {b['Năm học']}",
                  f"{b['Tỷ lệ nhập học (%)']}% ({b['Tỷ lệ nhập học (%)'] - a['Tỷ lệ nhập học (%)']:+.1f} điểm)")]
    return Result(t, kpis,
                  {"kind": "group", "data": t.melt("Năm học", ["Liên hệ", "Nộp hồ sơ (gồm nhập học)",
                                                               "Nhập học"],
                                                   var_name="Chỉ số", value_name="Số HS"),
                   "x": "Năm học", "y": "Số HS", "color": "Chỉ số"},
                  note="Không áp dụng bộ lọc ngày; chọn năm học và khối ở bộ lọc phía trên.")


def xh_khoi(c: Ctx) -> Result:
    df, nams = _xh(c)
    lh = _pivot_nam(df, "Khoi", nams, "Khối")
    nh = _pivot_nam(df, "Khoi", nams, "Khối", chi_nhap_hoc=True)
    order = _khoi_order(lh["Khối"]) + (["Chưa rõ"] if "Chưa rõ" in set(lh["Khối"]) else [])
    nh = nh.set_index("Khối").reindex(order, fill_value=0).reset_index()
    lh = lh.set_index("Khối").reindex(order, fill_value=0).reset_index()
    kpis = []
    bd = _bien_dong_so(nh.set_index("Khối"), nams)
    if len(bd):
        kpis = [("Khối tăng mạnh nhất", f"{bd.idxmax()} ({bd.max():+}%)"),
                ("Khối giảm mạnh nhất", f"{bd.idxmin()} ({bd.min():+}%)")]
    t = _bien_dong(_total_row(nh, "Khối"), nams)
    return Result(t, kpis,
                  {"kind": "group", "data": nh.melt("Khối", nams, var_name="Năm học",
                                                    value_name="Số HS"),
                   "x": "Khối", "y": "Số HS", "color": "Năm học"},
                  note="Số học sinh NHẬP HỌC theo khối qua các năm (sheet 2: số liên hệ). "
                       "Khối tăng/giảm mạnh nhất chỉ xét khối có từ 5 HS ở năm trước.",
                  sheets={"Nhap_hoc_theo_khoi": t,
                          "Lien_he_theo_khoi": _bien_dong(_total_row(lh, "Khối"), nams)})


def xh_nguon(c: Ctx) -> Result:
    df, nams = _xh(c)
    lh = _pivot_nam(df, "Nguon", nams, "Nguồn")
    nh = _pivot_nam(df, "Nguon", nams, "Nguồn", chi_nhap_hoc=True).set_index("Nguồn")
    t = lh.set_index("Nguồn")[nams].add_prefix("Liên hệ ").join(
        nh.reindex(lh["Nguồn"], fill_value=0)[nams].add_prefix("Nhập học "))
    for n in nams:
        t[f"Tỷ lệ NH {n} (%)"] = [_pct(a, b) for a, b in zip(t[f"Nhập học {n}"], t[f"Liên hệ {n}"])]
    t = t.assign(_s=lh.set_index("Nguồn")["Tổng"]).sort_values("_s", ascending=False) \
        .drop(columns="_s").rename_axis("Nguồn").reset_index()
    data = lh.melt("Nguồn", nams, var_name="Năm học", value_name="Số HS")
    return Result(t, [("Số nguồn", len(t))],
                  {"kind": "group", "data": data, "x": "Nguồn", "y": "Số HS", "color": "Năm học",
                   "horizontal": True},
                  note="Biểu đồ: số liên hệ theo nguồn từng năm; bảng có thêm nhập học và tỷ lệ "
                       "nhập học (NH) của từng năm.")


def xh_truongcu(c: Ctx) -> Result:
    df, nams = _xh(c)
    top = int(c.params.get("top", 30))
    df = df[df["TruongCu"].fillna("").str.strip() != ""]
    t = _pivot_nam(df, "TruongCu", nams, "Trường cũ").sort_values("Tổng", ascending=False)
    t = _bien_dong(t, nams).head(top)
    moi = t[(t[nams[0]] == 0)] if len(nams) >= 2 else t.iloc[0:0]
    return Result(t, [("Số trường (mọi năm)", df["TruongCu"].nunique()),
                      ("Trường mới xuất hiện trong top", len(moi))],
                  {"kind": "group", "data": t.head(12).melt("Trường cũ", nams, var_name="Năm học",
                                                            value_name="Số HS"),
                   "x": "Trường cũ", "y": "Số HS", "color": "Năm học", "horizontal": True},
                  note="Số liên hệ theo trường cũ từng năm, xếp theo tổng các năm.")


def xh_diaban(c: Ctx) -> Result:
    df, nams = _xh(c)
    cap = _cap(c)
    df, g = _ts_dia_ban(c, df)
    df = df.assign(_n=DB.nhan(g, cap))
    lh = _pivot_nam(df, "_n", nams, cap).set_index(cap)
    nh = _pivot_nam(df, "_n", nams, cap, chi_nhap_hoc=True).set_index(cap) \
        .reindex(lh.index, fill_value=0)
    t = pd.DataFrame(index=lh.index)
    for n in nams:
        t[f"Liên hệ {n}"] = lh[n]
        t[f"Nhập học {n}"] = nh[n]
        t[f"Tỷ lệ NH {n} (%)"] = [_pct(a, b) for a, b in zip(nh[n], lh[n])]
    if len(nams) >= 2:
        t[f"Liên hệ: tăng/giảm {nams[-1]}"] = [_pct_doi(a, b) for a, b in
                                              zip(lh[nams[-2]], lh[nams[-1]])]
        t[f"Nhập học: tăng/giảm {nams[-1]}"] = [_pct_doi(a, b) for a, b in
                                               zip(nh[nams[-2]], nh[nams[-1]])]
    t = _chua_ro_cuoi(t.assign(_s=lh["Tổng"]).sort_values("_s", ascending=False)
                      .drop(columns="_s").rename_axis(cap).reset_index(), cap)
    data = lh.reset_index().assign(_s=lh["Tổng"].values).sort_values("_s", ascending=False) \
        .head(12).melt(cap, nams, var_name="Năm học", value_name="Số HS")
    kpis = []
    bd = _bien_dong_so(lh[~lh.index.str.startswith(DB.CHUA_RO)], nams)
    if len(bd):
        kpis = [(f"Tăng mạnh nhất ({bd.max():+}%)", bd.idxmax().split(" · ")[0]),
                (f"Giảm mạnh nhất ({bd.min():+}%)", bd.idxmin().split(" · ")[0])]
    return Result(t, kpis,
                  {"kind": "group", "data": data, "x": cap, "y": "Số HS", "color": "Năm học",
                   "horizontal": True},
                  note="Theo trường cũ của liên hệ tuyển sinh. Biểu đồ: số liên hệ 12 địa bàn "
                       "nhiều nhất; bảng có nhập học, tỷ lệ và tăng/giảm. Tăng/giảm mạnh nhất chỉ "
                       "xét địa bàn có từ 5 liên hệ ở năm trước.")


def _thang_chu_ky(d: pd.Series, nam_hoc: pd.Series) -> pd.Series:
    """Số tháng tính từ tháng 9 năm khai giảng (âm = trước khai giảng)."""
    start = pd.to_numeric(nam_hoc.str[:4], errors="coerce")
    return (d.dt.year * 12 + d.dt.month) - (start * 12 + 9)


def xh_thang(c: Ctx) -> Result:
    df, nams = _xh(c)
    d = pd.to_datetime(df["NgayLienHe"], errors="coerce")
    df = df.assign(_k=_thang_chu_ky(d, df["NamHoc"]), _d=d).dropna(subset=["_k"])
    df["_k"] = df["_k"].astype(int)

    def nhan(k: int) -> str:
        m = (k + 8) % 12 + 1
        y = (k + 8) // 12  # 0 = năm khai giảng N
        return f"{m:02d}/N" + (f"{y:+d}" if y else "")

    t = pd.crosstab(df["_k"], df["NamHoc"]).reindex(columns=nams, fill_value=0).sort_index()
    luy_ke = t.cumsum().add_prefix("Lũy kế ")
    t = t.join(luy_ke)
    t.insert(0, "Tháng", [nhan(k) for k in t.index])
    t = t.reset_index(drop=True)
    data = t.melt("Tháng", nams, var_name="Năm học", value_name="Liên hệ mới")
    return Result(t, [],
                  {"kind": "line_cat", "data": data, "x": "Tháng", "y": "Liên hệ mới",
                   "color": "Năm học"},
                  note="Liên hệ mới theo tháng, căn theo chu kỳ tuyển sinh để so sánh các năm: "
                       "N = năm khai giảng (vd năm học 2026-2027 thì N = 2026; 05/N-1 = tháng 5/2025).")


# ------------------------------------------------------------------ xu hướng nhập học
HS_CU = ("Học tiếp", "Không học tiếp")   # học sinh cũ của trường (không qua tuyển sinh)
NGHI = ("Rút hồ sơ", "Không học tiếp")


def _xnh(c: Ctx) -> tuple[pd.DataFrame, list[str]]:
    """Hồ sơ nhập học các năm được chọn + lọc khối; thêm cột Loai (HS mới / HS cũ)."""
    nh = c.nh_all if c.nh_all is not None else c.nh
    df = nh.assign(NamHoc=_blank(nh["NamHoc"]), TT=_blank(nh["TinhTrangHS"]))
    nams = [n for n in (c.params.get("nam") or sorted(df["NamHoc"].unique())) if n != "Chưa rõ"]
    df = df[df["NamHoc"].isin(nams)]
    if c.params.get("khoi"):
        df = df[df["Khoi"].isin(c.params["khoi"])]
    df = df.assign(Loai=["HS cũ" if t in HS_CU else "HS mới" for t in df["TT"]])
    return df, sorted(nams)


def _dang_hoc(df: pd.DataFrame, c: Ctx) -> pd.DataFrame:
    """Học sinh đang học (bỏ rút hồ sơ / không học tiếp), lọc theo đối tượng đã chọn."""
    df = df[~df["TT"].isin(NGHI)]
    loai = c.params.get("loai") or "Tất cả"
    if loai == "HS mới":
        df = df[df["Loai"] == "HS mới"]
    elif loai == "HS cũ học tiếp":
        df = df[df["Loai"] == "HS cũ"]
    return df


def _dem(df: pd.DataFrame, nams: list[str]) -> pd.Series:
    return df["NamHoc"].value_counts().reindex(nams, fill_value=0)


def xnh_tongquan(c: Ctx) -> Result:
    df, nams = _xnh(c)
    if df.empty:
        return Result(pd.DataFrame(columns=["Năm học"]))
    moi, cu = df[df["Loai"] == "HS mới"], df[df["Loai"] == "HS cũ"]
    t = pd.DataFrame(index=nams)
    t["Hồ sơ nhập học"] = _dem(df, nams)
    t["HS mới (qua tuyển sinh)"] = _dem(moi, nams)
    t["HS mới rút hồ sơ"] = _dem(moi[moi["TT"] == "Rút hồ sơ"], nams)
    t["HS mới nhập học"] = t["HS mới (qua tuyển sinh)"] - t["HS mới rút hồ sơ"]
    t["HS cũ học tiếp"] = _dem(cu[cu["TT"] == "Học tiếp"], nams)
    t["HS cũ không học tiếp"] = _dem(cu[cu["TT"] == "Không học tiếp"], nams)
    t["Sĩ số (đang học)"] = t["HS mới nhập học"] + t["HS cũ học tiếp"]
    t["Tỷ lệ HS cũ ở lại (%)"] = [_pct(a, a + b) for a, b in
                                  zip(t["HS cũ học tiếp"], t["HS cũ không học tiếp"])]
    hoc = df[~df["TT"].isin(NGHI)]
    t["Nam"] = _dem(hoc[hoc["GioiTinh"] == "Nam"], nams)
    t["Nữ"] = _dem(hoc[hoc["GioiTinh"] == "Nữ"], nams)
    t["Nội trú"] = _dem(hoc[hoc["NoiTruBanTru"].fillna("").str.startswith("Nội trú")], nams)
    t["Bán trú"] = _dem(hoc[hoc["NoiTruBanTru"].fillna("").str.startswith("Bán trú")], nams)
    for col in ("Sĩ số (đang học)", "HS mới nhập học"):
        t[f"{col.split(' (')[0]} tăng/giảm"] = ["—"] + [_pct_doi(a, b) for a, b in
                                                     zip(t[col].iloc[:-1], t[col].iloc[1:])]
    t = t.rename_axis("Năm học").reset_index()
    kpis = []
    if len(t):
        b = t.iloc[-1]
        a = t.iloc[-2] if len(t) >= 2 else None
        doi = (lambda col: f" ({_pct_doi(a[col], b[col])})") if a is not None else (lambda col: "")
        kpis = [(f"Sĩ số {b['Năm học']}", f"{int(b['Sĩ số (đang học)'])}{doi('Sĩ số (đang học)')}"),
                ("HS mới nhập học", f"{int(b['HS mới nhập học'])}{doi('HS mới nhập học')}"),
                ("HS cũ học tiếp", int(b["HS cũ học tiếp"])),
                ("Rút / không học tiếp", int(b["HS mới rút hồ sơ"] + b["HS cũ không học tiếp"]))]
    note = ("Lấy từ danh sách HỒ SƠ NHẬP HỌC (không phải Data tuyển sinh). HS mới = hồ sơ tạo "
            "từ tuyển sinh; HS cũ = học sinh của trường lên lớp (tình trạng Học tiếp / Không học "
            "tiếp). Sĩ số = HS mới nhập học + HS cũ học tiếp.")
    if len(t) >= 2 and t.iloc[-1]["HS cũ học tiếp"] == 0 and t.iloc[-2]["HS cũ học tiếp"] > 0:
        note += (f" Năm {t.iloc[-1]['Năm học']} chưa có HS cũ học tiếp — có thể chưa nhập danh "
                 "sách lên lớp, nên sĩ số năm này còn thấp.")
    return Result(t, kpis,
                  {"kind": "stack_cat", "data": t.melt("Năm học", ["HS mới nhập học",
                                                                   "HS cũ học tiếp"],
                                                       var_name="Nhóm", value_name="Số HS"),
                   "x": "Năm học", "y": "Số HS", "color": "Nhóm"},
                  note=note)


def _lop_key(lop: str) -> tuple:
    so = "".join(ch for ch in str(lop) if ch.isdigit())[:2]
    return (int(so) if so else 99, str(lop))


def _xnh_pivot(df: pd.DataFrame, by: str, nams: list[str], label: str) -> pd.DataFrame:
    t = pd.crosstab(_blank(df[by]), df["NamHoc"]).reindex(columns=nams, fill_value=0) \
        if len(df) else pd.DataFrame(columns=nams)
    t["Tổng"] = t.sum(axis=1)
    t = t.rename_axis(label).reset_index()
    t.columns.name = None
    return _bien_dong(t, nams)


def xnh_khoi(c: Ctx) -> Result:
    df, nams = _xnh(c)
    hoc = _dang_hoc(df, c)
    if hoc.empty:
        return Result(pd.DataFrame(columns=["Khối"]))
    t = _xnh_pivot(hoc, "Khoi", nams, "Khối")
    order = _khoi_order(t["Khối"]) + (["Chưa rõ"] if "Chưa rõ" in set(t["Khối"]) else [])
    t = t.set_index("Khối").reindex(order).reset_index()
    t = _total_row(t, "Khối")
    lop = _xnh_pivot(hoc.assign(LopHoc=_blank(hoc["LopHoc"], "Chưa xếp lớp")), "LopHoc", nams,
                     "Lớp")
    lop = lop.sort_values("Lớp", key=lambda s: s.map(_lop_key))
    kpis = []
    bd = _bien_dong_so(t[t["Khối"] != "Tổng cộng"].set_index("Khối"), nams)
    if len(bd):
        kpis = [("Khối tăng mạnh nhất", f"{bd.idxmax()} ({bd.max():+}%)"),
                ("Khối giảm mạnh nhất", f"{bd.idxmin()} ({bd.min():+}%)")]
    return Result(t, kpis,
                  {"kind": "group", "data": t[t["Khối"] != "Tổng cộng"].melt(
                      "Khối", nams, var_name="Năm học", value_name="Số HS"),
                   "x": "Khối", "y": "Số HS", "color": "Năm học"},
                  note="Học sinh đang học theo khối (hồ sơ nhập học). Excel: thêm sheet sĩ số "
                       "theo lớp qua các năm.",
                  sheets={"Theo_khoi": t, "Theo_lop": lop})


def xnh_chedo(c: Ctx) -> Result:
    df, nams = _xnh(c)
    hoc = _dang_hoc(df, c)
    if hoc.empty:
        return Result(pd.DataFrame(columns=["Nhóm"]))
    phan = []
    for nhom, col, f in (("Chế độ", "NoiTruBanTru", lambda s: s.str.split(" ").str[:2].str.join(" ")),
                         ("Phân hệ", "PhanHe", lambda s: s.replace("", "Thường")),
                         ("Giới tính", "GioiTinh", lambda s: s)):
        d = hoc.assign(_g=f(hoc[col].fillna("").astype(str).str.strip()))
        x = _xnh_pivot(d, "_g", nams, "Nhóm")
        x.insert(0, "Tiêu chí", nhom)
        phan.append(x)
    t = pd.concat(phan, ignore_index=True)
    data = t[t["Tiêu chí"] != "Giới tính"].melt(["Tiêu chí", "Nhóm"], nams,
                                                var_name="Năm học", value_name="Số HS")
    return Result(t, [],
                  {"kind": "group", "data": data, "x": "Nhóm", "y": "Số HS", "color": "Năm học"},
                  note="Học sinh đang học theo nội trú/bán trú, phân hệ (IEP/ESL/thường) và giới "
                       "tính qua các năm.")


def xnh_diaban(c: Ctx) -> Result:
    df, nams = _xnh(c)
    cap = _cap(c)
    hoc = _dang_hoc(df, c)
    g = DB.dia_ban_nh(hoc, _bang_quan(c))
    m = _theo_tinh(g, c)
    hoc, g = hoc[m], g[m]
    if hoc.empty:
        return Result(pd.DataFrame(columns=[cap]))
    t = _chua_ro_cuoi(_xnh_pivot(hoc.assign(_n=DB.nhan(g, cap)), "_n", nams, cap)
                      .sort_values("Tổng", ascending=False), cap).reset_index(drop=True)
    kpis = []
    bd = _bien_dong_so(t[~t[cap].str.startswith(DB.CHUA_RO)].set_index(cap), nams)
    if len(bd):
        kpis = [(f"Tăng mạnh nhất ({bd.max():+}%)", bd.idxmax().split(" · ")[0]),
                (f"Giảm mạnh nhất ({bd.min():+}%)", bd.idxmin().split(" · ")[0])]
    chua = t[t[cap].str.startswith(DB.CHUA_RO)]
    if len(chua):
        kpis.append(("Chưa rõ địa bàn (các năm)", int(chua["Tổng"].sum())))
    return Result(t, kpis,
                  {"kind": "group", "data": t.head(12).melt(cap, nams, var_name="Năm học",
                                                            value_name="Số HS"),
                   "x": cap, "y": "Số HS", "color": "Năm học", "horizontal": True},
                  note="Theo CHỖ Ở HIỆN NAY của học sinh đang học (thiếu thì lấy hộ khẩu). Quận/"
                       "huyện cũ suy ra từ phường/xã mới.")


def xnh_truongcu(c: Ctx) -> Result:
    df, nams = _xnh(c)
    hoc = _dang_hoc(df, c)
    if hoc.empty:
        return Result(pd.DataFrame(columns=["Trường cũ"]))
    hoc = hoc[hoc["TruongCu"].fillna("").str.strip() != ""]
    t = _xnh_pivot(hoc, "TruongCu", nams, "Trường cũ").sort_values("Tổng", ascending=False) \
        .head(int(c.params.get("top", 30)))
    return Result(t, [("Số trường (mọi năm)", hoc["TruongCu"].nunique())],
                  {"kind": "group", "data": t.head(12).melt("Trường cũ", nams, var_name="Năm học",
                                                            value_name="Số HS"),
                   "x": "Trường cũ", "y": "Số HS", "color": "Năm học", "horizontal": True},
                  note="Trường cũ của học sinh đang học (theo hồ sơ nhập học).")


def xnh_doichieu(c: Ctx) -> Result:
    """Đối chiếu Data tuyển sinh (bước Nhập học) với danh sách hồ sơ nhập học HS mới."""
    ts, nams = _xh(c)
    df, nams_nh = _xnh(c)
    nams = sorted(set(nams) | set(nams_nh))
    moi = df[df["Loai"] == "HS mới"]
    lien_ket = set(moi["TuyenSinhID"].fillna("").astype(str)) - {""}
    ts_nh = ts[ts["TrangThai"] == "Nhập học"]
    chua_hs = ts_nh[~ts_nh["id"].astype(str).isin(lien_ket)]
    t = pd.DataFrame(index=nams)
    t["Liên hệ (tuyển sinh)"] = _dem(ts, nams)
    t["Nhập học theo Data tuyển sinh"] = _dem(ts_nh, nams)
    t["Hồ sơ nhập học HS mới"] = _dem(moi, nams)
    t["  trong đó rút hồ sơ"] = _dem(moi[moi["TT"] == "Rút hồ sơ"], nams)
    t["  trong đó chưa liên kết tuyển sinh"] = _dem(
        moi[moi["TuyenSinhID"].fillna("").astype(str) == ""], nams)
    t["Nhập học nhưng chưa có hồ sơ"] = _dem(chua_hs, nams)
    t["Tỷ lệ liên hệ → hồ sơ nhập học (%)"] = [
        _pct(a - r, b) for a, r, b in zip(t["Hồ sơ nhập học HS mới"], t["  trong đó rút hồ sơ"],
                                          t["Liên hệ (tuyển sinh)"])]
    t = t.rename_axis("Năm học").reset_index()
    ds = chua_hs[["NamHoc", "HoTenHS", "Khoi", "SDT", "NgayLienHe"]].rename(columns={
        "NamHoc": "Năm học", "HoTenHS": "Học sinh", "Khoi": "Khối", "SDT": "SĐT",
        "NgayLienHe": "Ngày liên hệ"}).sort_values(["Năm học", "Học sinh"])
    kpis = []
    if len(t):
        b = t.iloc[-1]
        kpis = [(f"Nhập học (Data TS) {b['Năm học']}", int(b["Nhập học theo Data tuyển sinh"])),
                ("Hồ sơ HS mới còn học", int(b["Hồ sơ nhập học HS mới"] - b["  trong đó rút hồ sơ"])),
                ("Nhập học chưa có hồ sơ", int(b["Nhập học nhưng chưa có hồ sơ"]))]
    return Result(t, kpis,
                  {"kind": "group", "data": t.melt("Năm học", ["Nhập học theo Data tuyển sinh",
                                                               "Hồ sơ nhập học HS mới"],
                                                   var_name="Nguồn số liệu", value_name="Số HS"),
                   "x": "Năm học", "y": "Số HS", "color": "Nguồn số liệu"},
                  note="So sánh số nhập học ghi ở Data tuyển sinh với số hồ sơ nhập học HS mới. "
                       "Lệch nhau nghĩa là còn học sinh chưa tạo hồ sơ hoặc hồ sơ chưa liên kết. "
                       "Excel: thêm danh sách học sinh đã nhập học nhưng chưa có hồ sơ.",
                  sheets={"Doi_chieu": t, "Chua_co_ho_so": ds})


REPORTS = [
    Report("ts_khoi", G_TS, "Phễu tuyển sinh theo khối",
           "Số học sinh ở từng bước (tư vấn → nộp hồ sơ → nhập học) theo khối, tỷ lệ chuyển đổi.",
           ts_khoi),
    Report("ts_nguon", G_TS, "Hiệu quả nguồn tuyển sinh",
           "Liên hệ, nộp hồ sơ, nhập học và tỷ lệ chuyển đổi của từng nguồn.", ts_nguon),
    Report("ts_thoigian", G_TS, "Liên hệ theo thời gian",
           "Số liên hệ mới theo tháng / tuần, lũy kế và số đã nhập học.", ts_thoigian, ("ky",)),
    Report("ts_tuvan", G_TS, "Kết quả theo người phụ trách",
           "Số học sinh mỗi người nhận hồ sơ phụ trách và kết quả chuyển đổi.", ts_tuvan),
    Report("ts_tinhtrang", G_TS, "Tình trạng tư vấn",
           "Phân bố tình trạng tư vấn (cần tư vấn thêm, đang cân nhắc…) theo khối.",
           ts_tinhtrang),
    Report("ts_truongcu", G_TS, "Trường cũ có nhiều học sinh",
           "Top trường cũ theo số liên hệ và số nhập học.", ts_truongcu, ("top",)),
    Report("ts_diaban", G_TS, "Địa bàn tuyển sinh",
           "Liên hệ, nộp hồ sơ, nhập học theo tỉnh/thành → quận/huyện (cũ) → phường/xã (mới) của "
           "trường cũ.", ts_diaban, ("cap", "tinh")),
    Report("ts_chedo", G_TS, "Chế độ & phân hệ nhập học",
           "Học sinh nhập học theo khối × nội trú/bán trú × IEP/ESL.", ts_chedo),
    Report("ts_quahan", G_TS, "Liên hệ tư vấn quá hạn",
           "Danh sách học sinh ở bước Tư vấn quá lâu chưa chuyển bước — cần gọi lại.",
           ts_quahan, ("ngay",)),
    Report("ts_rut", G_TS, "Danh sách rút hồ sơ", "Học sinh rút hồ sơ kèm lý do.", ts_rut),
    Report("kt_giucho", G_KT, "Giữ chỗ theo khối",
           "Số học sinh đã giữ chỗ, hủy, hoàn phí và tổng tiền giữ chỗ theo khối.", kt_giucho),
    Report("kt_hocphi", G_KT, "Tổng hợp học phí theo khối",
           "Tiền PHHS thanh toán, đã thu, còn lại và số học sinh còn nợ theo khối.", kt_hocphi),
    Report("kt_congno", G_KT, "Danh sách còn nợ học phí",
           "Học sinh còn số tiền phải thu.", kt_congno),
    Report("kt_hoanphi", G_KT, "Danh sách hủy giữ chỗ / hoàn phí",
           "Học sinh hủy giữ chỗ chờ hoàn phí và đã hoàn phí, kèm tài khoản nhận.", kt_hoanphi),
    Report("nh_siso", G_NH, "Sĩ số theo lớp",
           "Sĩ số, nam/nữ, nội trú/bán trú, IEP/ESL của từng lớp.", nh_siso),
    Report("nh_tinhtrang", G_NH, "Tình trạng hồ sơ nhập học",
           "Đang nhập hồ sơ / đang đóng phí / đã đóng phí / rút… theo khối.", nh_tinhtrang),
    Report("nh_hoanthien", G_NH, "Mức hoàn thiện hồ sơ (VEMIS)",
           "Số hồ sơ đủ thông tin tối thiểu theo lớp và các thông tin hay thiếu.", nh_hoanthien),
    Report("nh_giayto", G_NH, "Giấy tờ đã thu",
           "Giấy tờ cần nộp theo từng khối: đã nộp (bản gốc / bản sao), còn thiếu, danh sách "
           "học sinh chưa nộp đủ.",
           nh_giayto),
    Report("nh_xe", G_NH, "Đăng ký xe đưa đón",
           "Số học sinh nội trú theo tuyến xe và danh sách đón trả từng tuyến.", nh_xe),
    Report("nh_tohop", G_NH, "Tổ hợp môn lựa chọn (THPT)",
           "Số học sinh khối 10–12 theo tổ hợp môn lựa chọn 1.", nh_tohop),
    Report("nh_diaban", G_NH, "Địa bàn cư trú học sinh",
           "Học sinh đang học theo tỉnh/thành → quận/huyện (cũ) → phường/xã chỗ ở, kèm số theo "
           "khối.", nh_diaban, ("cap", "tinh")),
    Report("nh_hocluc", G_NH, "Kết quả học tập đầu vào",
           "Học lực kết quả 1 theo khối và điểm trung bình Toán / Văn / Anh.", nh_hocluc),
    Report("nh_danhsach", G_NH, "Danh sách học sinh theo lớp",
           "Danh sách từng lớp (STT, họ tên, ngày sinh, chế độ, liên lạc) để in.", nh_danhsach),
]
REPORTS += [
    Report("xh_tongquan", G_XH, "Phễu tuyển sinh qua các năm",
           "Data tuyển sinh: liên hệ, nộp hồ sơ, nhập học, rút hồ sơ, tỷ lệ chuyển đổi và "
           "tăng/giảm từng năm.", xh_tongquan, ("nam",)),
    Report("xh_khoi", G_XH, "Chốt nhập học theo khối qua các năm",
           "Data tuyển sinh: số học sinh ở bước Nhập học từng khối của các năm (Excel có thêm số "
           "liên hệ).", xh_khoi, ("nam",)),
    Report("xh_nguon", G_XH, "Nguồn tuyển sinh qua các năm",
           "Liên hệ, nhập học và tỷ lệ nhập học của từng nguồn theo năm.", xh_nguon, ("nam",)),
    Report("xh_truongcu", G_XH, "Trường cũ qua các năm",
           "Top trường cũ theo số liên hệ của từng năm, trường tăng/giảm.", xh_truongcu,
           ("nam", "top")),
    Report("xh_diaban", G_XH, "Địa bàn tuyển sinh qua các năm",
           "Liên hệ, nhập học, tỷ lệ theo tỉnh/thành → quận/huyện → phường/xã của trường cũ.",
           xh_diaban, ("nam", "cap", "tinh")),
    Report("xh_thang", G_XH, "Tiến độ tuyển sinh theo tháng qua các năm",
           "Liên hệ mới mỗi tháng và lũy kế, căn theo chu kỳ tuyển sinh để so cùng kỳ.",
           xh_thang, ("nam",)),
    Report("xnh_tongquan", G_XNH, "Sĩ số và học sinh mới qua các năm",
           "Hồ sơ nhập học: HS mới nhập học, HS cũ học tiếp / không học tiếp, rút hồ sơ, sĩ số, "
           "nam/nữ, nội trú/bán trú.", xnh_tongquan, ("nam",)),
    Report("xnh_doichieu", G_XNH, "Đối chiếu tuyển sinh và nhập học",
           "Số nhập học ở Data tuyển sinh so với số hồ sơ nhập học HS mới; danh sách đã nhập học "
           "nhưng chưa có hồ sơ.", xnh_doichieu, ("nam",)),
    Report("xnh_khoi", G_XNH, "Học sinh theo khối và lớp qua các năm",
           "Học sinh đang học theo khối (Excel: sĩ số từng lớp) và tăng/giảm.", xnh_khoi,
           ("nam", "loai")),
    Report("xnh_chedo", G_XNH, "Chế độ, phân hệ, giới tính qua các năm",
           "Nội trú/bán trú, IEP/ESL, nam/nữ của học sinh đang học.", xnh_chedo, ("nam", "loai")),
    Report("xnh_diaban", G_XNH, "Địa bàn cư trú qua các năm",
           "Học sinh đang học theo tỉnh/thành → quận/huyện → phường/xã chỗ ở.", xnh_diaban,
           ("nam", "loai", "cap", "tinh")),
    Report("xnh_truongcu", G_XNH, "Trường cũ của học sinh qua các năm",
           "Top trường cũ của học sinh đang học, tăng/giảm.", xnh_truongcu,
           ("nam", "loai", "top")),
]
BY_ID = {r.id: r for r in REPORTS}


# ------------------------------------------------------------------ xuất Excel
def excel(items: list[tuple[str, str, pd.DataFrame]], school: str, nam_hoc: str) -> bytes:
    """Workbook có tiêu đề trường, tên báo cáo, ngày lập; mỗi mục một sheet."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    buf = io.BytesIO()
    thin = Side(style="thin", color="CBD5E1")
    used = set()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for title, sheet, df in items:
            name = "".join(ch for ch in sheet if ch not in '[]:*?/\\')[:31] or "Sheet"
            base, i = name, 2
            while name in used:
                name = f"{base[:28]}_{i}"
                i += 1
            used.add(name)
            df = df.copy()
            for col in df.columns:
                if df[col].dtype == object:
                    df[col] = df[col].map(lambda v: v.strftime("%d/%m/%Y")
                                          if isinstance(v, (date, datetime)) and pd.notna(v)
                                          else "" if v is pd.NaT else v)
            df.to_excel(xw, sheet_name=name, index=False, startrow=4)
            ws = xw.sheets[name]
            ws["A1"] = school.upper()
            ws["A1"].font = Font(bold=True, size=11, color="1D4ED8")
            ws["A2"] = title.upper()
            ws["A2"].font = Font(bold=True, size=14)
            ws["A3"] = f"Năm học {nam_hoc} · Ngày lập {date.today():%d/%m/%Y}"
            ws["A3"].font = Font(italic=True, size=10, color="64748B")
            for c in ws[5]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="1D4ED8")
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.row_dimensions[5].height = 32
            for row in ws.iter_rows(min_row=5, max_row=ws.max_row, max_col=max(1, df.shape[1])):
                for c in row:
                    c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
                    if isinstance(c.value, (int, float)) and c.row > 5:
                        hdr = str(ws.cell(5, c.column).value or "")
                        c.number_format = "#,##0" if "(đ)" in hdr or "Tiền" in hdr else \
                            "0.0" if "(%)" in hdr or " TB" in hdr else "#,##0"
            last = ws.max_row
            if last > 5 and any(str(ws.cell(last, j).value) == "Tổng cộng" for j in (1, 2)):
                for c in ws[last]:
                    c.font = Font(bold=True)
                    c.fill = PatternFill("solid", fgColor="EEF2FF")
            for j, col in enumerate(df.columns, 1):
                w = max([len(str(col))] + [len(str(v)) for v in df[col].head(300)])
                ws.column_dimensions[get_column_letter(j)].width = min(max(8, w + 2), 50)
            ws.freeze_panes = "A6"
            ws.print_title_rows = "5:5"
            ws.page_setup.orientation = "landscape" if df.shape[1] > 6 else "portrait"
            ws.page_setup.fitToWidth = 1
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.fitToHeight = 0
    return buf.getvalue()


def export_items(rep: Report, res: Result) -> list[tuple[str, str, pd.DataFrame]]:
    if res.sheets:
        return [(f"{rep.title} — {k.replace('_', ' ')}", k, v) for k, v in res.sheets.items()]
    return [(rep.title, rep.title, res.table)]
