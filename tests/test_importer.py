import sys
from pathlib import Path

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from tuyensinh import importer, services  # noqa: E402
from tuyensinh.schema import TUYEN_SINH  # noqa: E402
from tuyensinh.storage.local import LocalStorage  # noqa: E402


def _df():
    return pd.DataFrame([
        {"Title": "A", "HoTenHS": "Nguyễn A", "SDT": "0901111111", "NgayLienHe": "2026-05-01",
         "Khoi": "10", "NamHoc": "2026-2027"},
        # trùng hệt dòng trên (có trong dữ liệu cũ) -> vẫn phải giữ đủ 2 dòng
        {"Title": "A", "HoTenHS": "Nguyễn A", "SDT": "0901111111", "NgayLienHe": "2026-05-01",
         "Khoi": "10", "NamHoc": "2026-2027"},
        # thiếu ngày liên hệ, không có năm học -> dùng năm học mặc định
        {"Title": "B", "HoTenHS": "Trần B", "SDT": "0902222222", "Khoi": "6"},
    ])


def test_import_chay_lai_khong_trung(tmp_path):
    st = LocalStorage(tmp_path / "t.sqlite3")
    r1 = importer.import_rows(st, TUYEN_SINH, _df(), "2025-2026", workers=3)
    assert (r1["ok"], r1["skipped"], r1["errors"]) == (3, 0, [])
    r2 = importer.import_rows(st, TUYEN_SINH, _df(), "2025-2026", workers=3)
    assert (r2["ok"], r2["skipped"]) == (0, 3)
    df = services.load_df(st, TUYEN_SINH)
    assert len(df) == 3 and set(df["NamHoc"]) == {"2026-2027", "2025-2026"}


def test_import_kiem_tra_bat_buoc(tmp_path):
    st = LocalStorage(tmp_path / "t.sqlite3")
    r = importer.import_rows(st, TUYEN_SINH, _df(), "2025-2026", keep_all=False, workers=1)
    assert r["ok"] == 2 and len(r["errors"]) == 1 and "Ngày liên hệ" in r["errors"][0]["Lỗi"]


def test_read_frame_nhan_ten_hien_thi(tmp_path):
    p = tmp_path / "x.xlsx"
    pd.DataFrame({"Họ tên HS": ["A"], "SĐT": ["0901"], "Nam hoc": ["2026-2027"],
                  "Cột lạ": [1]}).to_excel(p, index=False)
    df, skipped = importer.read_frame(str(p), TUYEN_SINH)
    assert list(df.columns) == ["HoTenHS", "SDT", "NamHoc"] and skipped == ["Cột lạ"]


def test_link_tuyen_sinh(tmp_path):
    from tuyensinh import importer
    from tuyensinh.schema import TUYEN_SINH
    from tuyensinh.storage.local import LocalStorage

    st = LocalStorage(str(tmp_path / "t.db"))
    a = st.create_item(TUYEN_SINH.name, {"HoTenHS": "Trần Gia Bảo", "SDT": "0903750252",
                                         "NamHoc": "2025-2026"})
    b = st.create_item(TUYEN_SINH.name, {"HoTenHS": "Hồ Duy", "NgaySinh": "2013-09-08",
                                         "NamHoc": "2025-2026"})
    rows = [{"HoTen": "TRẦN GIA BẢO", "DienThoaiMe": "0903 750 252", "NamHoc": "2026-2027"},
            {"HoTen": "Hồ Duy", "NgaySinh": "2013-09-08 00:00:00", "NamHoc": "2025-2026"},
            {"HoTen": "Người Khác", "DienThoaiSLL": "0903750252"}]
    assert importer.link_tuyen_sinh(st, rows) == 2
    assert rows[0]["TuyenSinhID"] == str(a["id"]) and rows[1]["TuyenSinhID"] == str(b["id"])
    assert "TuyenSinhID" not in rows[2]


def test_link_existing(tmp_path):
    from tuyensinh import importer
    from tuyensinh.schema import NHAP_HOC, TUYEN_SINH
    from tuyensinh.storage.local import LocalStorage

    st = LocalStorage(str(tmp_path / "t.db"))
    ts = st.create_item(TUYEN_SINH.name, {"HoTenHS": "An", "SDT": "0901234567", "NamHoc": "2026-2027"})
    nh = st.create_item(NHAP_HOC.name, {"HoTen": "An", "DienThoaiSLL": "0901234567",
                                        "NamHoc": "2026-2027"})
    res = importer.link_existing(st)
    assert res == {"total": 1, "chua_lien_ket": 1, "lien_ket": 1}
    assert str(st.get_item(NHAP_HOC.name, nh["id"])["TuyenSinhID"]) == str(ts["id"])


def test_fill_blanks(tmp_path):
    import pandas as pd
    from tuyensinh import importer
    from tuyensinh.schema import NHAP_HOC
    from tuyensinh.storage.local import LocalStorage

    st = LocalStorage(str(tmp_path / "t.db"))
    it = st.create_item(NHAP_HOC.name, {"HoTen": "An", "NgaySinh": "2012-01-02",
                                        "NamHoc": "2026-2027", "LopHoc": "7.1"})
    df = pd.DataFrame([{"HoTen": "An", "NgaySinh": "2012-01-02", "NamHoc": "2026-2027",
                        "Khoi": "7", "LopHoc": "7.9"},
                       {"HoTen": "Không có", "NgaySinh": "2012-01-02", "NamHoc": "2026-2027"}])
    res = importer.fill_blanks(st, NHAP_HOC, df, "2026-2027")
    assert res == {"updated": 1, "cells": 1, "not_found": 1, "total": 2}
    got = st.get_item(NHAP_HOC.name, it["id"])
    assert got["Khoi"] == "7" and got["LopHoc"] == "7.1"  # không ghi đè ô đã có


def test_file_mau_nhap_lai_duoc(tmp_path):
    """File mẫu: điền vài dòng (ngày gõ dd/mm/yyyy hoặc ô ngày thật) -> nhập lên list được."""
    import datetime
    import io

    from openpyxl import load_workbook

    from tuyensinh import importer
    from tuyensinh.schema import NHAP_HOC
    from tuyensinh.storage.local import LocalStorage

    wb = load_workbook(io.BytesIO(importer.file_mau(NHAP_HOC)))
    assert wb.sheetnames[0] == NHAP_HOC.name and "HuongDan" in wb.sheetnames
    ws = wb[NHAP_HOC.name]
    hdr = [c.value for c in ws[1]]
    lab = {f.key: f.label for f in NHAP_HOC.fields}
    assert lab["HoTen"] in hdr and "ID tuyển sinh" not in hdr
    for r, (ten, ns) in enumerate((("Nguyễn Văn Mẫu", "01/05/2012"),
                                   ("Trần Thị Thử", datetime.datetime(2013, 2, 3))), 2):
        ws.cell(r, hdr.index(lab["HoTen"]) + 1, ten)
        ws.cell(r, hdr.index(lab["NgaySinh"]) + 1, ns)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    df, skipped = importer.read_frame(buf, NHAP_HOC)
    assert not skipped and df["NgaySinh"].tolist() == ["2012-05-01", "2013-02-03"]
    st = LocalStorage(tmp_path / "m.sqlite3")
    assert importer.import_rows(st, NHAP_HOC, df, "2026-2027", keep_all=False)["ok"] == 2
    assert importer.import_rows(st, NHAP_HOC, df, "2026-2027", keep_all=False)["skipped"] == 2


def _dien(ld, rows):
    """Điền file mẫu của `ld` theo nhãn cột rồi đọc lại như khi tải lên."""
    import io

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(importer.file_mau(ld)))
    assert wb.sheetnames[0] == ld.name and "HuongDan" in wb.sheetnames
    ws = wb[ld.name]
    hdr = [c.value for c in ws[1]]
    for r, row in enumerate(rows, 2):
        for k, v in row.items():
            ws.cell(r, hdr.index(k) + 1, v)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    buf.name = "mau.xlsx"
    return buf, hdr


def test_mau_du_cho_moi_list():
    from tuyensinh.schema import ALL_LISTS
    for ld in ALL_LISTS:
        buf, hdr = _dien(ld, [])
        assert importer.guess_list([ld.name]) is ld
        # không có cột app tự quản lý
        assert not {"ID tuyển sinh", "Mật khẩu (đã mã hóa)", "Đánh giá từng học sinh (JSON)",
                    "ID phòng", "Học sinh trong phòng (ID hồ sơ nhập học)"} & set(hdr)


def test_nhap_giay_to_va_tai_khoan(tmp_path):
    from tuyensinh import tai_khoan
    from tuyensinh.schema import GIAY_TO, TAI_KHOAN
    st = LocalStorage(tmp_path / "g.sqlite3")
    buf, _ = _dien(GIAY_TO, [{"Khối (vd 10-IEP, 7)": "6", "Tên giấy tờ": "Khai sinh", "Thứ tự": 1},
                             {"Khối (vd 10-IEP, 7)": "6", "Tên giấy tờ": "Học bạ", "Thứ tự": 2}])
    df, skipped = importer.read_frame(buf, GIAY_TO)
    assert not skipped
    assert importer.import_rows(st, GIAY_TO, df, "2026-2027")["ok"] == 2
    assert importer.import_rows(st, GIAY_TO, df, "2026-2027")["skipped"] == 2

    buf, _ = _dien(TAI_KHOAN, [
        {"Họ tên": "Cô Hoa", "Tên đăng nhập": "Hoa", "Mật khẩu": "abc123",
         "Vai trò": "GV quản nhiệm"},
        {"Họ tên": "Thầy B", "Tên đăng nhập": "b", "Mật khẩu": "123", "Vai trò": "Kế toán"},
        {"Họ tên": "X", "Tên đăng nhập": "x", "Mật khẩu": "abcdef", "Vai trò": "Bảo vệ"}])
    df, skipped = importer.read_frame(buf, TAI_KHOAN)
    assert not skipped
    res = importer.nhap_tai_khoan(st, df, [])
    assert res["ok"] == 1 and [e["Dòng"] for e in res["errors"]] == [3, 4]
    items = st.list_items(TAI_KHOAN.name)
    assert items[0]["MatKhau"] != "abc123" and tai_khoan.dang_nhap(items, "hoa", "abc123")
    assert importer.nhap_tai_khoan(st, df.iloc[:1], items)["skipped"] == 1
