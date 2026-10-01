import io
import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from tuyensinh import danh_muc, export_vemis, services  # noqa: E402
from tuyensinh.schema import NHAP_HOC, TUYEN_SINH  # noqa: E402
from tuyensinh.storage.local import LocalStorage  # noqa: E402
from tuyensinh.storage.convert import ColumnMap, date_in  # noqa: E402

HS = {"NamHoc": "2026-2027", "NgayLienHe": "2026-08-06", "SDT": "0927 666 649",
      "HoTenHS": "Trần Huy Long", "Khoi": "10", "CheDo": "Nội trú", "GioiTinh": "Nam",
      "NgaySinh": "2011-05-02", "TruongCu": "THCS Tân Phú"}


@pytest.fixture
def storage(tmp_path):
    return LocalStorage(tmp_path / "t.sqlite3")


def test_save_validates_required(storage):
    with pytest.raises(ValueError, match="Họ tên HS"):
        services.save_tuyen_sinh(storage, {**HS, "HoTenHS": ""})
    with pytest.raises(ValueError, match="SĐT"):
        services.save_tuyen_sinh(storage, {**HS, "SDT": "123"})


def test_phone_normalized_and_default_status(storage):
    rec = services.save_tuyen_sinh(storage, HS)
    assert rec["SDT"] == "0927666649"
    assert rec["TrangThai"] == "Tư vấn"
    assert services.normalize_phone("+84 927 666 649") == "0927666649"


def test_nhap_hoc_creates_enrollment_record(storage):
    rec = services.save_tuyen_sinh(storage, HS)
    rec = services.set_trang_thai(storage, rec["id"], "Nhập học")
    assert rec["TrangThai"] == "Nhập học"
    nh = services.nhap_hoc_of(storage, rec["id"])
    assert nh["HoTen"] == "Trần Huy Long"
    assert nh["LopHoc"] == "10" and nh["NoiTruBanTru"] == "Nội trú"
    assert nh["DienThoaiSLL"] == "0927666649"
    assert nh["TruongCu"] == "THCS Tân Phú"
    # Chuyển lại không tạo trùng; hồ sơ bị ẩn khi rút hồ sơ
    services.set_trang_thai(storage, rec["id"], "Nhập học")
    assert len(storage.list_items(NHAP_HOC.name)) == 1
    rec = services.set_trang_thai(storage, rec["id"], "Rút hồ sơ", "Chuyển trường")
    assert "rút hồ sơ: Chuyển trường" in rec["GhiChu"]
    assert services.nhap_hoc_df(storage, "2026-2027").empty


def test_status_via_form_save(storage):
    rec = services.save_tuyen_sinh(storage, {**HS, "TrangThai": "Nhập học"})
    assert services.nhap_hoc_of(storage, rec["id"]) is not None


def test_duplicates(storage):
    services.save_tuyen_sinh(storage, HS)
    assert len(services.find_duplicates(storage, {**HS, "SDT": "0927666649"})) == 1


def test_giu_cho(storage):
    rec = services.save_tuyen_sinh(storage, {**HS, "TrangThai": "Nộp hồ sơ"})
    with pytest.raises(ValueError):
        services.xac_nhan_giu_cho(storage, rec["id"], 0, "Kế toán A")
    services.xac_nhan_giu_cho(storage, rec["id"], 2_000_000, "Kế toán A")
    r2 = services.save_tuyen_sinh(storage, {**HS, "HoTenHS": "B", "Khoi": "6"})
    services.xac_nhan_giu_cho(storage, r2["id"], 2_000_000, "Kế toán A", "Hủy giữ chỗ",
                              NganHang="Vietcombank", SoTaiKhoan="0071")
    ts = services.load_df(storage, TUYEN_SINH, "2026-2027")
    assert ts.set_index("id").loc[r2["id"], "NganHang"] == "Vietcombank"
    sm = services.giu_cho_summary(ts).set_index("Khối")
    assert sm.loc["10", "Tiền đã giữ chỗ"] == 2_000_000 and sm.loc["6", "Hủy giữ chỗ"] == 1


def test_scores_empty_vs_zero(storage):
    rec = services.save_tuyen_sinh(storage, {**HS, "Toan1": 8.5, "Van1": 0})
    row = services.load_df(storage, TUYEN_SINH).iloc[0]
    assert row["Toan1"] == 8.5 and row["Van1"] == 0 and pd.isna(row["Anh1"])
    services.save_tuyen_sinh(storage, row.to_dict(), rec["id"])
    assert storage.get_item(TUYEN_SINH.name, rec["id"])["Anh1"] == ""


def test_export_vemis_layout(storage):
    rec = services.save_tuyen_sinh(storage, {**HS, "TrangThai": "Nhập học"})
    nh = services.nhap_hoc_of(storage, rec["id"])
    services.save_nhap_hoc(storage, {**nh, "DoanVien": True, "ChoO_Tinh": "Thành phố Hồ Chí Minh"},
                           nh["id"])
    df = services.nhap_hoc_df(storage, "2026-2027")
    wb = load_workbook(io.BytesIO(export_vemis.to_bytes(df)))
    ws = wb["Sheet1"]
    assert len(export_vemis.COLUMNS) == 54
    assert ws.cell(7, 12).value == "Chỗ ở hiện nay" and ws.cell(8, 12).value == "SN/Xóm"
    assert ws.cell(7, 54).value == "Ghi chú"
    row = [c.value for c in ws[9]]
    assert row[0] == 1 and row[6] == "Trần Huy Long" and row[7] == "02/05/2011"
    assert row[14] == "Thành phố Hồ Chí Minh" and row[34] == "x"
    assert row[46] == "0927666649"


def test_every_enrollment_field_is_exported():
    exported = {k for _, _, k in export_vemis.COLUMNS if k}
    from tuyensinh.schema import NHAP_HOC_NGOAI_VEMIS

    assert exported == set(NHAP_HOC.keys) - set(NHAP_HOC_NGOAI_VEMIS)


def test_danh_muc_xa_depends_on_tinh():
    f = NHAP_HOC.get("ChoO_Xa")
    assert danh_muc.options_for(f, {}) == []
    xa = danh_muc.options_for(f, {"ChoO_Tinh": "Thành phố Hồ Chí Minh"})
    assert any("Tân Phú" in x for x in xa)
    assert "Kinh" in danh_muc.options_for(NHAP_HOC.get("DanToc"))


# Cột thật của list "Data tuyển sinh" (file data_tuyen_sinh.xlsx), tên nội bộ bị mã hóa
SP_COLUMNS = [{"name": "Title", "title": "Title", "type": "Text"}] + [
    {"name": f"c{i}_x00e0_", "title": t, "type": ty} for i, (t, ty) in enumerate([
        ("Ngày liên hệ", "DateTime"), ("SĐT", "Text"), ("Nguồn", "Choice"),
        ("Tài khoản FB", "Text"), ("Người giới thiệu", "Text"), ("Họ tên HS", "Text"),
        ("Khối", "Choice"), ("Giới tính", "Choice"), ("Chế độ", "Choice"), ("Trường cũ", "Text"),
        ("Trường cũ_Quận huyện", "Text"), ("Trường cũ_tỉnh", "Text"), ("Tình trạng", "Choice"),
        ("Nội dung đã trao đổi", "Note"), ("Bước", "Choice"), ("Phân hệ", "Text"),
        ("Người nhận hồ sơ", "User"), ("Số tiền xác nhận", "Number"),
        ("Người xác nhận", "Text"), ("Tên chủ tài khoản", "Text"), ("Ngân hàng", "Text"),
        ("Số tài khoản", "Text"), ("Toán 1", "Text"), ("Văn 1", "Text"), ("Anh 1", "Text"),
        ("TV 1", "Text"), ("Toán 2", "Text"), ("Văn 2", "Text"), ("Tiếng Anh 2", "Text"),
        ("TV 2", "Text"), ("Hạnh kiểm 1", "Text"), ("Hạnh kiểm 2", "Text"),
        ("Ngày sinh", "DateTime"), ("Nam hoc", "Text")])]


def test_column_map_matches_existing_list_by_display_name():
    cm = ColumnMap(TUYEN_SINH.name, SP_COLUMNS)
    assert cm.missing == ["TruongCu_DiaChiCu", "GiuCho"]  # 2 cột chỉ có ở list mới
    by_title = {c["title"]: c["name"] for c in SP_COLUMNS}
    assert cm.internal["NamHoc"] == by_title["Nam hoc"]
    assert cm.internal["TrangThai"] == by_title["Bước"]
    assert cm.internal["TenLienHe"] == by_title["Tài khoản FB"]
    out = cm.to_sp({"HoTenHS": "A", "NgaySinh": "2011-05-02", "Toan1": 8.0, "SoTienXacNhan": "2e6",
                    "NguoiNhanHoSo": "x", "Created": "…", "id": "5"})
    assert out == {by_title["Họ tên HS"]: "A", by_title["Ngày sinh"]: "2011-05-02T12:00:00Z",
                   by_title["Toán 1"]: "8", by_title["Số tiền xác nhận"]: 2_000_000.0,
                   "Title": "A"}  # cột Người (User) không ghi
    back = cm.from_sp({by_title["Ngày liên hệ"]: "2025-02-10T17:00:00Z",
                       by_title["Toán 1"]: "8,5", by_title["Người nhận hồ sơ"]: {"Title": "Uyên"},
                       "OData__x0020_": 1}, 7)
    assert back["NgayLienHe"] == "2025-02-11"  # 0h 11/2 giờ VN
    assert back["Toan1"] == 8.5 and back["NguoiNhanHoSo"] == "Uyên" and back["id"] == "7"


def test_date_in():
    assert date_in("2026-08-06") == "2026-08-06"
    assert date_in("2026-08-06T12:00:00Z") == "2026-08-06"


def test_to_df_types():
    df = services.to_df(TUYEN_SINH, [{"id": "1", "SoTienXacNhan": "100", "NgaySinh": "2026-01-02"}])
    assert df.loc[0, "SoTienXacNhan"] == 100 and str(df.loc[0, "NgaySinh"]) == "2026-01-02"
    assert isinstance(df, pd.DataFrame)


def test_empty_dates_roundtrip_from_dataframe(storage):
    rec = services.save_tuyen_sinh(storage, {**HS, "NgaySinh": ""})
    row = services.load_df(storage, TUYEN_SINH).iloc[0].to_dict()
    services.save_tuyen_sinh(storage, row, rec["id"])
    assert storage.get_item(TUYEN_SINH.name, rec["id"])["NgaySinh"] == ""


def test_column_map_new_list_uses_internal_names():
    # List mới tạo từ file chuẩn hóa: tên nội bộ = key, tên hiển thị tùy ý
    cols = [{"name": f.key, "title": "x" + f.key, "type": "Text"} for f in TUYEN_SINH.fields]
    cm = ColumnMap(TUYEN_SINH.name, cols)
    assert cm.missing == [] and cm.internal["TruongCu_PhuongXa"] == "TruongCu_PhuongXa"


def test_tao_ho_so_hang_loat(storage):
    a = services.save_tuyen_sinh(storage, {**HS, "TrangThai": "Nhập học"})  # đã có hồ sơ
    b = services.save_tuyen_sinh(storage, {**HS, "HoTenHS": "B"})
    c = services.save_tuyen_sinh(storage, {**HS, "HoTenHS": "C"})
    n = services.tao_ho_so_hang_loat(storage, [a, b, c], workers=2)
    assert n == 2 and len(storage.list_items(NHAP_HOC.name)) == 3
    assert services.tao_ho_so_hang_loat(storage, [a, b, c]) == 0


def test_phone_query():
    assert services.phone_query("0909 123 456") == "0909123456"
    assert services.phone_query("+84 909123456") == "0909123456"
    assert services.phone_query("Đồng Bộ 772") == ""
    assert services.phone_query("") == ""


def test_doc_ghi_ho_so():
    from datetime import date as _d
    text = ("Ngày đã nộp 23/06/2025\nPhiếu ĐK nhập học - Bản gốc\nGiấy khai sinh - Bản sao\n"
            "Học bạ THCS -")
    ngay, docs = services.doc_ho_so(text)
    assert ngay == _d(2025, 6, 23)
    assert docs == {"Phiếu đăng ký nhập học": "Bản gốc", "Giấy khai sinh": "Bản sao",
                    "Học bạ THCS": "Bản gốc"}
    assert services.doc_ho_so(services.ghi_ho_so(ngay, docs)) == (ngay, docs)


def test_rut_ho_so_nhap_hoc(storage):
    rec = services.save_tuyen_sinh(storage, dict(HS))
    services.set_trang_thai(storage, rec["id"], "Nhập học")
    nh = services.nhap_hoc_of(storage, rec["id"])
    services.rut_ho_so_nhap_hoc(storage, nh, "chuyển trường", "Cô Lan")
    assert storage.get_item(NHAP_HOC.name, nh["id"])["TinhTrangHS"] == "Rút hồ sơ"
    assert storage.get_item(TUYEN_SINH.name, rec["id"])["TrangThai"] == "Rút hồ sơ"


def test_giay_to_can_nop_theo_khoi():
    can6 = services.giay_to_can_nop("6", "")
    assert "Hoàn thành chương trình tiểu học" in can6 and len(can6) == 5
    can10 = services.giay_to_can_nop("10", "ESL")
    assert "Đơn xét tuyển 10" in can10 and "Phiếu đăng ký nhập học" in can10
    assert services.giay_to_can_nop("12", "") == services.giay_to_can_nop("12", "ESL")
    can, da, thieu = services.tien_do_giay_to(
        {"Khoi": "6", "PhanHe": "IEP", "HoSoDaNop": "Ngày đã nộp 01/07/2026\nGiấy khai sinh - Bản sao"})
    assert da == {"Giấy khai sinh": "Bản sao"} and len(thieu) == 4


def test_danh_muc_giay_to_luu_va_doc(storage):
    from tuyensinh.schema import GIAY_TO
    services.luu_giay_to_khoi(storage, "6", ["Giấy khai sinh", "Học bạ Tiểu học"], [])
    services.luu_giay_to_khoi(storage, "7", ["Phiếu đăng ký nhập học"], [])
    items = storage.list_items(GIAY_TO.name)
    bang = services.giay_to_tu_list(items)
    assert bang == {"6": ["Giấy khai sinh", "Học bạ Tiểu học"], "7": ["Phiếu đăng ký nhập học"]}
    # sửa: đổi thứ tự + bớt 1 giấy tờ
    services.luu_giay_to_khoi(storage, "6", ["Học bạ Tiểu học"], items)
    bang = services.giay_to_tu_list(storage.list_items(GIAY_TO.name))
    assert bang["6"] == ["Học bạ Tiểu học"]
    services.dat_giay_to(bang)
    try:
        assert services.giay_to_can_nop("6", "IEP") == ["Học bạ Tiểu học"]
    finally:
        services.dat_giay_to(None)
    assert len(services.giay_to_can_nop("6", "IEP")) == 5  # mặc định Book1


def test_bien_nhan_html():
    from datetime import date as _d
    from tuyensinh import bien_nhan
    rec = {"HoTen": "Trần An", "NgaySinh": "2014-03-02", "Khoi": "6", "PhanHe": "IEP",
           "NoiTruBanTru": "Nội trú", "TenMe": "Lê Hoa", "DienThoaiMe": "0909123456"}
    page = bien_nhan.html(rec, {"Giấy khai sinh": "Bản sao"}, ["Học bạ Tiểu học"],
                          _d(2026, 10, 21), "Lê Thiên An", "2026-2027", _d(2026, 10, 1))
    for s in ("BẢNG KÝ NHẬN HỒ SƠ HỌC SINH", "Trần An", "02/03/2014", "Lê Hoa", "0909123456",
              "1 hồ sơ hợp lệ", "Giấy khai sinh", "(Bản sao)", "1 hồ sơ", "21/10/2026",
              "Lê Thiên An", "window.print()"):
        assert s in page, s
