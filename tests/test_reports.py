"""Mọi mẫu báo cáo chạy được trên dữ liệu mẫu và dữ liệu rỗng, xuất được Excel."""
import io
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from tuyensinh import reports, services  # noqa: E402
from tuyensinh.schema import NHAP_HOC, TUYEN_SINH  # noqa: E402

TS = [
    {"id": "1", "NamHoc": "2026-2027", "HoTenHS": "An", "Khoi": "10", "TrangThai": "Nhập học",
     "Nguon": "Hotline", "NgayLienHe": "2026-06-01", "CheDo": "Nội trú", "PhanHe": "IEP",
     "GiuCho": "Đã giữ chỗ", "SoTienXacNhan": 2000000, "TruongCu": "THCS A",
     "TruongCu_Tinh": "Tây Ninh", "NguoiNhanHoSo": "Cô Lan", "TinhTrang": "Nhập học"},
    {"id": "2", "NamHoc": "2026-2027", "HoTenHS": "Bình", "Khoi": "6", "TrangThai": "Tư vấn",
     "Nguon": "", "NgayLienHe": "2026-01-10", "CheDo": "Bán trú"},
    {"id": "3", "NamHoc": "2025-2026", "HoTenHS": "Cúc", "Khoi": "7", "TrangThai": "Rút hồ sơ",
     "NgayLienHe": "2025-05-01", "GiuCho": "Hủy giữ chỗ", "SoTienXacNhan": 1000000},
]
NH = [
    {"id": "9", "NamHoc": "2026-2027", "HoTen": "An", "Khoi": "10", "PhanHe": "IEP",
     "LopHoc": "10A1", "GioiTinh": "Nam", "NoiTruBanTru": "Nội trú", "DangKyXe": "Vũng Tàu",
     "DiemDonTra": "Chợ", "LuaChon1": "Lý-Hóa-Địa-Tin học", "HocLuc1": "Giỏi", "Toan1": 9,
     "SoTienThanhToan": 10000000, "TongDaThu": 4000000, "SoTienConLai": 6000000,
     "HoSoDaNop": "Ngày đã nộp 01/07/2026\nGiấy khai sinh - Bản sao",
     "ChoO_Tinh": "Thành phố Hồ Chí Minh", "ChoO_Xa": "Phường Tân Phú",
     "TinhTrangHS": "Đang đóng phí"},
    {"id": "10", "NamHoc": "2026-2027", "HoTen": "Dũng", "Khoi": "7", "LopHoc": "",
     "GioiTinh": "Nữ", "NoiTruBanTru": "Bán trú", "TinhTrangHS": "Đang nhập hồ sơ"},
]


def _ctx(ts, nh):
    ts_all = services.to_df(TUYEN_SINH, ts)
    return reports.Ctx(ts=ts_all[ts_all["NamHoc"] == "2026-2027"].reset_index(drop=True),
                       nh=services.to_df(NHAP_HOC, nh), ts_all=ts_all, nam_hoc="2026-2027",
                       params={"ky": "Tháng", "top": 10, "ngay": 14})


def test_moi_bao_cao_chay_va_xuat_excel():
    for ctx in (_ctx(TS, NH), _ctx([], [])):
        items = []
        for rep in reports.REPORTS:
            res = rep.fn(ctx)
            assert isinstance(res.table, pd.DataFrame), rep.id
            items += reports.export_items(rep, res)
        wb = load_workbook(io.BytesIO(reports.excel(items, "Trường Tân Phú", "2026-2027")))
        assert len(wb.sheetnames) >= len(reports.REPORTS)


def test_so_lieu():
    ctx = _ctx(TS, NH)
    t = reports.ts_khoi(ctx).table
    tong = t[t["Khối"] == "Tổng cộng"].iloc[0]
    assert tong["Tổng"] == 2 and tong["Nhập học"] == 1 and tong["Tỷ lệ nhập học (%)"] == 50.0
    assert reports.kt_congno(ctx).table["Còn lại (đ)"].tolist() == [6000000]
    assert reports.nh_xe(ctx).kpis[0] == ("HS đăng ký xe", 1)
    siso = reports.nh_siso(ctx).table
    assert siso[siso["Lớp"] == "10A1"]["Sĩ số"].tolist() == [1]
    assert len(reports.xh_tongquan(ctx).table) == 2


def test_xu_huong_cac_nam():
    ctx = _ctx(TS + [{"id": "4", "NamHoc": "2025-2026", "HoTenHS": "Dê", "Khoi": "10",
                      "TrangThai": "Nhập học", "Nguon": "Hotline", "NgayLienHe": "2025-06-15",
                      "TruongCu": "THCS A", "TruongCu_Tinh": "Tây Ninh"}], NH)
    t = reports.xh_tongquan(ctx).table
    assert t["Năm học"].tolist() == ["2025-2026", "2026-2027"]
    assert t["Liên hệ"].tolist() == [2, 2] and t["Liên hệ tăng/giảm"].tolist() == ["—", "0.0%"]
    k = reports.xh_khoi(ctx).table
    hang10 = k[k["Khối"] == "10"].iloc[0]
    assert hang10["2025-2026"] == 1 and hang10["2026-2027"] == 1
    # lọc năm + khối qua params
    ctx.params.update(nam=["2026-2027"], khoi=["10"])
    assert reports.xh_tongquan(ctx).table["Liên hệ"].tolist() == [1]
    ctx.params.update(nam=None, khoi=None)
    th = reports.xh_thang(ctx).table
    # N = năm khai giảng: 01/2026 của năm học 2026-2027 là "01/N"; 09/2025 sẽ là "09/N-1"
    assert th["Tháng"].tolist() == ["01/N", "05/N", "06/N"]  # xếp theo thời gian trong chu kỳ
    assert th["2025-2026"].tolist() == [0, 1, 1] and th["2026-2027"].tolist() == [1, 0, 1]
    assert reports.xh_truongcu(ctx).table.iloc[0]["Tổng"] == 2


def test_dia_ban_quan_huyen_cu():
    from tuyensinh import dia_ban
    assert dia_ban.quan_cu("Quận Tân Phú, Thành phố Hồ Chí Minh") == "Quận Tân Phú"
    assert dia_ban.quan_cu("Thành phố Thủ Đức, Thành phố Hồ Chí Minh") == "Thành phố Thủ Đức"
    assert dia_ban.quan_cu("Tỉnh Tây Ninh") == ""
    hcm = "Thành phố Hồ Chí Minh"
    ts = [{"id": str(i), "NamHoc": "2026-2027", "TrangThai": "Nhập học" if i < 2 else "Tư vấn",
           "Khoi": "6", "TruongCu_Tinh": hcm, "TruongCu_PhuongXa": "Phường Tân Phú",
           "TruongCu_DiaChiCu": "Quận Tân Phú, Thành phố Hồ Chí Minh" if i < 3 else ""}
          for i in range(4)]
    nh = [{"id": "9", "NamHoc": "2026-2027", "Khoi": "6", "TinhTrangHS": "Đang đóng phí",
           "ChoO_Tinh": hcm, "ChoO_Xa": "Phường Tân Phú"},
          {"id": "8", "NamHoc": "2025-2026", "Khoi": "6", "TinhTrangHS": "Học tiếp"},
          {"id": "7", "NamHoc": "2025-2026", "Khoi": "6", "TinhTrangHS": "Không học tiếp"}]
    ctx = _ctx(ts, nh)
    ctx.params["cap"] = dia_ban.CAP[1]
    t = reports.ts_diaban(ctx).table
    # liên hệ thứ 4 không có địa chỉ cũ nhưng suy ra được từ phường/xã mới
    assert t.iloc[0]["Quận/Huyện (cũ)"].startswith("Quận Tân Phú") and t.iloc[0]["Liên hệ"] == 4
    n = reports.nh_diaban(ctx).table
    assert n.iloc[0]["Quận/Huyện (cũ)"].startswith("Quận Tân Phú") and n.iloc[0]["Số HS"] == 1
    ctx.nh_all = ctx.nh
    tq = reports.xnh_tongquan(ctx).table.set_index("Năm học")
    assert tq.loc["2025-2026", "HS cũ học tiếp"] == 1
    assert tq.loc["2025-2026", "Tỷ lệ HS cũ ở lại (%)"] == 50.0
    assert tq.loc["2026-2027", "HS mới nhập học"] == 1
    dc = reports.xnh_doichieu(ctx).table.set_index("Năm học")
    assert dc.loc["2026-2027", "Nhập học theo Data tuyển sinh"] == 2
    assert dc.loc["2026-2027", "Nhập học nhưng chưa có hồ sơ"] == 2  # hồ sơ chưa liên kết


def test_tuyen_moi_cung_ky():
    from datetime import date
    ts = [{"id": "1", "NamHoc": "2025-2026", "Khoi": "6", "TrangThai": "Nhập học",
           "NgayLienHe": "2025-05-01"},
          {"id": "2", "NamHoc": "2025-2026", "Khoi": "6", "TrangThai": "Nhập học",
           "NgayLienHe": "2025-08-01"},   # sau mốc 30/06/2025 -> không tính
          {"id": "3", "NamHoc": "2026-2027", "Khoi": "6", "TrangThai": "Nhập học",
           "NgayLienHe": "2026-06-30"},
          {"id": "4", "NamHoc": "2026-2027", "Khoi": "10", "TrangThai": "Tư vấn",
           "NgayLienHe": "2026-03-01"},
          {"id": "5", "NamHoc": "2026-2027", "Khoi": "10", "TrangThai": "Nhập học"}]  # thiếu ngày
    ctx = _ctx(ts, [])
    ctx.params.update(moc=date(2026, 6, 30))
    t = reports.xh_cungky(ctx).table.set_index("Khối")
    assert t.loc["6", "Nhập học 2025-2026"] == 1 and t.loc["6", "Nhập học 2026-2027"] == 1
    assert t.loc["10", "Liên hệ 2026-2027"] == 1 and t.loc["10", "Nhập học 2026-2027"] == 0
    assert t.loc["Tổng cộng", "Liên hệ 2026-2027"] == 2


def test_kpi_bien_dong_khong_bao_sai():
    import pandas as pd
    from tuyensinh import reports as R
    from tuyensinh import dia_ban as DB
    nams = ["2025-2026", "2026-2027"]
    # chỉ một nơi, năm mới 0 vì chưa rõ quận: không ghi "tăng mạnh nhất -100%"
    t = pd.DataFrame({"2025-2026": [40, 2], "2026-2027": [0, 60]},
                     index=["Quận Tân Phú · TP. Hồ Chí Minh", DB.CHUA_RO + " · TP. Hồ Chí Minh"])
    k = R._kpi_dia_ban_nam(t, nams, "Quận/Huyện (cũ)")
    assert len(k) == 1 and "Chưa rõ quận/huyện" in k[0][0] and "Phường/Xã" in k[0][1]
    # năm mới chưa có số liệu
    assert "chưa có số liệu" in R._kpi_bien_dong(t.assign(**{"2026-2027": 0}), nams)[0][1]
    # chỉ có tăng: không có dòng "giảm mạnh nhất"; tăng / giảm khác nơi
    t2 = pd.DataFrame({"2025-2026": [10, 10, 10], "2026-2027": [15, 5, 10]}, index=["A", "B", "C"])
    assert R._kpi_bien_dong(t2, nams) == [("Tăng mạnh nhất", "A (+50.0%)"),
                                         ("Giảm mạnh nhất", "B (-50.0%)")]
    assert [k for k, _ in R._kpi_bien_dong(t2.iloc[[0]], nams)] == ["Tăng mạnh nhất"]
