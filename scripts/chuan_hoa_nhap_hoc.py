"""Chuẩn hóa "Data nhập học" cũ (xuất từ SharePoint) để nhập lên list Data_NhapHoc.

- Cột theo biểu mẫu VEMIS (tuyensinh/schema.py: NHAP_HOC) + các cột bổ sung của trường
- Địa chỉ 3 cấp cũ (Tỉnh – Quận/Huyện – Phường/Xã) → 2 cấp mới (Tỉnh/Thành – Phường/Xã)
  theo bảng sáp nhập 2025; ghi chú địa chỉ cũ ở bảng báo cáo
- Trường cũ đối chiếu danh mục trường; thiếu thì lấy từ Data_TuyenSinh đã chuẩn hóa
  (cùng năm học, họ tên và SĐT / ngày sinh)
- Khối "10-IEP" → Khối 10 + Phân hệ IEP; dân tộc, tôn giáo, chế độ theo danh mục VEMIS;
  CCCD mất số 0 đầu được bù lại
- Cột không có chỗ trong list (điểm, học phí, tài khoản nhận…) giữ ở sheet riêng

    python scripts/chuan_hoa_nhap_hoc.py data_nhaphoc.xlsx \\
        --tuyen-sinh Data_TuyenSinh_chuan_hoa.xlsx -o Data_NhapHoc_chuan_hoa.xlsx
"""
import argparse
import re
from collections import Counter

import _common  # noqa: F401
import pandas as pd

import chuan_hoa_du_lieu_cu as cc
from tuyensinh import danh_muc, services
from tuyensinh.schema import (BOOL, DATE, NHAP_HOC, NHAP_HOC_NGOAI_VEMIS, NOTE, NUMBER,
                              PHAN_HE, TINH_TRANG_HS)

fold = danh_muc.fold

# Cột trong file cũ → key trong app
COLS = {
    "Họ tên HS": "HoTen", "Ngày sinh": "NgaySinh", "Giới tính": "GioiTinh", "Lớp mới": "LopHoc",
    "Lớp cũ": "LopCu", "Mã PSC": "MaHocSinh", "CCCD": "CanCuoc", "Ngày cấp": "NgayCapCanCuoc",
    "Mã số BHYT": "MaBHYT", "Ngày nhập học": "NgayVaoTruong", "SĐT SMS": "DienThoaiSLL",
    "Email": "EmailSLL", "Dân tộc": "DanToc", "Tôn giáo": "TonGiao", "Nơi sinh": "NoiSinh",
    "Chế độ": "NoiTruBanTru", "Khối": "Khoi", "Phân hệ": "PhanHe", "Trạng thái": "TinhTrangHS",
    "Tên cha": "TenCha", "Nghề của cha": "NgheNghiepCha", "SĐT cha": "DienThoaiBo",
    "Số CCCD Cha": "CanCuocCha", "Năm sinh cha": "NamSinhCha", "Email cha": "EmailCha",
    "Tên mẹ": "TenMe", "Nghề của mẹ": "NgheNghiepMe", "SĐT mẹ": "DienThoaiMe",
    "Số CCCD mẹ": "CanCuocMe", "Năm sinh mẹ": "NamSinhMe", "Email Mẹ": "EmailMe",
    "Người giám hộ": "NguoiGiamHo", "Năm sinh NGH": "NamSinhNGH", "Nghề của NGH": "NgheNghiepNGH",
    "SĐT NGH": "DienThoaiNGH", "Số CCCD NGH": "CanCuocNGH", "Email NGH": "EmailNGH",
    "Tên liên hệ khẩn cấp": "KhanCap_Ten", "SĐT liên hệ khẩn cấp": "KhanCap_SDT",
    "Mối quan hệ": "KhanCap_QuanHe", "Hồ sơ đã nộp": "HoSoDaNop", "Ghi chú": "GhiChu",
    "Nam hoc": "NamHoc",
}
# Địa chỉ: (tiền tố key, cột số nhà, đường, tổ, ấp/khóm, xã/phường, quận/huyện, tỉnh)
DIA_CHI = [
    ("ChoO", "Số nhà LL", "Đường LL", "Tổ LL", "Ấp_Khóm LL", "Xã_Phường LL", "Quận_Huyện LL",
     "Tỉnh_Thành LL"),
    ("HK", "Số nhà ĐKHK", "Đường ĐKHK", "Tổ ĐKHK", "Ấp_Khóm ĐKHK", "Xã_Phường ĐKHK",
     "Quận_Huyện ĐKHK", "Tỉnh_Thành ĐKHK"),
]
TRUONG_CU = ("Trường cũ", "Tỉnh_Trường cũ", "Trường cũ_Quận_Huyện", "Trường cũ_Phường")
XE = (("Đăng ký xe", "Đăng ký xe trong năm"), ("ĐỊa điểm đón trả", "Địa điểm đón 1"))
DA_DUNG = set(COLS) | {c for d in DIA_CHI for c in d[1:]} | set(TRUONG_CU) | \
    {c for pair in XE for c in pair} | {"Title", "Item Type", "Path"}

TON_GIAO = {"thien chua": "Công giáo", "thien chua giao": "Công giáo", "kito giao": "Công giáo",
            "ki to giao": "Công giáo", "cong giao": "Công giáo", "phat": "Phật giáo",
            "dao phat": "Phật giáo", "tin lanh": "Tin Lành", "cao dai": "Cao Đài",
            "khong": "Không", "khong co": "Không"}
QUAN_HE = {"ba": "Cha", "bo": "Cha", "cha": "Cha", "me": "Mẹ", "nguoi giam ho": "Người giám hộ"}


def s(v) -> str:
    v = "" if v is None or (isinstance(v, float) and v != v) else str(v)
    return re.sub(r"\s+", " ", v).strip()


def phone(v: str) -> str:
    p = services.normalize_phone(v)
    return p if re.fullmatch(r"0\d{9,10}", p) else s(v)


def can_cuoc(v: str) -> str:
    d = re.sub(r"\D", "", v)
    if not d:
        return s(v)
    return "0" + d if len(d) == 11 else d  # Excel làm mất số 0 đầu


def nam_sinh(v: str) -> str:
    m = re.search(r"(19|20)\d{2}", v)
    return m.group(0) if m else v


def ngay(v: str):
    if not v:
        return None
    d = pd.to_datetime(v, errors="coerce", dayfirst=not re.match(r"\d{4}-", v))
    return d.date() if pd.notna(d) else None


def cap_dau(v: str) -> str:
    return v[:1].upper() + v[1:] if v else v


# ------------------------------------------------------------------ địa chỉ 2 cấp
def _xa_core(xa_raw: str) -> str:
    x = fold(xa_raw)
    x = re.sub(r"^(?:phuu+ong\b|p\.|p\b)\s*", "phuong ", x)  # "P.12", "P 12", "Phuường 14"
    x = re.sub(r"^x\.\s*", "xa ", x)
    core = cc.strip_prefix(x)
    return re.sub(r"^0+(\d)", r"\1", core)  # "Phường 08" -> "8"


def dia_chi_moi(geo: cc.Geo, tinh_raw: str, quan_raw: str, xa_raw: str, region: set):
    """(tỉnh mới, phường/xã mới, ghi chú) từ địa chỉ cũ 3 cấp hoặc địa chỉ mới 2 cấp."""
    tinh, pc = geo.tinh(tinh_raw) if tinh_raw else ("", None)
    dist = geo.huyen(quan_raw, pc, tinh, allow_ward=False) if quan_raw else None
    if dist and not tinh:
        pc, tinh = dist[0], geo.old_to_new.get(dist[0], "")
    core = _xa_core(xa_raw) if xa_raw else ""
    if not core:
        return tinh, "", "không ghi phường/xã" if tinh else "không có địa chỉ"
    if not tinh:  # đoán tỉnh: tên phường/xã mới duy nhất trong vùng
        found = {t for t in region if geo.new_ward_by_core.get((t, core))}
        if len(found) == 1:
            tinh = found.pop()
        else:
            return "", "", "không xác định được tỉnh"
    new = geo.new_ward_by_core.get((tinh, core), [])
    old = geo.new_from_old.get((tinh, core), [])
    if dist:
        dname = cc.strip_prefix(dist[1])
        in_dist = geo.wards_of_district.get(dist, set())
        # xã cũ ghi kèm đúng quận/huyện này trong bảng sáp nhập ("Phường 14 (Quận Tân Bình)")
        cands = {xa for xa, q in old if q == dname or (not q and (tinh, xa) in in_dist)}
        if len(cands) == 1:
            xa = cands.pop()
            return tinh, xa, "giữ tên (địa chỉ mới)" if xa in new else "quy đổi từ địa chỉ cũ"
        if len(cands) > 1:
            return tinh, "", ("phường/xã cũ nay chia cho nhiều phường/xã mới — chọn theo số "
                              "nhà: " + " / ".join(sorted(cands)))
        near = [x for x in new if (tinh, x) in in_dist]
        if len(near) == 1:  # tên phường/xã mới + quận/huyện cũ
            return tinh, near[0], "giữ tên (địa chỉ mới)"
    if len(new) == 1:
        return tinh, new[0], "giữ tên (địa chỉ mới)"
    cands = {xa for xa, _ in old}
    if len(cands) == 1:
        return tinh, cands.pop(), "quy đổi từ địa chỉ cũ (không ghi quận/huyện)"
    if cands:
        return tinh, "", "phường/xã cũ trùng tên ở nhiều quận/huyện — cần ghi quận/huyện"
    return tinh, "", "không tìm thấy phường/xã"


def so_nha(so: str, duong: str) -> str:
    """SN/Xóm = số nhà + tên đường ("Số 262/1/9" + "Tân Kỳ Tân Quý")."""
    so = re.sub(r"^(số|so)\s+", "", so, flags=re.I)
    return " ".join(x for x in (so, duong) if x)


def khu_dan_cu(to: str, ap: str) -> str:
    """Khu dân cư = tổ + ấp/khóm/khu phố ("Tôổ 43" -> "Tổ 43")."""
    to = re.sub(r"^t[ôoổ]+\s*(?=\d)", "Tổ ", to, flags=re.I)
    return ", ".join(x for x in (to, ap) if x)


# ------------------------------------------------------------------ chính
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("file")
    ap.add_argument("-o", "--out", default="Data_NhapHoc_chuan_hoa.xlsx")
    ap.add_argument("--tuyen-sinh", help="Data_TuyenSinh_chuan_hoa.xlsx (lấy trường cũ khi thiếu)")
    ap.add_argument("--old-units", help="JSON danh mục hành chính trước sáp nhập (v2.4.1)")
    ap.add_argument("--bando", help="bando_co_dvch.sql (bảng sáp nhập)")
    args = ap.parse_args()

    geo = cc.Geo(cc.load_old_units(args.old_units))
    geo.load_sapnhap(cc.load_text(args.bando, cc.BANDO_URL))
    schools = cc.Schools()
    raw = pd.read_excel(args.file, dtype=str).fillna("")
    raw.columns = [c.strip() for c in raw.columns]
    missing = [c for c in list(COLS) + list(TRUONG_CU[:3]) if c not in raw.columns]
    if missing:
        raise SystemExit(f"File thiếu cột: {missing}")

    catalog = danh_muc.load()
    dan_toc, ton_giao = catalog["dan_toc"], catalog["ton_giao"]
    ntbt = catalog["noi_tru_ban_tru"]
    region_count = Counter()
    for c in ("Tỉnh_Thành LL", "Tỉnh_Thành ĐKHK", "Tỉnh_Trường cũ"):
        region_count.update(geo.tinh(t)[0] for t in raw[c] if t)
    region = {t for t, n in region_count.items() if t and n >= 5}

    # Data tuyển sinh đã chuẩn hóa: tra trường cũ theo (năm học, họ tên, SĐT / ngày sinh)
    ts_by = {}
    if args.tuyen_sinh:
        ts = pd.read_excel(args.tuyen_sinh, sheet_name="Data_TuyenSinh", dtype=str).fillna("")
        for r in ts.itertuples(index=False):
            r = r._asdict()
            if not r.get("TruongCu"):
                continue
            name = fold(r["HoTenHS"])
            ts_by.setdefault(("sdt", name, services.normalize_phone(r["SDT"])), r)
            ts_by.setdefault(("ns", name, str(r["NgaySinh"])[:10]), r)

    out, report, extra = [], [], []
    stats = Counter()
    for i, r in raw.iterrows():
        o = {k: s(r[c]) for c, k in COLS.items()}
        rec = {}
        rec["NamHoc"] = o["NamHoc"]
        rec["HoTen"] = cc.person(o["HoTen"])
        rec["Title"] = rec["HoTen"]
        rec["NgaySinh"] = ngay(o["NgaySinh"])
        rec["GioiTinh"] = cc.canon(o["GioiTinh"], ("Nam", "Nữ"))
        m = re.match(r"\s*(\d{1,2})\s*[-–_ ]?\s*([A-Za-z]*)", o["Khoi"])
        rec["Khoi"] = str(int(m.group(1))) if m else o["Khoi"]
        rec["PhanHe"] = cc.canon((m.group(2) if m and m.group(2) else o["PhanHe"]).upper(), PHAN_HE)
        rec["LopHoc"] = "" if fold(o["LopHoc"]).startswith("chua") else o["LopHoc"]
        rec["LopCu"] = o["LopCu"]
        rec["TinhTrangHS"] = cc.canon(o["TinhTrangHS"], TINH_TRANG_HS)
        rec["MaHocSinh"] = o["MaHocSinh"]
        rec["NgayVaoTruong"] = ngay(o["NgayVaoTruong"])
        rec["DanToc"] = cc.canon(o["DanToc"], dan_toc)
        rec["QuocTich"] = "Việt Nam" if rec["DanToc"] in dan_toc else ""
        tg = cc.canon(o["TonGiao"], ton_giao, TON_GIAO)
        if tg and tg not in ton_giao:
            stats[f"Tôn giáo không có trong danh mục (bỏ trống): {tg}"] += 1
            tg = ""
        rec["TonGiao"] = tg
        rec["NoiTruBanTru"] = cc.canon(o["NoiTruBanTru"], ntbt)
        rec["CanCuoc"] = can_cuoc(o["CanCuoc"])
        rec["NgayCapCanCuoc"] = ngay(o["NgayCapCanCuoc"])
        rec["MaBHYT"] = re.sub(r"\s+", "", o["MaBHYT"])
        rec["HoSoDaNop"] = str(r["Hồ sơ đã nộp"]).strip()
        rec["GhiChu"] = str(r["Ghi chú"]).strip()
        # gia đình
        for who in ("Cha", "Me", "NGH"):
            k_ten = {"Cha": "TenCha", "Me": "TenMe", "NGH": "NguoiGiamHo"}[who]
            rec[k_ten] = cc.person(o[k_ten])
            rec[f"NgheNghiep{who}"] = cap_dau(o[f"NgheNghiep{who}"])
            rec[f"NamSinh{who}"] = nam_sinh(o[f"NamSinh{who}"])
            rec[f"CanCuoc{who}"] = can_cuoc(o[f"CanCuoc{who}"])
            rec[f"Email{who}"] = o[f"Email{who}"].lower()
        rec["DienThoaiBo"], rec["DienThoaiMe"] = phone(o["DienThoaiBo"]), phone(o["DienThoaiMe"])
        rec["DienThoaiNGH"] = phone(o["DienThoaiNGH"])
        rec["DienThoaiSLL"] = phone(o["DienThoaiSLL"])
        rec["EmailSLL"] = o["EmailSLL"].lower()
        rec["KhanCap_Ten"] = cc.person(o["KhanCap_Ten"])
        rec["KhanCap_SDT"] = phone(o["KhanCap_SDT"])
        qh = o["KhanCap_QuanHe"]
        rec["KhanCap_QuanHe"] = QUAN_HE.get(fold(qh), cap_dau(qh))
        rec["DangKyXe"] = s(r[XE[0][0]]) or s(r[XE[0][1]])
        rec["DiemDonTra"] = s(r[XE[1][0]]) or s(r[XE[1][1]])

        # --- địa chỉ 2 cấp
        rp = {"Dòng": i + 2, "Họ tên": rec["HoTen"], "Năm học": rec["NamHoc"]}
        for key, c_so, c_duong, c_to, c_ap, c_xa, c_quan, c_tinh in DIA_CHI:
            tinh_raw, quan_raw, xa_raw = s(r[c_tinh]), s(r[c_quan]), s(r[c_xa])
            tinh, xa, note = dia_chi_moi(geo, tinh_raw, quan_raw, xa_raw, region)
            rec[f"{key}_SoNha"] = so_nha(s(r[c_so]), s(r[c_duong]))
            rec[f"{key}_KhuDanCu"] = khu_dan_cu(s(r[c_to]), s(r[c_ap]))
            rec[f"{key}_Tinh"], rec[f"{key}_Xa"] = tinh, xa
            label = "Chỗ ở" if key == "ChoO" else "Hộ khẩu"
            rp[f"{label} (cũ)"] = ", ".join(x for x in (xa_raw, quan_raw, tinh_raw) if x)
            rp[f"{label} (mới)"] = ", ".join(x for x in (xa, tinh) if x)
            rp[f"{label} - kết quả"] = note
            stats[f"{label}: {note}"] += 1
        # nơi sinh: tỉnh/thành mới + ghi nguyên văn
        rec["NoiSinh_ThongTin"] = o["NoiSinh"]
        rec["NoiSinh_Tinh"] = geo.tinh(o["NoiSinh"])[0] if o["NoiSinh"] else ""
        if o["NoiSinh"] and not rec["NoiSinh_Tinh"]:
            stats["Nơi sinh: không nhận ra tỉnh (giữ ở Thông tin nơi sinh)"] += 1

        # --- trường cũ
        ten_raw = cc.clean_school(s(r["Trường cũ"]))
        t_tinh, t_pc = geo.tinh(s(r["Tỉnh_Trường cũ"]))
        q_raw = s(r["Trường cũ_Quận_Huyện"])
        dist = geo.huyen(q_raw, t_pc, t_tinh) if q_raw else None
        if dist and not t_tinh:
            t_tinh = geo.old_to_new.get(dist[0], "")
        wards = geo.wards_of_district.get(dist) if dist else None
        hit, how = schools.match(ten_raw, t_tinh, wards, rec["Khoi"], region) if ten_raw \
            else (None, "")
        name = fold(rec["HoTen"])
        ts_hit = ts_by.get(("sdt", name, services.normalize_phone(rec["DienThoaiSLL"]))) or \
            ts_by.get(("ns", name, str(rec["NgaySinh"] or "")[:10]))
        if hit:
            rec["TruongCu"], rec["TruongCu_Tinh"], rec["TruongCu_PhuongXa"] = \
                hit["ten"], hit["tinh"], hit["xa"]
            tc = "khớp danh mục" if how == "exact" else "khớp gần đúng — nên kiểm tra"
        elif ts_hit and (not ten_raw or school_same(ten_raw, ts_hit["TruongCu"])):
            rec["TruongCu"] = ts_hit["TruongCu"]
            rec["TruongCu_Tinh"] = ts_hit["TruongCu_Tinh"]
            rec["TruongCu_PhuongXa"] = ts_hit["TruongCu_PhuongXa"]
            tc = "lấy từ Data tuyển sinh"
        elif ten_raw:
            rec["TruongCu"] = cc.full_school_name(ten_raw)
            rec["TruongCu_Tinh"] = t_tinh
            rec["TruongCu_PhuongXa"] = geo.ward_from_name(cc.school_key(ten_raw)[0], t_tinh,
                                                          dist) if t_tinh else ""
            tc = "chưa có trong danh mục — giữ tên đã viết đầy đủ"
        else:
            rec["TruongCu"] = rec["TruongCu_Tinh"] = rec["TruongCu_PhuongXa"] = ""
            tc = "không có thông tin trường cũ"
        stats[f"Trường cũ: {tc}"] += 1
        rp.update({"Trường cũ (gốc)": s(r["Trường cũ"]),
                   "Trường cũ (gốc) - quận, tỉnh": ", ".join(
                       x for x in (q_raw, s(r["Tỉnh_Trường cũ"])) if x),
                   "Trường cũ (chuẩn)": rec["TruongCu"],
                   "Trường cũ - phường/xã, tỉnh": ", ".join(
                       x for x in (rec["TruongCu_PhuongXa"], rec["TruongCu_Tinh"]) if x),
                   "Trường cũ - kết quả": tc,
                   "Khối (gốc)": o["Khoi"], "Khối": rec["Khoi"], "Phân hệ": rec["PhanHe"]})
        report.append(rp)
        extra.append({"Họ tên HS": rec["HoTen"], "Năm học": rec["NamHoc"],
                      **{c: r[c] for c in raw.columns if c not in DA_DUNG and s(r[c])}})
        out.append(rec)

    for rec in out:
        stats["Có Chỗ ở (Tỉnh + Phường/Xã mới)" if rec["ChoO_Xa"] else "Thiếu Phường/Xã chỗ ở"] += 1
        miss = services.missing_fields(rec)
        stats["Đủ thông tin tối thiểu VEMIS" if not miss else "Còn thiếu thông tin tối thiểu"] += 1
    write(args.out, out, report, extra, stats)
    print(f"Đã ghi {args.out}: {len(out)} dòng")
    for k, v in sorted(stats.items()):
        print(f"  {k}: {v}")


def school_same(a: str, b: str) -> bool:
    ka, kb = cc.school_key(a)[0], cc.school_key(b)[0]
    return bool(ka) and (ka == kb or ka in kb or kb in ka)


# ------------------------------------------------------------------ xuất Excel
def columns():
    return ["Title"] + [f.key for f in NHAP_HOC.fields if f.key != "TuyenSinhID"]


HUONG_DAN = [
    ("Mục đích", "Dữ liệu hồ sơ nhập học cũ đã chuẩn hóa để nhập lên list Data_NhapHoc."),
    ("Cách nhập", "App → Hệ thống → Cài đặt & đồng bộ → 3. Nhập dữ liệu từ Excel → chọn file "
                  "này, List: Hồ sơ nhập học, tích 'Giữ nguyên dữ liệu cũ' → Nhập. App tự "
                  "liên kết hồ sơ với Data tuyển sinh (cùng năm học + họ tên + SĐT/ngày sinh)."),
    ("Địa chỉ", "Theo địa giới 2 cấp (Tỉnh/Thành – Phường/Xã, hiệu lực 01/07/2025). Địa chỉ "
                "cũ (có quận/huyện) được quy đổi theo bảng sáp nhập; đối chiếu ở sheet "
                "Bao_cao_chuan_hoa. Dòng không quy đổi được để trống Phường/Xã để bổ sung tay."),
    ("SN/Xóm", "Số nhà + tên đường. Khu dân cư = Tổ + Ấp/Khóm/Khu phố."),
    ("Trường cũ", "Tên chuẩn theo danh mục trường của app (trùng với gợi ý khi nhập)."),
    ("Cột ngoài VEMIS", ", ".join(k for k in NHAP_HOC_NGOAI_VEMIS if k != "TuyenSinhID")
     + " — dùng trong app, không xuất ra file VEMIS."),
    ("Cot_khac", "Các cột của file cũ không có trong list (điểm, học phí, tài khoản nhận hoàn "
                 "phí, nơi cấp CCCD cha mẹ…) — giữ để tra cứu."),
]


def write(path, rows, report, extra, stats):
    cols = columns()
    df = pd.DataFrame(rows).reindex(columns=cols)
    fields = {f.key: f for f in NHAP_HOC.fields}
    struct = pd.DataFrame([{
        "Tên nội bộ": k, "Tên hiển thị": fields[k].label if k in fields else "Tiêu đề (Họ tên)",
        "Kiểu": {DATE: "Ngày", NUMBER: "Số", NOTE: "Nhiều dòng", BOOL: "Có/Không"}.get(
            fields[k].type, "Văn bản / lựa chọn") if k in fields else "Văn bản",
        "Nhóm": fields[k].group if k in fields else "",
        "Có trong VEMIS": "" if k in NHAP_HOC_NGOAI_VEMIS or k == "Title" else "Có",
        "Số dòng có dữ liệu": int((df[k].astype(str).str.strip().replace("None", "") != "").sum()),
    } for k in cols])
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame(HUONG_DAN, columns=["Mục", "Nội dung"]).to_excel(
            xw, sheet_name="Huong_dan", index=False)
        df.to_excel(xw, sheet_name="Data_NhapHoc", index=False)
        struct.to_excel(xw, sheet_name="Cau_truc_cot", index=False)
        pd.DataFrame(report).to_excel(xw, sheet_name="Bao_cao_chuan_hoa", index=False)
        pd.DataFrame(sorted(stats.items()), columns=["Hạng mục", "Số dòng"]).to_excel(
            xw, sheet_name="Thong_ke", index=False)
        pd.DataFrame(extra).to_excel(xw, sheet_name="Cot_khac", index=False)
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        for ws in xw.book.worksheets:
            ws.freeze_panes = "A2" if ws.title != "Data_NhapHoc" else "C2"
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="1D4ED8")
                c.alignment = Alignment(vertical="center", wrap_text=True)
            for j, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 200)), 1):
                w = max((len(str(c.value)) for c in col if c.value is not None), default=8)
                ws.column_dimensions[get_column_letter(j)].width = min(max(10, w + 2), 60)
            ws.auto_filter.ref = ws.dimensions
        ws = xw.book["Data_NhapHoc"]
        for j, k in enumerate(cols, 1):
            if k in fields and fields[k].type == DATE:
                for (c,) in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                    c.number_format = "DD/MM/YYYY"
            else:  # CCCD, SĐT, mã: giữ dạng chữ (không mất số 0 đầu)
                for (c,) in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                    c.number_format = "@"


if __name__ == "__main__":
    main()
